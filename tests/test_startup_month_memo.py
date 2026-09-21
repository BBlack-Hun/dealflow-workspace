"""스타트업 명단의 **월별 칸 = 여러 줄 메모**, 머리글에 `(내용 기입)`.

사용자가 정한 것 그대로다.

> "스타트업 메뉴의 **월별 리마인드 컬럼의 이름을 `월별 리마인드(내용 기입)`**
>  이렇게 바꿔주고 **o,x 가 아니라 메모처럼 2-3줄의 메모**를 남길 수 있게
>  수정해줘"

세 칸(문자 · TEL · 카톡 연결)은 **그대로 둔다** — 셋을 하나로 합치면 그 달에
무엇을 한 것인지가 한 칸에 섞인다. 바뀌는 것은 **고르는 칸 → 글 칸**과
**머리글에 붙는 말**뿐이다.

## 여기서 잠그는 것 — 조용히 깨지는 순서대로

  1. **저장된 칸 이름에는 `(내용 기입)` 이 안 붙는다.** 이것이 제일 조용하다 —
     붙는 순간 시트를 다시 올릴 때 임포터가 그 칸을 못 찾아(`import_startup_
     sheet.py` 의 `if label not in columns`) **같은 달 칸이 두 벌** 서고, 값은
     옛 칸에 남아 화면에서 사라진 것처럼 보인다.
  2. **화면에는 옛 칸에도 붙는다.** `month_seed` 에 적었다면 달 칸이 하나도
     없는 새 명단에만 붙고, 지금 쓰는 명단에는 영영 안 붙었을 것이다.
  3. **고르는 칸이 아니다** — 보기가 뜨면 그 달에 있었던 일이 한 낱말로 줄어든다.
  4. **필터가 빠진다** — 줄마다 다른 글은 고를 것이 아니다. 필터를 열면 줄 수
     만큼 항목이 생긴다.
  5. **표 줄 높이가 안 무너진다** — 서른두 줄짜리 표에서 모든 칸이 통째로
     펴지면 훑을 수가 없다. 이 표에 이미 있는 길(`.clamp2` + `title` + 전문은
     수정창)을 그대로 쓴다.
  6. **이미 들어 있는 값이 그대로 보인다** — `O` 도 `7/24 o` 도 메모 칸에서는
     그냥 적힌 글이다. 안 보이면 사람은 지워진 줄 알고 다시 적는다.

이름·회사·번호는 전부 지어낸 값이다 — 저장소가 공개다.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest

from .conftest import DEMO_PASSWORD

ROOT = Path(__file__).resolve().parent.parent

LIST = "샘플 스타트업(9)"

# 사용자가 정한 꼬리말. **여기에 다시 적는 것이 이 검사의 전부다** — 앱에서
# 읽어 오면 앱이 바뀔 때 검사도 같이 바뀌어 아무것도 못 막는다.
SUFFIX = "(내용 기입)"

# 운영에 실제로 서 있는 칸 이름들(2026-09-21 개발 DB 실측). 시트마다 앞말이
# 제각각이라 **이름으로 갈라 보기를 붙이던 기전이 있었다** — 이제 셋 다 같은
# 메모 칸이라 가를 것이 없다. 폭을 재는 데 쓴다.
REAL_LABELS = [
    "9월 리마인드 문자",
    "9월 리마인드 TEL",
    "9월 카톡 연결",
    "7월 리마인드 문자 (7/28)",
    "카톡방 연결여부",
    # 가장 긴 이름. 팀이 정한 세 모양이 아니라 시트에서 딸려 온 칸인데,
    # **그 칸도 같은 메모 칸**이다.
    "리마인드 카톡(월1회-수or목or금)",
]

# 이미 들어 있는 값. 고르는 칸이던 시절의 `O` 와, 그 전부터 사람이 손으로 적던
# 글이 섞여 있다(개발 DB 실측: `O` 71칸 · `7/24 o` 13칸 · `o` 1칸).
OLD_VALUES = ["O", "7/24 o", "7/30 문자 및 명함 발송"]

# 사람이 새로 적는 메모. **세 줄**이다.
MEMO = "9/2 문자 발송\n9/5 부재중 — 다시 걸기로\n9/8 통화 완료, 자료 보내기로"


def _month(offset: int = 0) -> int:
    from app import clock

    return (clock.today().month - 1 + offset) % 12 + 1


def _url(sheet: str = LIST, show_all: bool = False) -> str:
    """그 명단의 화면.

    `show_all` 은 **접힌 달까지 펴서** 본다(`?months=all`). 표에는 최근 몇
    달만 서므로(`monthly_columns.VISIBLE_MONTHS`), 시트에서 딸려 온 옛 칸까지
    보려면 펴야 한다 — 거기가 꼬리말이 조용히 빠질 수 있는 자리다.
    """
    from urllib.parse import quote

    from app.services import contact_columns as cc

    tail = "&months=all" if show_all else ""
    return f"/{cc.page_of(cc.STARTUP)}?sheet={quote(sheet)}{tail}"


def _th(html: str, label: str) -> str:
    """그 이름을 가진 머리글의 속성. 없으면 `None`.

    빈 글자가 아니라 `None` 을 돌려준다 — 속성이 하나도 없는 머리글과
    "그런 머리글이 없다" 를 가려야 한다.
    """
    for attrs, cell in re.findall(r"<th\b([^>]*)>(.*?)</th>", html, re.S):
        flat = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", cell)).strip()
        if flat == label:
            return attrs
    return None


def _cell(html: str, key: str) -> str:
    """표에서 그 칸을 그린 `<div>` 통째로. 없으면 빈 글자."""
    m = re.search(r'<div class="[^"]*"\s+data-field="' + re.escape(key)
                  + r'"[^>]*>.*?</div>', html, re.S)
    return m.group(0) if m else ""


# ── 밑자리 ──────────────────────────────────────────────────────────────────

@pytest.fixture()
def sheets(client, db, users):
    """스타트업 명단 하나 — 팀이 정한 세 칸 + 시트에서 딸려 온 칸 하나.

    **옛 값을 일부러 심는다.** 메모 칸이 되었다고 그것이 화면에서 사라지면
    안 된다.
    """
    from app.models import ContactColumn, SheetOwner, VcContact
    from app.services import contact_columns as cc

    u1 = users["u1"]
    db.add(SheetOwner(label=LIST, user_id=u1.id, layout=cc.STARTUP, is_hidden=0))
    labels = [f"{_month()}월 리마인드 문자", f"{_month()}월 리마인드 TEL",
              f"{_month()}월 카톡 연결", "리마인드 카톡(월1회-수or목or금)"]
    for pos, label in enumerate(labels):
        db.add(ContactColumn(sheet=LIST, label=label, position=pos))
    db.flush()
    cols = {c.label: c for c in cc.month_columns(db, LIST)}

    db.add(VcContact(
        user_id=u1.id, source_sheet=LIST, name="김샘플1", firm="샘플기업1",
        phone="01000000101", email="sample1@example.com",
        notes=cc.dump_notes({
            cc.note_key(cols[labels[0]].id): OLD_VALUES[0],
            cc.note_key(cols[labels[1]].id): OLD_VALUES[1],
            cc.note_key(cols[labels[2]].id): OLD_VALUES[2],
        })))
    db.add(VcContact(
        user_id=u1.id, source_sheet=LIST, name="김샘플2", firm="샘플기업2",
        phone="01000000102", email="sample2@example.com"))
    db.commit()
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    return client


# ── 1. 저장된 이름에는 꼬리말이 안 붙는다 (제일 조용한 자리) ────────────────

def test_저장된_칸_이름에는_꼬리말이_안_붙는다(sheets, db):
    """**시트를 다시 올릴 때 그 글자로 칸을 찾는다.**

    임포터는 `ContactColumn.label` 과 시트 머리글을 **그대로** 맞춘다
    (`scripts/import_startup_sheet.py` — `if label not in columns`). 저장된
    이름에 `(내용 기입)` 이 붙으면 다음 업로드에서 같은 달 칸이 두 벌 서고,
    값은 옛 칸에 남아 화면에서 사라진 것처럼 보인다.
    """
    from app.services import contact_columns as cc

    for column in cc.month_columns(db, LIST):
        assert SUFFIX not in column.label, (
            f"저장된 칸 이름이 바뀌었습니다: {column.label}\n"
            "  ★ 시트를 다시 올리면 그 칸을 못 찾아 새 칸이 생깁니다")


def test_새_달에_저절로_생기는_칸도_시트_이름_그대로다(sheets, db):
    """새 달 칸은 직전 달 칸을 본떠 만든다(`monthly_columns.relabel`).

    본이 된 이름에 꼬리말이 실려 있으면 **그 꼬리말이 매달 다시 실린다** —
    `9월 리마인드 문자(내용 기입)(내용 기입)` 까지 가는 길이다.
    """
    from app.services import contact_columns as cc
    from app.services import monthly_columns

    today = date(2026, _month(), 15)
    nxt = date(today.year + (_month() == 12), _month(1), 1)
    made = monthly_columns.ensure_contact(db, LIST, today=nxt,
                                          seed=cc.STARTUP_LAYOUT.month_seed)
    assert made, "다음 달 칸이 만들어지지 않았습니다"
    for label in made:
        assert SUFFIX not in label, f"새 달 칸 이름에 꼬리말이 실렸습니다: {label}"


def test_밑자리로_심는_이름에도_꼬리말이_없다():
    """달 칸이 하나도 없는 명단에 처음 세우는 이름(`month_seed`)도 시트 모양이다.

    여기에 적었다면 **새 명단에만** 붙고 지금 쓰는 명단에는 영영 안 붙는다 —
    이미 칸이 선 명단은 직전 달 칸을 본떠 이름을 짓기 때문이다.
    """
    from app.services import contact_columns as cc

    for tail in cc.STARTUP_LAYOUT.month_seed:
        assert SUFFIX not in tail, f"밑자리에 꼬리말이 적혔습니다: {tail}"


# ── 2. 화면에는 **옛 칸에도** 붙는다 ────────────────────────────────────────

def test_표_머리글에_내용_기입이_붙는다(sheets, db):
    from app.services import contact_columns as cc

    html = sheets.get(_url(show_all=True)).text
    months = cc.month_columns(db, LIST)
    assert months, "월별 칸이 없습니다"
    for column in months:
        want = column.label + SUFFIX
        assert _th(html, want) is not None, (
            f"머리글이 `{want}` 로 안 섰습니다\n"
            f"  `{column.label}` 로 선 머리글 "
            f"{_th(html, column.label) is not None}")


def test_시트에서_딸려_온_옛_칸에도_같은_꼬리말이_붙는다(sheets, db):
    """팀이 정한 세 모양이 아닌 칸도 이 명단의 월별 칸이다.

    이름으로 갈라 붙이면 **가른 쪽만** 붙고, 시트마다 앞말이 다른 칸
    (`리마인드 카톡(월1회-…)`)은 조용히 빠진다.
    """
    html = sheets.get(_url(show_all=True)).text
    assert _th(html, "리마인드 카톡(월1회-수or목or금)" + SUFFIX) is not None, (
        "시트에서 딸려 온 칸에는 꼬리말이 안 붙었습니다")


def test_수정창_이름도_표_머리글과_같은_글자다(sheets, db):
    """두 곳에 따로 적어 두면 같은 칸인지 알아볼 수 없게 된다."""
    from app.services import contact_columns as cc

    html = sheets.get(_url()).text
    months = cc.month_columns(db, LIST)
    for column in months:
        key = cc.note_key(column.id)
        block = re.search(
            r"<label class=\"field[^\"]*\">\s*<span>([^<]*)</span>\s*"
            r"<textarea[^>]*data-note=\"" + re.escape(key) + r"\"", html)
        assert block, f"수정창에 `{column.label}` 칸이 없습니다"
        assert block.group(1) == column.label + SUFFIX, (
            f"수정창 이름이 표 머리글과 다릅니다: {block.group(1)}")


def test_칸_이름을_고치는_자리는_저장된_이름_그대로_보여_준다(sheets, db):
    """⚙ 명단·칸 설정에서 고치는 것은 **저장된 글자**다.

    거기까지 꼬리말이 실리면, 이름을 안 고치고 [저장]만 눌러도 저장된 이름에
    꼬리말이 박힌다 — 위 1번이 막는 그 사고가 화면 조작 한 번으로 난다.
    """
    from app.services import contact_columns as cc

    html = sheets.get(_url(show_all=True)).text
    for column in cc.month_columns(db, LIST):
        assert f'value="{column.label}"' in html, (
            f"칸 설정에 저장된 이름이 안 실렸습니다: {column.label}")
        assert f'value="{column.label}{SUFFIX}"' not in html, (
            f"칸 설정에 꼬리말이 실렸습니다 ★ [저장] 한 번에 이름이 박힙니다")


# ── 3. 고르는 칸이 아니다 ───────────────────────────────────────────────────

def test_월별_칸은_여러_줄_메모다():
    from app.services import contact_columns as cc

    layout = cc.STARTUP_LAYOUT
    assert layout.month_kind == "long", (
        "월별 칸이 여러 줄 칸이 아닙니다 ★ 그 달에 있었던 일이 한 낱말로 줄어듭니다")
    assert layout.month_choices == "", "월별 칸에 고를 거리가 남아 있습니다"
    assert layout.month_label_suffix == SUFFIX


def test_표에도_수정창에도_고를_거리가_안_뜬다(sheets, db):
    """`data-choices` 가 남아 있으면 칸을 눌렀을 때 옛 보기가 그대로 뜬다."""
    from app.services import contact_columns as cc

    html = sheets.get(_url(show_all=True)).text
    months = cc.month_columns(db, LIST)
    for column in months:
        key = cc.note_key(column.id)
        cell = _cell(html, key)
        assert cell, f"`{column.label}` 칸이 표에 없습니다"
        assert 'data-type="long"' in cell, (
            f"`{column.label}` 이 여러 줄 칸으로 안 섰습니다: {cell}")
        assert "data-choices" not in cell, (
            f"`{column.label}` 에 고를 거리가 남아 있습니다: {cell}")
        assert f'opts-panel-{key}' not in html, (
            f"수정창의 `{column.label}` 에 고를 거리가 남아 있습니다")

    # 배치가 내놓는 칸에서도 같은 것을 본다 — 화면만 고치고 배치가 그대로면
    # 엑셀·수정창처럼 같은 목록을 읽는 다른 자리가 옛 모양으로 남는다.
    keys = {cc.note_key(m.id) for m in months}
    for column in cc.table_columns(cc.STARTUP_LAYOUT, months):
        if column.key in keys:
            assert column.kind == "long" and not column.choices, (
                f"`{column.label}` 이 배치에서 아직 고르는 칸입니다")


# ── 4. 필터가 빠진다 ────────────────────────────────────────────────────────

def test_월별_칸에는_필터가_안_붙는다(sheets, db):
    """줄마다 다른 글은 고를 것이 아니다 — 필터를 열면 줄 수만큼 항목이 생긴다.

    선언(머리글 `data-filters`)과 행이 싣는 값(`data-f-*`) **둘 다** 빠져야
    한다. 한쪽만 빠지면 `tests/test_filter_columns.py` 가 잡는 어긋남이 된다.
    """
    from app.services import contact_columns as cc

    html = sheets.get(_url(show_all=True)).text
    months = cc.month_columns(db, LIST)
    filterable = {c.key for c in cc.filter_columns(cc.STARTUP_LAYOUT, months)}
    for column in months:
        key = cc.note_key(column.id)
        assert key not in filterable, f"`{column.label}` 에 필터가 남아 있습니다"
        attrs = _th(html, column.label + SUFFIX)
        assert "data-filters" not in (attrs or ""), (
            f"`{column.label}` 머리글에 필터 선언이 남아 있습니다: {attrs}")
        assert f"data-f-{key}=" not in html, (
            f"행이 `{column.label}` 값을 아직 싣습니다 ★ 아무도 안 보는 죽은 속성")


# ── 5. 표 줄 높이 ───────────────────────────────────────────────────────────

def test_표에서는_두_줄까지만_보이고_전문은_수정창에서_본다(sheets, db):
    """서른두 줄짜리 표에서 모든 줄이 통째로 펴지면 훑을 수가 없다.

    이 표에 **이미 있는 길**을 그대로 쓴다 — `메모 ( 통화내용 …)` 칸이
    `.clamp2`(두 줄) + `title`(마우스를 올려 전문) + 수정창의 `<textarea>` 로
    푼 그 문제다. 새 기전을 만들지 않는다.
    """
    from app.services import contact_columns as cc

    html = sheets.get(_url(show_all=True)).text
    for column in cc.month_columns(db, LIST):
        cell = _cell(html, cc.note_key(column.id))
        assert "clamp2" in cell, (
            f"`{column.label}` 이 두 줄로 안 접힙니다 ★ 줄 높이가 무너집니다: {cell}")
        assert "ellipsis" not in cell, (
            f"`{column.label}` 이 한 줄로 잘립니다 — 메모 칸이 아닙니다: {cell}")
    # 잘린 글은 마우스를 올려 읽는다. 칸을 감싼 `<td>` 가 그 값을 진다.
    assert html.count('<td title="') >= len(cc.month_columns(db, LIST))


def test_수정창은_세_줄짜리_메모_칸이다(sheets, db):
    """사용자가 말한 "2-3줄" 이 실제로 보이는 자리다.

    표의 `.clamp2` 는 두 줄까지만 **보여 주는** 것이고, 적는 것은 이 칸이다.
    `메모 ( 통화내용 …)` 와 같은 `rows="3"` 이라 같은 창에서 두 칸이 다른
    높이로 서지 않는다.
    """
    from app.services import contact_columns as cc

    html = sheets.get(_url()).text
    for column in cc.month_columns(db, LIST):
        key = cc.note_key(column.id)
        assert re.search(
            r'<textarea rows="3"[^>]*data-note="' + re.escape(key) + r'"', html), (
            f"수정창의 `{column.label}` 이 세 줄짜리 메모 칸이 아닙니다")


def test_머리글이_두_줄로_접히지_않는다():
    """이름이 길어졌다 — **폭을 다시 쟀다.**

    이 표의 머리글은 반복문으로 세워서 `tests/test_ui_layout.py` 의 정적
    검사가 조용히 건너뛴다. 운영에 실제로 서 있는 이름으로 같은 자를 댄다.
    필터 꼬리표(`(1) ▾`)는 안 잰다 — 이제 이 칸에는 필터가 안 붙는다.
    """
    from app.services import contact_columns as cc

    from .test_ui_layout import _text_px

    width = cc.STARTUP_LAYOUT.month_width
    problems = []
    for label in REAL_LABELS:
        need = round(_text_px(cc.month_label(cc.STARTUP_LAYOUT, label)) + 18)
        if need > width:
            problems.append(f"{label}{SUFFIX} ({width}px → {need}px 필요)")
    assert not problems, "머리글이 두 줄로 접힙니다:\n  " + "\n  ".join(problems)


# ── 6. 값 — 옛 것도 새 것도 ─────────────────────────────────────────────────

def test_이미_들어_있는_값이_글자로_그대로_보인다(sheets, db):
    """`O` 도 `7/24 o` 도 메모 칸에서는 그냥 적힌 글이다.

    잃는 것은 없지만 **보이지 않으면 사람은 지워진 줄 알고 다시 적는다** —
    그때는 이미 원래 무엇이 적혀 있었는지 화면 어디에서도 알 수 없다.
    """
    html = sheets.get(_url(show_all=True)).text
    for value in OLD_VALUES:
        assert f">{value}</div>" in html, f"옛 값이 표에서 사라졌습니다: {value}"


def test_여러_줄_메모가_저장되고_되읽힌다(sheets, db):
    """스키마 · 저장 · 되읽기 — 한 곳만 빠져도 증상이 조용하다.

    **줄바꿈이 그대로 살아야 한다.** 어딘가에서 공백으로 펴지면 세 줄로 적은
    메모가 한 줄로 뭉개진다.
    """
    from app.models import VcContact
    from app.services import contact_columns as cc

    row = db.query(VcContact).filter(VcContact.source_sheet == LIST).first()
    months = cc.month_columns(db, LIST)
    sent = {cc.note_key(m.id): MEMO for m in months}

    res = sheets.patch(f"/api/contacts/{row.id}", json={"notes": sent})
    assert res.status_code == 200, res.text
    got = sheets.get(f"/api/contacts/{row.id}").json()["contact"]["notes"]
    for key in sent:
        assert got.get(key) == MEMO, (
            f"`{key}` 의 여러 줄 메모가 되읽기에서 갈렸습니다: {got.get(key)!r}")


def test_엑셀에도_화면과_같은_이름으로_나간다(sheets, db):
    """엑셀은 **화면에서 쓰는 칸 그대로** 내보낸다(`data_io._export_sheet`).

    이 명단의 파일에는 이미 화면 이름이 실린다(`계약서 수신여부` — 시트는
    `계산서 수신`). 머리글만 다른 규칙으로 두면 같은 칸인지 알아볼 수 없다.
    """
    import io
    from urllib.parse import quote

    from openpyxl import load_workbook

    res = sheets.get(f"/api/export/contacts.xlsx?sheet={quote(LIST)}")
    assert res.status_code == 200, res.text
    book = load_workbook(io.BytesIO(res.content))
    head = [c.value for c in next(book.active.iter_rows(max_row=1))]
    got = [h for h in head if h and SUFFIX in h]
    assert len(got) >= 4, f"엑셀 머리글에 꼬리말이 안 실렸습니다: {head}"


# ── 7. 다른 화면이 안 깨진다 ────────────────────────────────────────────────

@pytest.mark.skipif(shutil.which("node") is None,
                    reason="node 미설치 — 브라우저 로직 테스트 생략")
def test_칸을_눌러_여러_줄을_치는_길이_실제로_도는가():
    """파이썬으로는 `data-type="long"` 이 붙어 있는지까지만 볼 수 있다.

    칸을 눌렀을 때 정말 `<textarea>` 가 뜨는지, 줄바꿈이 그대로 나가는지는
    `inline_edit.js` 를 돌려 봐야 한다.
    로컬에서는 `node tests/js/startup_month_memo_test.js` 로도 돈다.
    """
    script = ROOT / "tests" / "js" / "startup_month_memo_test.js"
    out = subprocess.run([shutil.which("node"), str(script)],
                         capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stdout + out.stderr
