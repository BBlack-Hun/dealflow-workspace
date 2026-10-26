"""`소유 펀드`(#1) — 회사 단위 값 · 같은 회사 줄 전파 · 엑셀 올리기 매칭 · 칸 자리.

회사 · 펀드 · 사람 이름은 전부 지어낸 것이다 — 저장소가 공개다.
"""
from __future__ import annotations

import io
import re

import pytest

from tests.conftest import DEMO_PASSWORD


def _contact(db, user_id, name, firm, sheet=None):
    from app.models import VcContact

    row = VcContact(user_id=user_id, name=name, firm=firm, source_sheet=sheet)
    db.add(row)
    db.commit()
    return row


@pytest.fixture()
def member(client, users):
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    return client


@pytest.fixture()
def admin(client, db, users):
    users["u2"].role = "admin"
    db.commit()
    client.post("/login", data={"phone": "01000000002", "password": DEMO_PASSWORD})
    return client


def _xlsx(rows):
    import openpyxl

    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── 1. 회사를 가르는 열쇠 ──────────────────────────────────────────────────

@pytest.mark.parametrize("a, b", [
    ("가상벤처스", "(주)가상벤처스"),
    ("가상벤처스", "가상 벤처스"),
    ("가상벤처스", "가상벤처스 주식회사"),
    ("Sample Partners", "sample partners"),
    ("㈜가상벤처스", "가상벤처스(주)"),
])
def test_같은_회사는_같은_열쇠다(a, b):
    from app.services import investor_firms as f

    assert f.firm_key(a) == f.firm_key(b)


def test_다른_회사는_다른_열쇠다():
    from app.services import investor_firms as f

    assert f.firm_key("가상벤처스") != f.firm_key("가상인베스트")
    assert f.firm_key("") == "" and f.firm_key(None) == ""


# ── 2. 한 줄에서 고치면 같은 회사의 모든 줄에 ─────────────────────────────

def test_한_줄에서_고치면_같은_회사의_다른_사람_줄에도_같은_값이_선다(member, db, users):
    from app.routers.contacts import contact_rows

    mine = _contact(db, users["u1"].id, "가나다", "(주)가상벤처스")
    also = _contact(db, users["u1"].id, "라마바", "가상 벤처스")
    other = _contact(db, users["u1"].id, "사아자", "가상인베스트")

    r = member.patch(f"/api/contacts/{mine.id}", json={"owned_funds": "가상 1호 조합"})
    assert r.status_code == 200, r.text
    assert r.json()["owned_funds"] == "가상 1호 조합"
    assert r.json()["firm_key"]

    rows = {row["name"]: row for row in contact_rows(db, users["u1"])}
    assert rows["가나다"]["owned_funds"] == "가상 1호 조합"
    assert rows["라마바"]["owned_funds"] == "가상 1호 조합"
    assert rows["사아자"]["owned_funds"] == ""
    # 수정창이 다시 열 때도 같은 값이다.
    assert member.get(f"/api/contacts/{also.id}").json()["contact"]["owned_funds"] == "가상 1호 조합"
    assert member.get(f"/api/contacts/{other.id}").json()["contact"]["owned_funds"] == ""


def test_회사_단위_값은_한_줄에만_저장된다(member, db, users):
    from app.models import InvestorFirm

    a = _contact(db, users["u1"].id, "가나다", "가상벤처스")
    _contact(db, users["u1"].id, "라마바", "가상벤처스")
    member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호"})
    member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 2호"})
    firms = db.query(InvestorFirm).all()
    assert len(firms) == 1
    assert firms[0].owned_funds == "가상 2호"
    assert firms[0].updated_by == users["u1"].id


def test_소유_펀드를_안_보내면_안_바뀐다(member, db, users):
    from app.services import investor_firms as f

    a = _contact(db, users["u1"].id, "가나다", "가상벤처스")
    member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호"})
    member.patch(f"/api/contacts/{a.id}", json={"memo": "통화 완료"})
    assert f.funds_of(db, "가상벤처스") == "가상 1호"


def test_회사명이_없는_줄에는_적을_수_없다(member, db, users):
    a = _contact(db, users["u1"].id, "가나다", None)
    r = member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호"})
    assert r.status_code == 400
    assert "투자사명" in r.json()["detail"]


def test_남의_줄은_소유_펀드도_못_고친다(member, db, users):
    theirs = _contact(db, users["u2"].id, "가나다", "가상벤처스")
    r = member.patch(f"/api/contacts/{theirs.id}", json={"owned_funds": "가상 1호"})
    assert r.status_code == 404


def test_새_줄을_넣으며_적어도_회사에_들어간다(member, db, users):
    from app.services import investor_firms as f

    r = member.post("/api/contacts", json={"name": "가나다", "firm": "가상벤처스",
                                           "owned_funds": "가상 3호"})
    assert r.status_code == 200, r.text
    assert f.funds_of(db, "가상 벤처스") == "가상 3호"


def test_고친_일은_수정_로그에_공용으로_남는다(member, db, users):
    from app.models import EditLog

    a = _contact(db, users["u1"].id, "가나다", "가상벤처스")
    member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호"})
    rows = [r for r in db.query(EditLog).all() if r.table_name == "investor_firms"]
    assert rows and rows[0].scope == "shared"


# ── 3. 표 · 수정창 · 엑셀에서 칸의 자리 ─────────────────────────────────────

def test_표에서_라운드_사이즈_바로_오른쪽에_선다(member, db, users):
    _contact(db, users["u1"].id, "가나다", "가상벤처스")
    html = member.get("/contacts?sheet=all").text
    head = html[html.index('id="contacts-table"'):html.index("</thead>")]
    names = [re.sub(r"<[^>]+>", "", cell).strip()
             for cell in re.findall(r"<th\b[^>]*>(.*?)</th>", head, re.S)]
    at = names.index("라운드 사이즈(투자운영금액)")
    assert names[at + 1] == "소유 펀드"
    # 줄에서도 같은 자리다 — 머리글과 칸이 한 칸이라도 밀리면 안 된다.
    row = re.search(r'id="contacts-table".*?<tbody>(.*?)</tr>', html, re.S).group(1)
    fields = re.findall(r'data-field="([a-z_]+)"', row)
    assert fields[fields.index("round_size") + 1] == "owned_funds"
    assert 'data-firm-key="' in row


def test_달마다_늘어나는_명단의_수정창에도_라운드_사이즈_뒤에_선다():
    from app.services import contact_columns as cc

    keys = [c.key for c in cc.INVESTOR_MONTHLY_LAYOUT.extra]
    assert keys[keys.index("round_size") + 1] == "owned_funds"


def test_엑셀_내려받기에도_실린다(member, db, users):
    import openpyxl

    a = _contact(db, users["u1"].id, "가나다", "가상벤처스")
    member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호"})
    r = member.get("/api/export/contacts.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    head = [c.value for c in ws[1]]
    at = head.index("소유 펀드")
    assert head[at - 1] == "라운드 사이즈(투자운영금액)"
    values = [row[at] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "가상 1호" in values


# ── 4. 엑셀 올리기 ─────────────────────────────────────────────────────────

def test_머리글은_이름으로_찾는다():
    from app.services import investor_firms as f

    rows = [["메모", ""], ["", "소유 펀드", "투자사명"], ["", "가상 1호", "가상벤처스"],
            ["", "", ""]]
    parsed = f.parse_upload(rows)
    assert [(p.name, p.funds) for p in parsed] == [("가상벤처스", "가상 1호")]
    with pytest.raises(f.UploadError):
        f.parse_upload([["회사", "펀드"]])


def test_맞춘_회사와_못_맞춘_이름을_가른다(db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "가상벤처스")
    _contact(db, users["u2"].id, "라마바", "가상벤처스")
    _contact(db, users["u1"].id, "사아자", "가상인베스트")
    f.set_funds(db, "가상인베스트", "기존 펀드", users["u1"].id)
    db.commit()
    rows = f.parse_upload([
        ["투자사명", "소유 펀드"],
        ["(주)가상벤처스", "가상 1호"],
        ["없는투자사", "가상 2호"],
        ["가상인베스트", ""],
        ["가상 인베스트", "기존 펀드"],
    ])
    plan = f.match(db, rows).as_dict()
    assert [m["firm"] for m in plan["matched"]] == ["가상벤처스", "가상인베스트"]
    assert plan["matched"][0]["rows"] == 2 and plan["matched"][0]["changes"]
    assert plan["matched"][1]["changes"] is False      # 이미 같은 값
    assert [u["name"] for u in plan["unmatched"]] == ["없는투자사"]
    assert [b["name"] for b in plan["blank"]] == ["가상인베스트"]
    assert plan["counts"]["changes"] == 1


def test_스타트업_명단의_기업_이름에는_맞추지_않는다(db, users):
    from app.models import SheetOwner
    from app.services import contact_columns as cc
    from app.services import investor_firms as f

    db.add(SheetOwner(label="가상 스타트업 명단", user_id=users["u1"].id,
                      layout=cc.STARTUP))
    db.commit()
    _contact(db, users["u1"].id, "대표", "가상스타트업", sheet="가상 스타트업 명단")
    rows = f.parse_upload([["투자사명", "소유 펀드"], ["가상스타트업", "가상 1호"]])
    plan = f.match(db, rows).as_dict()
    assert plan["matched"] == []
    assert [u["name"] for u in plan["unmatched"]] == ["가상스타트업"]


def test_같은_회사가_두_번이면_아래_줄을_쓴다(db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "가상벤처스")
    rows = f.parse_upload([["투자사명", "소유 펀드"], ["가상벤처스", "옛 값"],
                           ["가상 벤처스", "새 값"]])
    plan = f.match(db, rows).as_dict()
    assert [m["after"] for m in plan["matched"]] == ["새 값"]
    assert plan["duplicates"][0]["kept_line"] == 3


def test_미리보기는_안_쓰고_반영은_회사마다_바꾼다(admin, db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "가상벤처스")
    f.set_funds(db, "가상벤처스", "옛 펀드", users["u1"].id)
    db.commit()
    data = _xlsx([["투자사명", "소유 펀드"], ["가상벤처스", "가상 1호, 가상 2호"],
                  ["없는투자사", "가상 3호"]])
    files = {"file": ("funds.xlsx", data,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = admin.post("/api/import/owned-funds", files=files, data={"dry_run": "true"})
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["counts"]["changes"] == 1 and body["counts"]["unmatched"] == 1
    assert body["matched"][0]["before"] == "옛 펀드"
    db.expire_all()
    assert f.funds_of(db, "가상벤처스") == "옛 펀드", "미리보기가 썼다"

    done = admin.post("/api/import/owned-funds", files=files, data={"dry_run": "false"})
    assert done.status_code == 200, done.text
    assert done.json()["changed"] == 1
    db.expire_all()
    assert f.funds_of(db, "가상벤처스") == "가상 1호, 가상 2호"
    assert f.funds_of(db, "없는투자사") == "", "못 맞춘 이름을 넣었다"


def test_올리기와_양식은_관리자만(member, db, users):
    data = _xlsx([["투자사명", "소유 펀드"], ["가상벤처스", "가상 1호"]])
    files = {"file": ("funds.xlsx", data, "application/octet-stream")}
    assert member.post("/api/import/owned-funds", files=files).status_code == 403
    assert member.get("/api/sample/owned-funds.xlsx").status_code == 403
    # 화면의 단추도 안 선다.
    assert 'id="funds-btn"' not in member.get("/contacts").text


def test_양식은_두_칸_머리글이다(admin):
    import openpyxl

    r = admin.get("/api/sample/owned-funds.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:2] == ["투자사명", "소유 펀드"]
    assert 'id="funds-btn"' in admin.get("/contacts").text


# ── 5. 소유 펀드 필터 (펀드마다 줄 수) ─────────────────────────────────────

@pytest.mark.parametrize("text, funds", [
    ("가상 1호, 가상 2호", ["가상 1호", "가상 2호"]),
    ("  가상   1호 ,가상 2호  ,", ["가상 1호", "가상 2호"]),      # 앞뒤·겹친 공백 · 끝 쉼표
    ("가상 1호\n가상 2호", ["가상 1호", "가상 2호"]),           # 엑셀 칸 안 줄바꿈
    ("가상 1호, 가상  1호", ["가상 1호"]),                      # 다듬은 뒤 같은 이름은 한 번
    ("Sample Fund I, sample fund i", ["Sample Fund I", "sample fund i"]),  # 글자는 안 바꾼다
    ("", []), (None, []), (" , ", []),
])
def test_펀드_이름은_쉼표로_나누고_공백만_다듬는다(text, funds):
    from app.services import investor_firms as f

    assert f.split_funds(text) == funds


def test_줄마다_자기_회사의_펀드를_필터_값으로_싣는다(member, db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "가상벤처스")
    _contact(db, users["u1"].id, "라마바", "가상 벤처스")
    _contact(db, users["u1"].id, "사아자", "가상인베스트")
    _contact(db, users["u1"].id, "차카타", "가상캐피탈")
    f.set_funds(db, "가상벤처스", "가상 1호,  가상 2호", users["u1"].id)
    f.set_funds(db, "가상인베스트", "가상 2호", users["u1"].id)
    db.commit()

    html = member.get("/contacts?sheet=all").text
    assert 'data-filters="fund:소유 펀드"' in html
    funds = dict(re.findall(r'data-name="([^"]+)"[^>]*?data-f-fund="([^"]*)"', html, re.S))
    assert funds == {"가나다": "가상 1호|가상 2호", "라마바": "가상 1호|가상 2호",
                     "사아자": "가상 2호", "차카타": ""}


def test_한_줄에서_고치면_응답이_나눈_펀드도_싣는다(member, db, users):
    a = _contact(db, users["u1"].id, "가나다", "가상벤처스")
    r = member.patch(f"/api/contacts/{a.id}", json={"owned_funds": "가상 1호, 가상 2호"})
    assert r.status_code == 200, r.text
    assert r.json()["owned_fund_list"] == ["가상 1호", "가상 2호"]


def test_필터_건수와_고친_뒤_같은_회사_줄_따라가기_node():
    import shutil
    import subprocess
    from pathlib import Path

    node = shutil.which("node")
    if node is None:
        pytest.skip("node 미설치 — `node tests/js/contacts_fund_filter_test.js`")
    script = Path(__file__).resolve().parent / "js" / "contacts_fund_filter_test.js"
    out = subprocess.run([node, str(script)], capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stdout + out.stderr
