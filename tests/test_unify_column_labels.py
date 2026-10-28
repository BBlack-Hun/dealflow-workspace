"""#63 1단계 — **같은 칸은 같은 이름으로**(사용자 결정 2026-10-20 "추천대로 ㄱㄱ").

같은 저장 칸이 화면마다 다른 이름으로 서 있었다. 투자사 쪽 칸은 투자사 관리 현황의
말로, 기업 쪽 칸은 IR 기업 현황의 말로 맞춘다. **값 · 동기화 범위는 그대로**다 —
이름만 바뀐다.

  ① 딜 소싱 머리글 — `선호 투자분야` · `라운드 사이즈(투자운영금액)` · `전자 메일 주소`
     (#87 · 10/29 로 화면은 두 화면 다 기업 쪽과 같은 `이메일` — 엑셀 머리글은 그대로).
     이어진 줄 표시(`SHARED_LABELS`)도 같은 말.
  ② 딜 소싱 분야 칸은 투자사 관리 현황과 **같은 고르는 부품**(`data-type="sectors"`).
  ③ 투자사 딜공유 배치 — `회사` · `라운드 사이즈(투자운영금액)`.
  ④ 엑셀 — 내보내기 · 샘플 양식이 새 이름을 쓰고, **올리기는 옛 이름도 받는다.**
  ⑤ Series A 이상 탭 — `연락처` · `Pre Value`(스타트업DB 의 말).
  ⑥ 옛 이름으로 자리를 고른 추가 칸은 새 이름 칸 뒤에 그대로 선다
     (`tests/test_custom_columns.py` 에 함께 둔다 — 그쪽 도구를 쓴다).

이름 · 회사 · 번호는 전부 지어낸 값이다(공개 저장소).
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import pytest

from .conftest import DEMO_PASSWORD

ROOT = Path(__file__).resolve().parent.parent / "app"
CONTACTS_HTML = (ROOT / "templates" / "contacts.html").read_text(encoding="utf-8")

SECTORS = "선호 투자분야"
ROUND = "라운드 사이즈(투자운영금액)"
# 엑셀(내려받기 · 올리기) 머리글은 옛 이름 그대로다 — 받아 쓰는 수식이 있다.
EMAIL = "전자 메일 주소"
# 화면 머리글은 #87(10/29) 로 기업 쪽과 같은 `이메일`.
EMAIL_SCREEN = "이메일"
STAGES = "선호 투자단계"
FIRM = "회사"


def _contacts_heads():
    """투자사 관리 현황(투자사 명함) 표의 머리글 이름 — 템플릿에 손으로 적혀 있다."""
    return [re.sub(r"<[^>]+>", "", m).strip()
            for m in re.findall(r"<th[^>]*>(.*?)</th>", CONTACTS_HTML, re.S)]


@pytest.fixture()
def seeded(client, db, users):
    from app.models import SourcingContact

    db.add(SourcingContact(bucket="시리즈 A 이상", position=0, name="홍길동",
                           title="심사역", firm="가나벤처스",
                           sectors="AI, 옛날에 적은 분야", round_size="10~30억"))
    db.commit()
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    return client


# ── ① 딜 소싱 머리글 ─────────────────────────────────────────────────────────

def test_딜_소싱의_같은_칸은_투자사_관리_현황의_말이다():
    from app.routers.sourcing import COLUMNS

    labels = {field: label for field, label, _w, _f in COLUMNS}
    assert labels["sectors"] == SECTORS
    assert labels["round_size"] == ROUND
    assert labels["email"] == EMAIL_SCREEN
    # 투자사 관리 현황 표에 **그 글자 그대로** 서 있는 이름이어야 한다.
    heads = _contacts_heads()
    for word in (SECTORS, ROUND, EMAIL_SCREEN, FIRM, "휴대폰", "직함"):
        assert word in heads, f"투자사 관리 현황 머리글에 '{word}' 가 없다: {heads}"
    for field in ("firm", "phone", "title"):
        assert labels[field] in heads


def test_딜_소싱_화면에_새_이름이_서고_옛_이름은_머리글에_없다(seeded):
    html = seeded.get("/sourcing").text
    heads = [re.sub(r"<[^>]+>", "", m).strip()
             for m in re.findall(r"<th[^>]*>(.*?)</th>", html, re.S)]
    for word in (SECTORS, ROUND, EMAIL_SCREEN):
        assert word in heads, heads
    for old in ("투자분야", "라운드 사이즈", "이메일 주소", EMAIL):
        assert old not in heads, f"옛 머리글 '{old}' 이 남았다"
    # 필터 선언도 새 이름 — 단추 글자가 이 이름에서 선다(`filters.js`).
    assert f'data-filters="sectors:{SECTORS}"' in html
    assert f'data-filters="round_size:{ROUND}"' in html


def test_이어진_줄_표시도_같은_말이다():
    """`🔗 투자사 관리 현황` 풍선말 · `다른 칸` 이 딜 소싱 머리글과 같은 이름을 쓴다."""
    from app.routers.sourcing import COLUMNS
    from app.services import sourcing_sync

    labels = {field: label for field, label, _w, _f in COLUMNS}
    for field in sourcing_sync.SHARED_FIELDS:
        assert sourcing_sync.SHARED_LABELS[field] == labels[field], field


def test_넓힌_칸이_한_줄에_들고_표_폭도_같이_올랐다():
    """분야 · 라운드 칸 폭은 투자사 관리 현황의 같은 칸과 같다(134 · 204px) — 필터
    단추(`선호 투자분야 (1) ▾`)와 긴 이름이 한 줄에 든다."""
    from app.routers.sourcing import COLUMNS

    widths = {field: width for field, _l, width, _f in COLUMNS}
    assert widths["sectors"] == "134px"
    assert widths["round_size"] == "204px"
    assert 'style="width:134px" data-filters="sector:선호 투자분야"' in CONTACTS_HTML
    assert 'style="width:204px" data-filters="round:라운드 사이즈(투자운영금액)"' in CONTACTS_HTML


# ── ② 분야 칸은 같은 고르는 부품 ─────────────────────────────────────────────

def test_딜_소싱_분야_칸은_분야_기준_목록에서_고른다(seeded):
    html = seeded.get("/sourcing").text
    cell = re.search(r'<div class="cell[^"]*" data-field="sectors"[^>]*>', html).group(0)
    assert 'data-type="sectors"' in cell
    # **구분자는 그대로 안 준다** — 옛 자유 글에는 괄호 안 쉼표가 있다
    # (`tests/test_filter_columns.py`).
    assert "data-filter-sep" not in cell
    # 고를 거리 · 부품 — 투자사 관리 현황과 같은 것. 부품이 칸 고치기보다 먼저.
    assert 'id="sector-taxonomy"' in html
    assert html.index("js/sector_picker.js") < html.index("js/inline_edit.js")


def test_목록_밖_옛_값은_칸에_그대로_남는다(seeded):
    """자유 글로 적던 값은 지우지 않는다 — 고르는 창이 `목록 밖 값` 으로 보여 준다
    (`sector_picker.js`). 화면은 적힌 그대로다."""
    html = seeded.get("/sourcing").text
    assert "AI, 옛날에 적은 분야" in html


def test_고른_값은_그대로_저장되고_이어진_규칙도_그대로다(seeded, db):
    """부품이 보내는 값은 쉼표로 이은 글자 하나다 — 저장 칸 · API 는 그대로."""
    from app.models import SourcingContact

    row = db.query(SourcingContact).filter_by(name="홍길동").one()
    r = seeded.patch(f"/api/sourcing/{row.id}", json={"sectors": "AI, 핀테크"})
    assert r.status_code == 200, r.text
    db.expire_all()
    assert db.get(SourcingContact, row.id).sectors == "AI, 핀테크"


# ── ③ 투자사 딜공유 배치 ─────────────────────────────────────────────────────

def test_투자사_딜공유_배치도_같은_말이다():
    from app.services import contact_columns as cc

    layout = cc.INVESTOR_MONTHLY_LAYOUT
    labels = {c.key: c.label for c in [*layout.head, *layout.tail, *layout.extra]}
    assert labels["firm"] == FIRM
    assert labels["round_size"] == ROUND
    assert labels["sectors"] == SECTORS
    assert labels["email"] == EMAIL_SCREEN
    # 사람이 주인공이라는 판정은 이름이 아니라 차례에서 읽는다 — 그대로다.
    assert cc.firm_leads(cc.INVESTOR_MONTHLY) is False


# ── ④ 엑셀 ───────────────────────────────────────────────────────────────────

def test_엑셀_내려받기_머리글이_새_이름이다(client, users, db):
    from app.models import VcContact
    from app.routers import data_io

    heads = data_io.CONTACT_HEADERS
    for word in (FIRM, ROUND, STAGES, SECTORS, EMAIL):
        assert word in heads, word
    for old in ("투자사", "라운드 규모", "선호 단계", "선호 분야"):
        assert old not in heads, old
    assert FIRM in data_io.JOB_HEADERS and "투자사" not in data_io.JOB_HEADERS

    db.add(VcContact(user_id=users["u1"].id, name="홍길동", firm="가나벤처스",
                     round_size="10~30억", stages="SeriesA", sectors="AI"))
    db.commit()
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    import openpyxl

    r = client.get("/api/export/contacts.xlsx")
    assert r.status_code == 200
    rows = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(
        values_only=True))
    head, body = list(rows[0]), list(rows[1])
    assert body[head.index(FIRM)] == "가나벤처스"
    assert body[head.index(ROUND)] == "10~30억"
    assert body[head.index(STAGES)] == "SeriesA"
    assert body[head.index(SECTORS)] == "AI"


def test_샘플_양식도_새_이름이다():
    from app.routers import data_io

    head = data_io.SAMPLE_HEADER_ROW
    assert FIRM in head and ROUND in head and SECTORS in head
    assert "투자사명" not in head and "라운드 규모" not in head
    # 칸 수가 맞아야 양식이 밀리지 않는다.
    assert len(data_io.SAMPLE_WIDTHS) == len(head)
    assert all(len(r) == len(head) for r in data_io.SAMPLE_ROWS)


OLD_HEAD = ["NO", "그룹", "이름", "직책", "투자사명", "부서", "담당자", "휴대폰",
            "이메일", "선호 분야", "라운드 규모", "선호 단계", "메모"]
NEW_HEAD = ["NO", "그룹", "이름", "직함", "회사", "부서", "담당자", "휴대폰",
            "전자 메일 주소", "선호 투자분야", "라운드 사이즈(투자운영금액)",
            "선호 투자단계", "메모"]
ROW = ["1", "A", "홍길동", "심사역", "가나벤처스", "투자1본부", "김담당", "010-0000-0001",
       "hong@example.com", "AI, 헬스케어", "10~30억", "SeriesA", "샘플 메모"]


@pytest.mark.parametrize("head", [OLD_HEAD, NEW_HEAD,
                                  # `NO` 칸 없이 `이름` + `회사` 만으로도 머리글을 찾는다.
                                  [h for h in NEW_HEAD if h != "NO"]],
                         ids=["old", "new", "new-no-NO"])
def test_올리기는_옛_이름도_새_이름도_받는다(head):
    from app.services import sheet_import

    row = ROW if "NO" in head else ROW[1:]
    parsed = sheet_import.parse_sheet_a([head, row], 2026)
    assert parsed.skipped == []
    (c,) = parsed.contacts
    assert c.firm == "가나벤처스"
    assert c.title == "심사역"
    assert c.email == "hong@example.com"
    assert c.sectors == ["AI", "헬스케어"]
    assert c.round_size == "10~30억"
    assert c.stages == "SeriesA"
    assert c.phone == "010-0000-0001"


@pytest.mark.parametrize("head", [OLD_HEAD, NEW_HEAD], ids=["old", "new"])
def test_올리기_화면에서도_두_이름이_다_반영된다(client, users, db, head):
    """파서만이 아니라 [현황 업로드] 길 전체 — 반영까지 한다."""
    import openpyxl

    from app.models import VcContact

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "딜소개 현황"
    ws.append(head)
    ws.append(ROW)
    buf = io.BytesIO()
    wb.save(buf)
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    r = client.post(
        "/api/import/contacts",
        files={"file": ("현황.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"dry_run": "false", "year": "2026"})
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 1, r.json()
    db.expire_all()
    c = db.query(VcContact).filter_by(name="홍길동").one()
    assert (c.firm, c.email, c.round_size) == ("가나벤처스", "hong@example.com", "10~30억")
    assert [s.strip() for s in (c.sectors or "").split(",")] == ["AI", "헬스케어"]


def test_비교_표의_칸_이름도_같은_말이다():
    """[현황 업로드] 미리보기의 `다른데 안 덮은 칸` 이 부르는 이름."""
    from app.services import sheet_import

    assert sheet_import._DIFF_LABELS["round_size"] == ROUND
    assert sheet_import._DIFF_LABELS["email"] == EMAIL_SCREEN
    assert sheet_import._DIFF_LABELS["stages"] == STAGES
    assert sheet_import._DIFF_LABELS["sectors"] == SECTORS


# ── ⑤ Series A 이상 탭 ──────────────────────────────────────────────────────

def test_Series_A_탭은_스타트업DB_의_말이다():
    html = (ROOT / "templates" / "companies_series_a.html").read_text(encoding="utf-8")
    db_tab = (ROOT / "templates" / "companies.html").read_text(encoding="utf-8")
    heads = [re.sub(r"<[^>]+>", "", m).strip()
             for m in re.findall(r"<th[^>]*>(.*?)</th>", html, re.S)]
    assert "연락처" in heads and "Pre Value" in heads
    assert "핸드폰" not in heads and "프리밸류" not in heads
    # 스타트업DB 탭이 같은 칸을 그 이름으로 부른다.
    assert ">연락처</th>" in db_tab and ">Pre Value</th>" in db_tab
