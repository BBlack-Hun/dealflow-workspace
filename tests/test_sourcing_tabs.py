"""딜 소싱 [+ 탭 추가] — 사람이 없어도 서는 새 갈래(탭).

여기서 못 박는 것.

  ① 만든 탭은 사람이 없어도 선다(빈 탭 안내) · 시트에서 온 탭들 뒤 · 사람을 넣어도 그 자리
  ② 이름 — 앞뒤 공백을 떼고, 비거나 40자를 넘거나 띄어쓰기만 다른 이름이 이미 있으면
     막는다. 잇는 갈래 이름(짧은 이름 `M&A` 포함) · `전체` 도 막는다
  ③ 권한은 [이름 저장]과 같다 — 이 화면을 보는 팀원 누구나
  ④ 새 탭은 투자사 관리 현황과 잇지 않는다 · 이름을 바꾸면 빈 탭도 따라간다

이름은 전부 가상값이다(공개 저장소).
"""
from __future__ import annotations

from urllib.parse import unquote

import pytest

from .conftest import DEMO_PASSWORD

NEW = "바이오 딜 찾는 투자사"


@pytest.fixture()
def page(client, db, users):
    from app.models import SourcingContact

    db.add_all([
        SourcingContact(bucket="시리즈 A 이상", position=0, name="홍길동"),
        SourcingContact(bucket="M&A 찾는 투자사", position=1000, name="김서연"),
    ])
    db.commit()
    client.post("/login", data={"phone": "01000000002", "password": DEMO_PASSWORD})
    return client       # u2 — 관리자가 아닌 팀원


def _add(client, name, tab="all"):
    return client.post("/sourcing/buckets/add", data={"name": name, "tab": tab},
                       follow_redirects=False)


def _tabs(db):
    from app.routers.sourcing import buckets

    db.expire_all()
    return [(t["label"], t["count"]) for t in buckets(db)]


def test_a_teammate_makes_an_empty_tab_that_shows(page, db):
    r = _add(page, f"  {NEW} ")
    assert r.status_code == 303
    assert unquote(r.headers["location"]) == f"/sourcing?tab={NEW}"
    # 시트에서 온 탭들 뒤에 선다
    assert _tabs(db) == [("시리즈 A 이상", 1), ("M&A 찾는 투자사", 1), (NEW, 0)]

    body = page.get("/sourcing", params={"tab": NEW}).text
    assert f"{NEW} <span>0</span>" in body
    assert 'data-testid="sourcing-empty-tab"' in body
    assert f"{NEW}에 추가" in body                  # 바로 사람을 넣을 수 있다
    assert "잇지 않습니다" in body                   # 만드는 창의 안내


def test_adding_people_keeps_the_tab_in_its_place(page, db):
    from app.models import SourcingContact

    _add(page, NEW)
    r = page.post("/api/sourcing", params={"bucket": NEW}, json={"name": "박도윤"})
    assert r.status_code == 200, r.text
    row = db.get(SourcingContact, r.json()["id"])
    assert row.bucket == NEW
    # 새 줄의 position 은 1 이지만 탭은 맨 앞으로 뛰지 않는다
    assert _tabs(db)[-1] == (NEW, 1)
    # 잇지 않는 탭 — 투자사 관리 현황 쪽은 그대로다
    from app.models import VcContact
    assert row.vc_contact_id is None and db.query(VcContact).count() == 0


@pytest.mark.parametrize("bad, why", [
    ("", "입력하세요"),
    ("   ", "입력하세요"),
    ("가" * 41, "40자"),
    ("시리즈A이상", "이미 있는 탭"),                 # 띄어쓰기만 다르다
    ("M&A", "잇는 갈래"),                           # 잇는 갈래의 짧은 이름
    ("구주매각 / 후속투자사 찾는 투자사 / 공동투자사 / 앵커투자사", "잇는 갈래"),
    # 대표 · 개인참여도 잇는 갈래다(2026-10-22) — 갈래 글자(띄어쓰기만 달라도) · 짧은 이름
    ("딜 소싱 참여 투자사 대표", "잇는 갈래"),
    ("딜소싱  개인참여 심사역", "잇는 갈래"),
    ("대표", "잇는 갈래"),
    ("개인 참여", "잇는 갈래"),
    ("전체", "전체"),
    ("all", "전체"),
])
def test_bad_names_are_refused(page, db, bad, why):
    before = _tabs(db)
    r = _add(page, bad, tab="시리즈 A 이상")
    assert r.status_code == 303
    where = unquote(r.headers["location"])
    assert where.startswith("/sourcing?tab=시리즈 A 이상&msg=")
    assert why in where
    assert _tabs(db) == before


def test_the_same_new_name_twice_is_refused(page, db):
    _add(page, NEW)
    r = _add(page, NEW.replace(" ", ""))
    assert "이미 있는 탭" in unquote(r.headers["location"])
    assert [t for t in _tabs(db) if t[0] == NEW] == [(NEW, 0)]


def test_renaming_an_empty_tab_moves_it(page, db):
    _add(page, NEW)
    r = page.post("/sourcing/buckets/rename", data={"old": NEW, "new": "헬스케어 찾는 투자사"},
                  follow_redirects=False)
    assert "msg=" not in r.headers["location"]
    assert _tabs(db)[-1] == ("헬스케어 찾는 투자사", 0)
    assert NEW not in [t[0] for t in _tabs(db)]


@pytest.mark.parametrize("synced", ["M&A", "대표", "개인참여"])
def test_a_tab_cannot_be_renamed_into_a_synced_name(page, db, synced):
    r = page.post("/sourcing/buckets/rename", data={"old": "시리즈 A 이상", "new": synced},
                  follow_redirects=False)
    assert "잇는 갈래" in unquote(r.headers["location"])
    assert ("시리즈 A 이상", 1) in _tabs(db)


def test_the_add_tab_hint_lists_every_synced_branch(page, db):
    """[+ 탭 추가]의 안내가 잇는 갈래를 서버 목록(`BRANCHES`)으로 읽는다."""
    body = page.get("/sourcing").text
    assert "Series A · M&amp;A · 구주·후속 · 대표 · 개인참여 뿐입니다" in body


def test_the_empty_tab_has_its_own_sheet(page, db):
    import io

    import openpyxl

    _add(page, NEW)
    r = page.get("/api/sourcing/export.xlsx", params={"every": 1})
    book = openpyxl.load_workbook(io.BytesIO(r.content))
    assert book.sheetnames[-1] == NEW
    assert book[NEW].max_row == 1                  # 머리글만


# ── [탭 삭제] (#78) — 사람이 0명인 탭만 · 관리자만 ─────────────────────────────

@pytest.fixture()
def admin(page, db, users):
    users["u2"].role = "admin"
    db.commit()
    return page


def _delete(client, label):
    return client.post("/sourcing/buckets/delete", data={"label": label},
                       follow_redirects=False)


def test_an_admin_deletes_an_empty_tab_with_its_texts(admin, db, users):
    from app.models import MessageTemplate, TemplateChoice

    _add(admin, NEW)
    text = MessageTemplate(kind="sourcing_intro", name=NEW, body="안내")
    db.add(text)
    db.flush()
    db.add(TemplateChoice(user_id=users["u2"].id, kind="sourcing_intro", variant=NEW,
                          template_id=text.id))
    db.commit()
    assert 'data-testid="sourcing-tab-delete"' in admin.get(
        "/sourcing", params={"tab": NEW}).text

    r = _delete(admin, NEW)
    assert r.status_code == 303 and "지웠습니다" in unquote(r.headers["location"])
    assert NEW not in [t[0] for t in _tabs(db)]
    assert db.query(MessageTemplate).filter_by(name=NEW).count() == 0
    assert db.query(TemplateChoice).filter_by(variant=NEW).count() == 0
    assert _add(admin, NEW).status_code == 303 and (NEW, 0) in _tabs(db)  # 다시 만들 수 있다


def test_a_tab_with_people_is_not_deleted(admin, db):
    body = admin.get("/sourcing", params={"tab": "시리즈 A 이상"}).text
    assert 'data-testid="sourcing-tab-delete"' not in body
    r = _delete(admin, "시리즈 A 이상")
    assert "옮겨 주세요" in unquote(r.headers["location"])
    assert ("시리즈 A 이상", 1) in _tabs(db)


def test_only_an_admin_deletes(page, db):
    _add(page, NEW)
    assert 'data-testid="sourcing-tab-delete"' not in page.get(
        "/sourcing", params={"tab": NEW}).text
    assert _delete(page, NEW).status_code != 303
    assert (NEW, 0) in _tabs(db)
