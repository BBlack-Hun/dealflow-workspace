"""IR 기업 현황 메뉴의 셋째 탭 `Series A 이상 딜 소싱`.

사용자 요청: 시트 `Series A 이상 기업` 을 기반으로 탭을 더하고, 다른 탭과 같이
**칸 일부를 고정**하고 **필터**를 단다.

  ① 화면 — 세 탭이 한 줄로 서고, 표는 수신 일자 최신순이며, 왼쪽 두 칸
     (`NO` · `기업명`)이 고정이고, 고르는 칸에 머리글 필터가 선다.
  ② 고치기 · 더하기 · 지우기 — 옆 두 탭과 **같은 판정**(지우기는
     `can_delete_company`, 투자컨설턴트는 화면부터 막힌다).
  ③ 가져오기 — 시트 머리글로 칸을 찾고, 가져온 그때의 `(기업명, 수신 일자,
     공유 심사역, 회사)`(`import_key`)로 한 줄을 알아보며, 몇 번을 돌려도
     결과가 같다.
  ④ 브라우저 — 검색 · 필터 · 정렬 · 줄 번호 · 저장 뒤처리(node).

값은 전부 지어낸 것이다 — 저장소가 공개다.
"""
from __future__ import annotations

import re
import shutil
import sqlite3
import subprocess
from pathlib import Path

import pytest

from .conftest import DEMO_PASSWORD
from .test_migrations import _alembic

ROOT = Path(__file__).resolve().parent.parent

# 시트 머리글 **그대로**(원본에 `기업구분 ` 처럼 뒤에 공백이 붙은 것까지).
SHEET_HEAD = ["NO", "수신 일자", "Type", "담당팀원", "공유 심사역", "회사", "기업명",
              "사업분야 대분류", "소분류", "기업구분 ", "상세 내용", "딜소개 문구"]


def _sheet_rows():
    return [
        SHEET_HEAD,
        ["1", "2026-07-01", "딜 소싱 네트워크", "강민준", "박영희 이사", "가나벤처스",
         "샘플애그", "애그테크", "스마트팜", "Series A (누적투자금 10억이상)",
         "스마트팜 관제\n후속 라운드 준비", ""],
        ["2", "2026.8.25", "후속 투자 연계", "윤서아", "김철수 수석", "다라인베스트",
         "(주)샘플메디", "헬스케어", "의료AI", "Series B", "의료 영상 판독", ""],
        # 번호가 빈 줄 — 원본에 실제로 있다. 열쇠가 아니라 그대로 들어온다.
        ["", "2026-01-01", "딜 소싱 네트워크", "강민준", "이영수 상무", "가나벤처스",
         "샘플페이", "", "", "", "", ""],
        # 기업명이 빈 줄(합계 · 메모) — 건너뛴다.
        ["", "", "", "", "", "", "", "", "", "", "메모 줄", ""],
    ]


def _write_csv(path: Path, rows) -> Path:
    import csv

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)
    return path


@pytest.fixture()
def rows(db):
    from app.models import IrCompany, SeriesACompany

    db.add_all([
        # IR 기업 현황에도 있는 기업 — 이름 표기가 달라도(`(주)`) 이어져야 한다.
        IrCompany(name="샘플애그"),
        SeriesACompany(name="(주)샘플애그", received_at="2026-07-01",
                       deal_type="딜 소싱 네트워크", assignee_name="강민준",
                       shared_by="박영희 이사", source_firm="가나벤처스",
                       sector_major="애그테크", sector_minor="스마트팜",
                       series="Series A (누적투자금 10억이상, 년매출액 10억이상)",
                       detail="스마트팜 관제"),
        SeriesACompany(name="샘플메디", received_at="2026-08-25",
                       deal_type="후속 투자 연계", assignee_name="윤서아",
                       shared_by="김철수 수석", source_firm="다라인베스트",
                       series="Series B"),
        # 같은 날 둘 — 이름순으로 선다.
        SeriesACompany(name="가온샘플", received_at="2026-08-25"),
        # 날짜가 아닌 줄 — 맨 아래.
        SeriesACompany(name="나루샘플", received_at="날짜 미정"),
    ])
    db.commit()
    return db


def _body_rows(html: str) -> list:
    table = re.search(r'<table[^>]*id="sa-table".*?</table>', html, re.S).group(0)
    return re.findall(r"<tr data-id=.*?</tr>", table, re.S)


def _names(html: str) -> list:
    return [re.search(r'data-field="name">([^<]*)<', r).group(1) for r in _body_rows(html)]


# ── ① 화면 ──────────────────────────────────────────────────────────────────

def test_탭이_열리고_세_탭이_한_줄로_선다(logged_in, rows):
    res = logged_in.get("/companies?tab=series-a")
    assert res.status_code == 200
    html = res.text
    tabs = re.search(r'<nav class="sheet-tabs co-tabs">.*?</nav>', html, re.S).group(0)
    assert "IR 기업 현황" in tabs and "스타트업DB" in tabs and "Series A 이상 딜 소싱" in tabs
    active = re.findall(r'class="sheet-tab active" href="([^"]+)"', tabs)
    assert active == ["/companies?tab=series-a"], active
    # 셋째 탭의 수 = 이 표의 줄 수.
    assert re.search(r"Series A 이상 딜 소싱 <span[^>]*>4</span>", tabs), tabs


@pytest.mark.parametrize("path", ["/companies", "/companies?tab=db"])
def test_옆_두_탭에도_셋째_탭이_선다(logged_in, rows, path):
    """탭 줄은 한 조각(`_company_tabs.html`)이다 — 어느 탭에서든 같은 셋이 선다."""
    html = logged_in.get(path).text
    tabs = re.search(r'<nav class="sheet-tabs co-tabs">.*?</nav>', html, re.S).group(0)
    assert 'href="/companies?tab=series-a"' in tabs
    assert re.search(r"Series A 이상 딜 소싱 <span[^>]*>4</span>", tabs), tabs
    assert "IR 기업 현황 <span title=\"등록된 기업\">1</span>" in tabs


def test_수신_일자_최신순이고_날짜_없는_줄은_맨_아래(logged_in, rows):
    names = _names(logged_in.get("/companies?tab=series-a").text)
    assert names == ["가온샘플", "샘플메디", "(주)샘플애그", "나루샘플"], names


def test_수신_일자_머리글은_처음부터_최신순으로_서_있다(logged_in, rows):
    """처음 차례가 수신 일자 최신순이다 — 머리글도 그렇다고 말해야 첫 클릭이
    뜻밖의 오래된 순이 되지 않는다(`table_sort.js` 가 `data-sort-default` 를 읽는다)."""
    html = logged_in.get("/companies?tab=series-a").text
    head = re.search(r'<th[^>]*data-sort="received_at"[^>]*>', html).group(0)
    assert 'data-sort-default="desc"' in head, head


def test_왼쪽_두_칸이_가로로_고정이다(logged_in, rows):
    """IR 기업 현황과 **같은 규칙 · 같은 CSS**(`.stick-a` · `.stick-b`)."""
    html = logged_in.get("/companies?tab=series-a").text
    table = re.search(r'<table[^>]*id="sa-table"[^>]*>', html).group(0)
    assert "--stick-a:0px; --stick-b:36px" in table, table
    heads = re.findall(r"<th\b([^>]*)>(.*?)</th>", re.search(
        r'id="sa-table".*?</thead>', html, re.S).group(0), re.S)
    assert "stick-a" in heads[0][0] and heads[0][1].strip() == "NO"
    assert "stick-b" in heads[1][0] and heads[1][1].strip() == "기업명"
    assert 'width:36px' in heads[0][0], "`--stick-b` 와 `NO` 의 폭이 갈렸습니다"
    first = _body_rows(html)[0]
    assert re.search(r'<td class="rowno[^"]*stick stick-a"', first)
    assert re.search(r'<td class="co-name stick stick-b"', first)


FILTERS = {"type": "Type", "assignee": "담당자", "sharer": "공유 심사역",
           "firm": "회사", "sector": "대분류", "minor": "소분류", "series": "투자라운드"}


def test_고르는_칸에_머리글_필터가_선다(logged_in, rows):
    html = logged_in.get("/companies?tab=series-a").text
    for key, label in FILTERS.items():
        assert f'data-filters="{key}:{label}"' in html, key
        assert f"data-f-{key}=" in _body_rows(html)[0], key
    # 같은 필터 부품을 같은 방식으로 연다(`DealflowFilters.init`).
    src = (ROOT / "app" / "static" / "js" / "series_a.js").read_text("utf-8")
    assert re.search(r'DealflowFilters\.init\(\{[\s\S]{0,200}table: "#sa-table"', src)


def test_투자라운드는_괄호_앞만_보이고_고칠_때는_전체다(logged_in, rows):
    html = logged_in.get("/companies?tab=series-a").text
    row = next(r for r in _body_rows(html) if "샘플애그" in r)
    assert 'data-f-series="Series A"' in row
    assert 'data-value="Series A (누적투자금 10억이상, 년매출액 10억이상)">Series A<' in row


def test_IR_기업_현황에도_있는_기업에는_딱지가_선다(logged_in, rows):
    """이름 표기가 달라도(`(주)샘플애그` ↔ `샘플애그`) `deal_history._key` 로 맞는다."""
    html = logged_in.get("/companies?tab=series-a").text
    by_name = {re.search(r'data-field="name">([^<]*)<', r).group(1): r
               for r in _body_rows(html)}
    agri = re.search(r"<a class=\"room-badge ok ir-link\"[^>]*>", by_name["(주)샘플애그"]).group(0)
    assert "hidden" not in agri and 'href="/companies?q=%EC%83%98%ED%94%8C%EC%95%A0%EA%B7%B8"' in agri
    medi = re.search(r"<a class=\"room-badge ok ir-link\"[^>]*>", by_name["샘플메디"]).group(0)
    assert "hidden" in medi


def test_검색어는_보이는_칸으로_걸린다(logged_in, rows):
    row = next(r for r in _body_rows(logged_in.get("/companies?tab=series-a").text)
               if "샘플메디" in r)
    hay = re.search(r'data-search="([^"]*)"', row).group(1)
    for word in ("샘플메디", "다라인베스트", "김철수", "후속 투자 연계", "series b"):
        assert word.lower() in hay, word


def test_칸을_눌러_고친다(logged_in, rows):
    html = logged_in.get("/companies?tab=series-a").text
    assert 'data-inline-url="/api/series-a-companies"' in html
    for field in ("name", "received_at", "deal_type", "assignee_name", "shared_by",
                  "source_firm", "sector_major", "sector_minor", "series", "detail",
                  "deal_pitch"):
        assert f'data-field="{field}"' in html, field


# ── ② 고치기 · 더하기 · 지우기 ───────────────────────────────────────────────

def _row_id(db, name):
    from app.models import SeriesACompany

    return next(r.id for r in db.query(SeriesACompany).all() if r.name == name)


def test_고치면_서버만_아는_값을_돌려준다(logged_in, rows):
    rid = _row_id(rows, "샘플메디")
    res = logged_in.patch(f"/api/series-a-companies/{rid}",
                          json={"series": "Series C (누적투자금 50억 이상)",
                                "received_at": "2026.9.3", "name": "  "})
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["series_short"] == "Series C"
    # 가져오기와 같은 모양으로 맞춘다 — 다르면 다음 가져오기가 다른 줄로 읽는다.
    assert data["received_at"] == "2026-09-03" and data["received_sort"] == "2026-09-03"
    # 기업명은 비울 수 없다.
    assert data["name"] == "샘플메디"


def test_고친_칸은_서버가_저장한_값과_검색_글자를_돌려준다(logged_in, rows):
    """화면은 응답으로 칸을 다시 그리고(`series_a.js`) 검색 글자도 바꾼다.

    응답에 고친 칸이 없으면 서버가 받지 않은 값(빈 기업명)이 화면에만 남고,
    검색 글자가 없으면 고친 값이 새로고침 전까지 검색에 안 걸린다.
    """
    rid = _row_id(rows, "샘플메디")
    data = logged_in.patch(f"/api/series-a-companies/{rid}",
                           json={"source_firm": "  마바캐피탈 ", "name": ""}).json()
    assert data["name"] == "샘플메디", "빈 이름을 받지 않았으면 옛 이름을 돌려줘야 한다"
    assert data["source_firm"] == "마바캐피탈"
    for field in ("deal_type", "assignee_name", "shared_by", "sector_major",
                  "sector_minor", "detail", "deal_pitch"):
        assert field in data, field
    assert "마바캐피탈" in data["search"] and "다라인베스트" not in data["search"]


def test_이름을_고치면_IR_딱지도_다시_선다(logged_in, rows):
    rid = _row_id(rows, "샘플메디")
    res = logged_in.patch(f"/api/series-a-companies/{rid}", json={"name": "샘플 애그"})
    assert res.json()["ir_name"] == "샘플애그"


def test_더하면_오늘_받은_딜로_선다(logged_in, rows):
    from app.clock import today

    assert logged_in.post("/api/series-a-companies", json={"name": " "}).status_code == 400
    res = logged_in.post("/api/series-a-companies", json={"name": "샘플새기업"})
    assert res.status_code == 200, res.text
    assert res.json()["received_at"] == today().isoformat()
    names = _names(logged_in.get("/companies?tab=series-a").text)
    assert names[0] == "샘플새기업", "오늘 받은 딜이 맨 위가 아닙니다"


@pytest.fixture()
def portal(db, users):
    """관리자 · 팀원 · 투자컨설턴트를 **따로** 로그인한 클라이언트."""
    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.models import User
    from app.services import auth as auth_svc

    pw = auth_svc.hash_password(DEMO_PASSWORD)
    db.add_all([
        User(id=91, name="관리자시험", phone="01000000091", role="admin", password_hash=pw),
        User(id=92, name="컨설턴트시험", phone="01000000092", role="consultant",
             password_hash=pw),
    ])
    db.commit()
    app = create_app()

    def sign_in(phone):
        c = TestClient(app)
        c.post("/login", data={"phone": phone, "password": DEMO_PASSWORD})
        return c

    return {"admin": sign_in("01000000091"), "member": sign_in("01000000001"),
            "consultant": sign_in("01000000092")}


def test_지우기는_옆_탭과_같은_판정이다(portal, rows):
    """관리자 · 팀원은 지우고, 투자컨설턴트는 화면부터 막힌다."""
    from app.routers.companies import can_delete_company

    a, b = _row_id(rows, "가온샘플"), _row_id(rows, "나루샘플")
    html = portal["member"].get("/companies?tab=series-a").text
    assert "js-sa-delete" in html, "팀원에게 [삭제] 가 안 보입니다"
    assert portal["member"].delete(f"/api/series-a-companies/{a}").status_code == 200
    assert portal["admin"].delete(f"/api/series-a-companies/{b}").status_code == 200
    assert portal["member"].delete(f"/api/series-a-companies/{a}").status_code == 404

    left = _row_id(rows, "샘플메디")
    consultant = portal["consultant"]
    assert consultant.get("/companies?tab=series-a").status_code in (303, 403)
    assert consultant.patch(f"/api/series-a-companies/{left}",
                            json={"detail": "x"}).status_code == 403
    assert consultant.delete(f"/api/series-a-companies/{left}").status_code == 403
    # 판정은 한 함수다 — 역할을 여기서 새로 적지 않는다.
    src = (ROOT / "app" / "routers" / "series_a.py").read_text("utf-8")
    assert "require_delete_company(user)" in src and "can_delete_company(user)" in src
    assert can_delete_company.__module__ == "app.routers.companies"


def test_고친_것은_수정_로그에_남는다(logged_in, rows):
    from app.models import EditLog

    rid = _row_id(rows, "샘플메디")
    logged_in.patch(f"/api/series-a-companies/{rid}", json={"source_firm": "마바캐피탈"})
    rows.expire_all()
    logs = [r for r in rows.query(EditLog).all() if r.table_name == "series_a_companies"]
    assert logs, "고친 것이 수정 로그에 안 남았습니다"


# ── ③ 가져오기 ──────────────────────────────────────────────────────────────

def test_머리글로_칸을_찾고_옛_이름도_읽는다():
    from app.services import series_a_import as svc

    parsed = svc.parse(_sheet_rows())
    assert len(parsed.records) == 3 and parsed.no_name == 1
    first = parsed.records[0]
    assert first["series"] == "Series A (누적투자금 10억이상)", "`기업구분 `(뒤 공백)을 못 읽었습니다"
    assert first["sector_major"] == "애그테크"
    assert first["shared_by"] == "박영희 이사" and first["source_firm"] == "가나벤처스"
    assert first["detail"] == "스마트팜 관제\n후속 라운드 준비", "줄바꿈이 사라졌습니다"
    assert parsed.records[1]["received_at"] == "2026-08-25", "날짜 모양을 맞추지 않았습니다"
    assert parsed.missing == []

    # 화면 이름으로 머리글을 고친 시트도 읽는다.
    renamed = [[h.replace("사업분야 대분류", "대분류").replace("기업구분 ", "투자라운드")
                for h in _sheet_rows()[0]]] + _sheet_rows()[1:]
    again = svc.parse(renamed)
    assert again.records[0]["series"].startswith("Series A")
    assert again.records[0]["sector_major"] == "애그테크"


def test_머리글이_없으면_멈춘다():
    from app.services import series_a_import as svc

    with pytest.raises(ValueError):
        svc.parse([["기업명", "대표자"], ["샘플애그", "김가나"]])


def test_몇_번을_돌려도_같고_앱에서_고친_값은_안_덮는다(db):
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    records = svc.parse(_sheet_rows()).records
    first = svc.upsert(db, records)
    db.commit()
    assert (first.created, first.updated) == (3, 0)

    again = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    assert (again.created, again.updated, again.unchanged) == (0, 0, 3)
    assert db.query(SeriesACompany).count() == 3

    # 앱에서 고친 값 · 앱에서 채운 칸.
    medi = next(r for r in db.query(SeriesACompany).all() if "샘플메디" in r.name)
    medi.source_firm = "앱에서 고친 투자사"
    pay = next(r for r in db.query(SeriesACompany).all() if r.name == "샘플페이")
    pay.detail = "앱에서 적은 내용"
    db.commit()

    # 시트에는 샘플페이의 분야가 새로 생겼다(빈 칸 채우기).
    sheet = _sheet_rows()
    sheet[3][7] = "핀테크"
    kept = svc.upsert(db, svc.parse(sheet).records)
    db.commit()
    assert kept.filled == 1 and kept.kept == 1 and kept.overwritten == 0
    db.refresh(medi), db.refresh(pay)
    assert medi.source_firm == "앱에서 고친 투자사", "앱에서 고친 값을 덮었습니다"
    assert pay.detail == "앱에서 적은 내용", "시트의 빈 칸이 앱 값을 지웠습니다"
    assert pay.sector_major == "핀테크"

    forced = svc.upsert(db, svc.parse(sheet).records, overwrite=True)
    db.commit()
    db.refresh(medi)
    assert forced.overwritten == 1 and medi.source_firm == "다라인베스트"


def test_이름_표기가_달라도_같은_딜이다(db):
    """`(주)샘플메디` 와 `샘플메디` 는 같은 날 같은 딜 — 두 줄이 되면 안 된다."""
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    db.add(SeriesACompany(name="샘플메디", received_at="2026-08-25"))
    db.commit()
    result = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    assert result.created == 2 and result.updated == 1
    assert db.query(SeriesACompany).count() == 3


def test_같은_기업이_다른_날_또_오면_다른_줄이다(db):
    from app.services import series_a_import as svc

    sheet = _sheet_rows()
    sheet.append(["", "2026-09-30", "후속 투자 연계", "강민준", "박영희 이사",
                  "가나벤처스", "샘플애그", "", "", "", "", ""])
    sheet.append(list(sheet[1]))                   # 시트 안에서 겹친 줄
    result = svc.upsert(db, svc.parse(sheet).records)
    assert result.created == 4 and result.duplicates == 1


def _deal(day, sharer, firm, name="샘플애그"):
    """시트 한 줄 — `SHEET_HEAD` 차례."""
    return ["", day, "딜 소싱 네트워크", "강민준", sharer, firm, name,
            "", "", "", "", ""]


def test_같은_날_같은_기업이라도_넘겨_준_곳이_다르면_다른_딜이다(db):
    """열쇠가 `(기업명, 수신 일자)` 뿐이면 두 번째 딜이 겹친 줄로 **버려졌다.**

    같은 날 두 투자사에게서 같은 기업을 받는 일이 있다. 한 줄은 받은 딜 한
    건이라, 넘겨 준 사람 · 회사까지 같아야 같은 줄이다(`import_key`).
    """
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    sheet = [SHEET_HEAD,
             _deal("2026-07-01", "박영희 이사", "가나벤처스"),
             _deal("2026-07-01", "김철수 수석", "다라인베스트")]
    result = svc.upsert(db, svc.parse(sheet).records)
    db.commit()
    assert (result.created, result.duplicates) == (2, 0), result
    assert db.query(SeriesACompany).count() == 2

    again = svc.upsert(db, svc.parse(sheet).records)
    db.commit()
    assert (again.created, again.unchanged) == (0, 2), again
    assert db.query(SeriesACompany).count() == 2


def test_열쇠는_띄어쓰기_법인_표기가_달라도_같다():
    from app.services import series_a_import as svc

    a = svc.import_key({"name": "(주)샘플애그", "received_at": "2026-07-01",
                        "shared_by": "박영희 이사", "source_firm": "가나벤처스"})
    b = svc.import_key({"name": "샘플 애그", "received_at": "2026-07-01",
                        "shared_by": "박영희이사", "source_firm": "(주)가나벤처스 "})
    assert a == b
    c = svc.import_key({"name": "샘플애그", "received_at": "2026-07-01",
                        "shared_by": "김철수 수석", "source_firm": "가나벤처스"})
    assert a != c


def test_앱에서_기업명과_수신_일자를_고쳐도_다시_가져오면_같은_줄이다(logged_in, db):
    """열쇠는 **가져온 그때의 시트 값**이다 — 앱에서 고쳐도 다시 셈하지 않는다.

    고칠 때마다 열쇠가 바뀌면 다음 가져오기에서 같은 딜이 **새 줄로 또** 선다.
    """
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    medi = next(r for r in db.query(SeriesACompany).all() if "샘플메디" in r.name)
    key = medi.import_key
    assert key, "가져온 줄에 열쇠가 안 적혔다"

    res = logged_in.patch(f"/api/series-a-companies/{medi.id}",
                          json={"name": "샘플메디플러스", "received_at": "2026-08-26"})
    assert res.status_code == 200, res.text
    db.expire_all()
    assert db.get(SeriesACompany, medi.id).import_key == key, "고쳤더니 열쇠가 바뀌었다"

    again = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    assert again.created == 0, "앱에서 고친 딜이 다시 가져오기에서 새 줄로 섰다"
    assert db.query(SeriesACompany).count() == 3
    medi = db.get(SeriesACompany, medi.id)
    assert (medi.name, medi.received_at) == ("샘플메디플러스", "2026-08-26"), \
        "앱에서 고친 값을 덮었다"
    assert again.kept == 2


def test_열쇠가_없던_옛_줄은_한_번_맞춰_열쇠를_적는다(db):
    """`import_key` 칸이 생기기 전에 들어온 줄 · 앱에서 더한 줄은 열쇠가 없다.

    그런 줄은 예전 열쇠 `(기업명, 수신 일자)` 로 맞추고, 맞으면 열쇠를 적는다 —
    그 뒤로는 앱에서 이름을 고쳐도 열쇠로 맞는다.
    """
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    old = SeriesACompany(name="샘플애그", received_at="2026-07-01")
    db.add(old)
    db.commit()
    result = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    assert (result.created, result.linked) == (2, 1), result
    db.refresh(old)
    assert old.import_key and old.shared_by == "박영희 이사", "빈 칸을 안 채웠다"

    old.name = "샘플애그테크"
    db.commit()
    again = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.commit()
    assert (again.created, again.linked) == (0, 0), again
    assert db.query(SeriesACompany).count() == 3


def test_옛_줄은_넘겨_준_곳까지_같은_시트_줄이_먼저_가져간다(db):
    """옛 줄 하나에 같은 `(기업명, 수신 일자)` 의 시트 줄이 둘이면, **넘겨 준
    곳까지 같은 줄**이 그 옛 줄이다 — 시트 차례로 먼저 온 줄이 아니다."""
    from app.models import SeriesACompany
    from app.services import series_a_import as svc

    old = SeriesACompany(name="샘플애그", received_at="2026-07-01",
                         shared_by="김철수 수석", source_firm="다라인베스트")
    db.add(old)
    db.commit()
    sheet = [SHEET_HEAD,
             _deal("2026-07-01", "박영희 이사", "가나벤처스"),
             _deal("2026-07-01", "김철수 수석", "다라인베스트")]
    result = svc.upsert(db, svc.parse(sheet).records)
    db.commit()
    assert (result.created, result.linked) == (1, 1), result
    db.refresh(old)
    assert old.import_key == svc.import_key(svc.parse(sheet).records[1])
    rows = db.query(SeriesACompany).all()
    assert sorted(r.source_firm for r in rows) == ["가나벤처스", "다라인베스트"]


def test_옛_줄이_어느_것인지_모르면_건너뛰고_알린다(db, tmp_path, capsys):
    """같은 `(기업명, 수신 일자)` 의 옛 줄이 둘인데 시트 줄과 딱 맞는 것이 없으면
    고르지 않는다 — 아무거나 고르면 남의 딜에 값을 채우고, 새로 만들면 겹친다."""
    from app.models import SeriesACompany
    from app.services import series_a_import as svc
    from scripts import import_series_a as script

    db.add_all([SeriesACompany(name="샘플페이", received_at="2026-01-01"),
                SeriesACompany(name="샘플페이", received_at="2026-01-01",
                               detail="앱에서 더한 줄")])
    db.commit()
    result = svc.upsert(db, svc.parse(_sheet_rows()).records)
    db.rollback()
    assert (result.created, result.ambiguous) == (2, 1), result

    path = _write_csv(tmp_path / "sheet.csv", _sheet_rows())
    assert script.main([str(path), "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "어느 줄인지 몰라 건너뛴 시트 줄 1개" in out, out


def test_스크립트는_파일을_받고_dry_run_은_저장하지_않는다(db, tmp_path, capsys):
    from app.models import SeriesACompany
    from scripts import import_series_a as script

    path = _write_csv(tmp_path / "sheet.csv", _sheet_rows())
    assert script.main([str(path), "--dry-run"]) == 0
    db.expire_all()
    assert db.query(SeriesACompany).count() == 0, "--dry-run 인데 저장했습니다"
    out = capsys.readouterr().out
    assert "새 줄 3" in out and "미리보기" in out
    # 출력에 기업명 · 사람 이름을 늘어놓지 않는다.
    for word in ("샘플애그", "박영희", "가나벤처스"):
        assert word not in out

    assert script.main([str(path)]) == 0
    db.expire_all()
    assert db.query(SeriesACompany).count() == 3
    assert script.main([str(path)]) == 0
    db.expire_all()
    assert db.query(SeriesACompany).count() == 3, "다시 돌렸더니 줄이 늘었습니다"


@pytest.mark.parametrize("name, content", [
    ("broken.xlsx", b"this is not a zip"),        # 깨진 xlsx(zip 이 아니다)
    ("old.xls", b"\xd0\xcf\x11\xe0 old excel"),     # openpyxl 이 못 읽는 옛 형식
])
def test_스크립트는_못_읽는_파일에_traceback_대신_한_줄로_알린다(tmp_path, capsys, name, content):
    from scripts import import_series_a as script

    path = tmp_path / name
    path.write_bytes(content)
    assert script.main([str(path)]) == 1
    out = capsys.readouterr().out
    assert "읽지 못했습니다" in out, out


def test_xlsx_는_이_시트로_보이는_탭을_고른다(tmp_path):
    import openpyxl

    from app.services import series_a_import as svc

    book = openpyxl.Workbook()
    other = book.active
    other.title = "IR 기업 현황"
    other.append(["NO", "기업명", "사업분야 대분류"])
    other.append(["1", "샘플다른", "바이오"])
    ws = book.create_sheet("Series A 이상 기업")
    for row in _sheet_rows():
        ws.append(row)
    from datetime import datetime
    ws.cell(2, 2).value = datetime(2026, 7, 1)          # 날짜 칸으로 온 값
    path = tmp_path / "book.xlsx"
    book.save(path)

    parsed = svc.parse(svc.read_rows(path))
    assert parsed.records[0]["name"] == "샘플애그"
    assert parsed.records[0]["received_at"] == "2026-07-01"

    with pytest.raises(ValueError):
        svc.read_rows(path, sheet="없는 탭")


def test_저장소에_시트_주소가_없다():
    """공개 저장소다 — 시트 주소 · 아이디는 어디에도 적지 않는다(파일 경로를 받는다)."""
    for path in (ROOT / "scripts" / "import_series_a.py",
                 ROOT / "app" / "services" / "series_a_import.py"):
        text = path.read_text("utf-8")
        assert "docs.google.com" not in text and "spreadsheets/d/" not in text


# ── 판(0102) ────────────────────────────────────────────────────────────────

def _columns(db) -> set:
    con = sqlite3.connect(db)
    try:
        return {r[1] for r in con.execute('PRAGMA table_info("series_a_companies")')}
    finally:
        con.close()


def test_가져오기_열쇠_판은_0101_뒤에_칸을_더하고_되돌리면_뺀다(tmp_path):
    """0102 는 0101(확정한 공유 심사역) 바로 뒤 — 올리면 칸이 서고, 0101 로
    내리면 이 칸만 빠진다(0101 의 칸은 그대로)."""
    db = tmp_path / "m.db"
    done = _alembic(db, "upgrade", "0101_series_a_shared_by")
    assert done.returncode == 0, done.stdout + done.stderr
    con = sqlite3.connect(db)
    try:
        # 운영 DB 는 이 칸이 없는 표에서 출발한다(0001 의 create_all 은 이미 칸이 있다)
        cols = [r[1] for r in con.execute('PRAGMA table_info("series_a_companies")')
                if r[1] != "import_key"]
        keep = ", ".join(cols)
        con.execute(f"CREATE TABLE sa_old AS SELECT {keep} FROM series_a_companies")
        con.execute("DROP TABLE series_a_companies")
        con.execute("ALTER TABLE sa_old RENAME TO series_a_companies")
        con.commit()
    finally:
        con.close()
    assert "import_key" not in _columns(db)

    done = _alembic(db, "upgrade", "0102_series_a_import_key")
    assert done.returncode == 0, done.stdout + done.stderr
    assert {"import_key", "shared_by_contact_id"} <= _columns(db)

    done = _alembic(db, "downgrade", "0101_series_a_shared_by")
    assert done.returncode == 0, done.stdout + done.stderr
    cols = _columns(db)
    assert "import_key" not in cols and "shared_by_contact_id" in cols


# ── ④ 브라우저 ──────────────────────────────────────────────────────────────

@pytest.mark.skipif(shutil.which("node") is None,
                    reason="node 미설치 — 브라우저 로직 테스트 생략")
def test_브라우저_쪽은_파일을_그대로_돌려_본다():
    script = ROOT / "tests" / "js" / "series_a_test.js"
    out = subprocess.run([shutil.which("node"), str(script)],
                         capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stdout + out.stderr
