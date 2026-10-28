"""투자사 관리 현황의 `딜 소싱` 칸 — 갈래(Series A · M&A · 구주·후속 · 대표 · 개인참여)를 고른다.

예전 `Series A 딜 소싱`(O/X) 칸을 갈래를 고르는 칸 하나로 바꿨다(사용자 결정
2026-10-21 · `services/sourcing_sync.BRANCHES`). 2026-10-22 대표 · 개인참여 갈래도 잇는다
(⑧ — 발송은 그대로). 여기서 못 박는 것.

  ① 체크 — 그 갈래에서 같은 분을 찾아 잇거나 새 줄(`딜 소싱 참여 제안 대기중`).
     풀기 — 그 줄을 `불참여` 로(지우지 않는다). 한 분이 두 갈래에 다 있을 수 있다.
  ② 이어진 줄들의 같은 칸(이름 · 번호 · 투자분야 …)은 어느 쪽에서 고쳐도 모두 같이.
  ③ `(딜 소개 X)` · `딜 소개 보류` 인 분은 넣지 않는다(빼는 것은 된다).
  ④ 엑셀 — 내보내기는 `Series A, M&A 불참여`, 올리기는 새 칸과 예전
     `Series A 딜 소싱`(O/X) 머리글을 다 받는다.
  ⑤ M&A 갈래의 [저절로 잇기] · 후보 [이 분으로 확정] — `참여` 칸은 그대로.
  ⑥ M&A 월간 리마인드는 M&A 문구(저장 문구가 있으면 그것).
  ⑦ 이주 0114 — 색인을 (vc_contact_id, bucket) 으로. 겹친 줄이 있으면 내리기를 거절.
  ⑧ 대표 · 개인참여 — 체크 · 풀기 · 같은 칸 · [저절로 잇기] · 엑셀이 나머지 셋과 같다.
     방(투자사 방 이어받기) · 문구 · Series A 딜 공유는 예전 그대로.

이름 · 회사 · 번호는 전부 가상값이다(공개 저장소).
"""
from __future__ import annotations

import io
import sqlite3

import pytest

from tests.test_migrations import _alembic
from tests.test_sourcing_vc_sync import (  # noqa: F401 - 픽스처를 함께 쓴다
    MNA, SHEET, SYNC, board, sc, vc)

FOLLOW = "구주매각 / 후속투자사 찾는 투자사 / 공동투자사 / 앵커투자사"
CEO = "딜 소싱  참여 투자사 대표"          # 운영 시트 탭처럼 두 칸 공백
SOLO = "딜 소싱  개인참여 심사역"
#: 잇지 않는 갈래 — [+ 탭 추가]로 만든 탭 같은 것.
OTHER = "해외 LP 찾는 투자사"
PENDING = "딜 소싱 참여 제안 대기중"


def pick(client, contact_id, value):
    return client.patch(f"/api/contacts/{contact_id}", json={"sourcing_branches": value})


def rows_of(db, bucket):
    from app.models import SourcingContact

    db.expire_all()
    return db.query(SourcingContact).filter_by(bucket=bucket).order_by(SourcingContact.id).all()


def row_of(client, contact_id):
    rows = client.get("/api/contacts").json()["contacts"]
    return next(r for r in rows if r["id"] == contact_id)


# ── ① 체크 · 풀기 ───────────────────────────────────────────────────────────

def test_checking_m_and_a_makes_a_pending_row_and_a_chip(board, db, users):
    contact = vc(db, users)

    r = pick(board, contact.id, "M&A")
    assert r.status_code == 200, r.text
    got = r.json()
    assert "M&A 찾는 투자사 에 '딜 소싱 참여 제안 대기중' 으로 새로" in got["sourcing_msg"]
    assert got["sourcing_branches"] == "M&A"
    assert got["sourcing_filter"] == "M&A"
    assert [(c["text"], c["kind"], c["on"]) for c in got["sourcing_chips"]] == [
        ("M&A·제안 대기중", "pending", True)]
    assert got["sourcing_join"] == ""             # 시리즈 A 갈래는 그대로 빈칸

    [row] = rows_of(db, MNA)
    assert (row.vc_contact_id, row.participation) == (contact.id, PENDING)
    assert rows_of(db, SYNC) == []


def test_checking_links_the_same_person_already_in_that_branch(board, db, users):
    contact = vc(db, users)
    existing = sc(db, bucket=MNA, participation="참여")   # 이름 · 번호가 둘 다 같다

    r = pick(board, contact.id, "M&A")
    assert r.status_code == 200, r.text
    assert "있던 '홍길동' 줄과 이었습니다" in r.json()["sourcing_msg"]
    [row] = rows_of(db, MNA)
    assert row.id == existing.id and row.vc_contact_id == contact.id
    assert row.participation == "참여"               # 잇기만 — 참여는 그대로


def test_one_investor_in_two_branches_and_unchecking_keeps_the_row(board, db, users):
    contact = vc(db, users)
    assert pick(board, contact.id, "Series A, M&A").status_code == 200
    [sa], [mna] = rows_of(db, SYNC), rows_of(db, MNA)
    assert sa.vc_contact_id == mna.vc_contact_id == contact.id
    assert row_of(board, contact.id)["sourcing_branches"] == "Series A, M&A"

    # 고르는 창은 고른 갈래 전부를 보낸다 — M&A 를 풀었다
    r = pick(board, contact.id, "Series A")
    assert r.status_code == 200, r.text
    [sa], [mna] = rows_of(db, SYNC), rows_of(db, MNA)
    assert (sa.participation, mna.participation) == (PENDING, "불참여")
    assert mna.vc_contact_id == contact.id           # 줄도 잇기도 남는다
    got = row_of(board, contact.id)
    assert got["sourcing_branches"] == "Series A"
    assert got["sourcing_filter"] == "Series A|M&A 불참여"
    assert [(c["text"], c["on"]) for c in got["sourcing_chips"]] == [
        ("Series A·제안 대기중", True), ("M&A·불참여", False)]

    # 다시 체크하면 다시 묻는다(제안 대기중) · 안 바뀐 갈래는 건드리지 않는다
    sa.participation = "참여"
    db.commit()
    assert pick(board, contact.id, "Series A, M&A").status_code == 200
    [sa], [mna] = rows_of(db, SYNC), rows_of(db, MNA)
    assert (sa.participation, mna.participation) == ("참여", PENDING)

    # 표에는 딱지 둘 — 빠진 갈래는 흐리게(`off`)
    pick(board, contact.id, "M&A")
    body = board.get(f"/contacts?sheet={SHEET}").text
    assert 'class="branch-chip left off"' in body and "Series A·불참여</span>" in body
    assert "M&amp;A·제안 대기중</span>" in body
    assert 'data-f-sourcing="Series A 불참여|M&amp;A"' in body


def test_unknown_branch_names_are_refused(board, db, users):
    contact = vc(db, users)
    r = pick(board, contact.id, "M&A, 시리즈 B")
    assert r.status_code == 400
    assert "Series A · M&A · 구주·후속 · 대표 · 개인참여 만 적습니다" in r.json()["detail"]
    assert rows_of(db, MNA) == []                    # 한 갈래라도 막히면 통째로


def test_the_follow_on_branch_and_the_edit_panel(board, db, users):
    contact = vc(db, users)
    assert pick(board, contact.id, "구주·후속").status_code == 200
    [row] = rows_of(db, FOLLOW)
    assert row.vc_contact_id == contact.id
    detail = board.get(f"/api/contacts/{contact.id}").json()["contact"]
    assert detail["sourcing_branches"] == "구주·후속"


# ── ② 같은 칸은 이어진 줄 모두 ──────────────────────────────────────────────

def test_shared_fields_follow_to_every_linked_row(board, db, users):
    from app.models import SourcingContact, VcContact

    contact = vc(db, users)
    pick(board, contact.id, "Series A, M&A")

    # 투자사 쪽에서 고치면 두 줄 다
    r = board.patch(f"/api/contacts/{contact.id}",
                    json={"name": "홍길순", "sectors": "AI,핀테크"})
    assert r.status_code == 200, r.text
    [sa], [mna] = rows_of(db, SYNC), rows_of(db, MNA)
    assert (sa.name, mna.name) == ("홍길순", "홍길순")
    assert (sa.sectors, mna.sectors) == ("AI,핀테크", "AI,핀테크")

    # M&A 줄에서 고치면 투자사 줄과 **시리즈 A 줄**도
    r = board.patch(f"/api/sourcing/{mna.id}", json={"phone": "010-5555-6666"})
    assert r.status_code == 200, r.text
    db.expire_all()
    assert db.get(VcContact, contact.id).phone == "010-5555-6666"
    assert db.get(SourcingContact, sa.id).phone == "010-5555-6666"


# ── ③ 딜 소개 X · 딜 소개 보류 ──────────────────────────────────────────────

@pytest.mark.parametrize("kw", [{"name": "홍길동 (딜 소개 X)"}, {"group_name": "딜 소개 보류"}])
def test_no_intro_investors_are_never_added(board, db, users, kw):
    contact = vc(db, users, **kw)
    r = pick(board, contact.id, "M&A")
    assert r.status_code == 400
    assert "딜 소싱에 넣지 않습니다" in r.json()["detail"]
    assert rows_of(db, MNA) == []

    # 이미 들어 있던 분을 빼는 것은 된다
    sc(db, name=contact.name, bucket=MNA, vc_contact_id=contact.id)
    assert pick(board, contact.id, "").status_code == 200
    assert rows_of(db, MNA)[0].participation == "불참여"
    # 다시 넣는 것은 막는다
    assert pick(board, contact.id, "M&A").status_code == 400


# ── ④ 엑셀 ──────────────────────────────────────────────────────────────────

def test_excel_export_lists_branches(board, db, users):
    import openpyxl

    contact = vc(db, users)
    pick(board, contact.id, "Series A, M&A")
    pick(board, contact.id, "Series A")
    r = board.get("/api/export/contacts.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    head = next(row for row in rows if row and "이름" in row)
    line = next(row for row in rows if row and row[head.index("이름")] == "홍길동")
    assert line[head.index("딜 소싱")] == "Series A, M&A 불참여"


def _upload(client, rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET                      # 그 명단으로 올린다(`board` 의 투자사 명단)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return client.post(
        "/api/import/contacts",
        files={"file": ("현황.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"dry_run": "false", "sheet": SHEET})


def test_excel_import_takes_the_new_column_and_the_old_header(board, db, users):
    a = vc(db, users, name="가투자", phone="010-1111-0001")
    b = vc(db, users, name="나투자", phone="010-1111-0002")
    c = vc(db, users, name="다투자", phone="010-1111-0003")
    sc(db, name="나투자", phone="010-1111-0002", vc_contact_id=b.id, participation="참여")

    # 새 칸 — 적힌 갈래만 넣거나 빼고, 안 적힌 갈래는 그대로
    r = _upload(board, [["이름", "투자사명", "딜 소싱", "딜소싱 참여 투자사"],
                        ["가투자", "가나다벤처스", "Series A, M&A", "아무 글"],
                        ["나투자", "가나다벤처스", "M&A", ""],
                        ["다투자", "가나다벤처스", "", ""]])
    assert r.status_code == 200, r.text
    assert {x.vc_contact_id for x in rows_of(db, SYNC)} == {a.id, b.id}
    assert {x.vc_contact_id for x in rows_of(db, MNA)} == {a.id, b.id}
    assert next(x for x in rows_of(db, SYNC) if x.vc_contact_id == b.id).participation == "참여"
    assert c.id not in {x.vc_contact_id for x in rows_of(db, SYNC) + rows_of(db, MNA)}

    # 예전 머리글 `Series A 딜 소싱` — O 는 Series A, X 는 Series A 불참여
    r = _upload(board, [["이름", "투자사명", "Series A 딜 소싱"],
                        ["나투자", "가나다벤처스", "X"],
                        ["다투자", "가나다벤처스", "O"]])
    assert r.status_code == 200, r.text
    got = {x.vc_contact_id: x.participation for x in rows_of(db, SYNC)}
    assert got == {a.id: PENDING, b.id: "불참여", c.id: PENDING}


def test_excel_import_reports_a_blocked_row_and_goes_on(board, db, users):
    vc(db, users, name="가투자 (딜 소개 X)", phone="010-1111-0001")
    ok = vc(db, users, name="나투자", phone="010-1111-0002")
    r = _upload(board, [["이름", "투자사명", "딜 소싱"],
                        ["가투자 (딜 소개 X)", "가나다벤처스", "M&A"],
                        ["나투자", "가나다벤처스", "M&A, 그런갈래"],
                        ["나투자", "가나다벤처스", "구주·후속"]])
    assert r.status_code == 200, r.text
    notes = " ".join(r.json()["notes"])
    assert "딜 소싱에 넣지 않습니다" in notes and "그런갈래" in notes
    assert rows_of(db, MNA) == []
    assert [x.vc_contact_id for x in rows_of(db, FOLLOW)] == [ok.id]


# ── ⑤ M&A 갈래 — 저절로 잇기 · 후보 확정 ────────────────────────────────────

def test_auto_link_pass_links_m_and_a_rows_by_name_and_phone(board, db, users):
    exact = vc(db, users, name="가투자", phone="010-1111-0001")
    phone_only = vc(db, users, name="나투자", phone="010-1111-0002")
    # 시리즈 A 에 이미 이어진 분도 M&A 에서 또 이어진다(갈래마다 한 줄)
    sc(db, name="가투자", phone="010-1111-0001", vc_contact_id=exact.id)
    a = sc(db, name="가투자", phone="01011110001", bucket=MNA, participation="참여")
    b = sc(db, name="나투자씨", phone="010-1111-0002", bucket=MNA, participation="참여")
    c = sc(db, name="라투자", phone=None, bucket=MNA, participation="불참여")
    sc(db, name="마해외", phone="010-1111-0009", bucket=OTHER)

    page = board.get("/sourcing", params={"tab": MNA}).text
    assert 'data-testid="sourcing-vc-auto-link"' in page
    assert 'data-testid="sourcing-vc-auto-link"' not in board.get(
        "/sourcing", params={"tab": OTHER}).text

    r = board.post("/sourcing/vc-auto-link", data={"tab": MNA}, follow_redirects=False)
    assert r.status_code == 303
    assert "1%EB%AA%85" in r.headers["location"]          # '1명을 … 이었습니다'
    db.expire_all()
    got = {x.id: (x.vc_contact_id, x.participation) for x in rows_of(db, MNA)}
    assert got == {a.id: (exact.id, "참여"), b.id: (None, "참여"), c.id: (None, "불참여")}

    # 번호만 같은 줄은 후보로 — [이 분으로 확정]
    from app.services import sourcing_sync
    assert [x["id"] for x in sourcing_sync.link_candidates(db, [b])[b.id]] == [phone_only.id]
    r = board.post(f"/api/sourcing/{b.id}/link", json={"contact_id": phone_only.id})
    assert r.status_code == 200, r.text
    assert rows_of(db, MNA)[1].vc_contact_id == phone_only.id
    assert rows_of(db, MNA)[1].participation == "참여"
    assert row_of(board, phone_only.id)["sourcing_branches"] == "M&A"


def test_candidates_skip_investors_already_linked_in_the_same_branch(board, db, users):
    from app.services import sourcing_sync

    contact = vc(db, users)
    sc(db, bucket=MNA, vc_contact_id=contact.id, name="홍길동")
    other = sc(db, bucket=MNA, name="홍길동", phone="010-9999-0000")
    in_sa = sc(db, name="홍길동", phone="010-9999-0000")
    picks = sourcing_sync.link_candidates(db, [other, in_sa])
    assert picks[other.id] == []                       # M&A 에서는 이미 이어진 분
    assert [x["id"] for x in picks[in_sa.id]] == [contact.id]   # 시리즈 A 에서는 후보


def test_m_and_a_rows_do_not_borrow_the_investor_room(board, db, users):
    """M&A 줄도 시리즈 A 처럼 — 자기 방이 없으면 `방 없음`, 투자사 방을 안 이어받는다."""
    from app.services import sourcing_link

    contact = vc(db, users)
    pick(board, contact.id, "Series A, M&A")
    sa, mna = rows_of(db, SYNC)[0], rows_of(db, MNA)[0]
    assert sourcing_link.linked_rooms(db, [sa, mna]) == {}


# ── ⑥ M&A 월간 리마인드 ─────────────────────────────────────────────────────

def test_m_and_a_remind_reads_its_own_template(db, users):
    from app.models import MessageTemplate
    from app.services import sourcing_msg

    user = users["u1"]
    body = sourcing_msg.remind_body_for(db, user, MNA)
    assert body.startswith("M&A 딜 네트워크에 함께해 주셔서 감사합니다.")
    assert "{월}" not in body
    # 구주·후속 · 대표 갈래는 예전 뼈대 그대로
    assert sourcing_msg.remind_body_for(db, user, FOLLOW).startswith("Deal Sourcing 네트워크")

    db.add(MessageTemplate(user_id=None, kind=sourcing_msg.REMIND_KIND, name=MNA,
                           is_active=1, body="M&A 저장 문구 {월}월"))
    db.commit()
    assert sourcing_msg.remind_body_for(db, user, MNA).startswith("M&A 저장 문구")


# ── ⑦ 이주 0114 ─────────────────────────────────────────────────────────────

NOW = "2026-10-21T00:00:00+09:00"


def _indexes(path) -> set:
    con = sqlite3.connect(path)
    try:
        return {r[1] for r in con.execute('PRAGMA index_list("sourcing_contacts")')}
    finally:
        con.close()


def _add(path, rows):
    con = sqlite3.connect(path)
    try:
        con.executemany(
            "INSERT INTO sourcing_contacts (id, bucket, position, name, room_verified, "
            "vc_contact_id, participation, vc_link_confirmed, created_at, updated_at) "
            f"VALUES (?, ?, 1, ?, 'unverified', ?, '참여', 0, '{NOW}', '{NOW}')", rows)
        con.commit()
    finally:
        con.close()


def test_migration_widens_the_index_and_refuses_to_go_down_with_twins(tmp_path):
    path = tmp_path / "m.db"
    done = _alembic(path, "upgrade", "0113_send_via_agent")
    assert done.returncode == 0, done.stdout + done.stderr
    # 운영처럼 옛 색인만 있게(새로 깐 DB 는 0001 이 지금 모델의 새 색인도 만든다)
    con = sqlite3.connect(path)
    con.execute("DROP INDEX IF EXISTS uq_sourcing_contacts_vc_contact_bucket")
    con.commit()
    con.close()
    assert "uq_sourcing_contacts_vc_contact" in _indexes(path)
    _add(path, [(1, SYNC, "가투자", 11), (2, MNA, "나투자", 12)])

    up = _alembic(path, "upgrade", "head")
    assert up.returncode == 0, up.stdout + up.stderr
    idx = _indexes(path)
    assert "uq_sourcing_contacts_vc_contact_bucket" in idx
    assert "uq_sourcing_contacts_vc_contact" not in idx

    # 같은 투자사 줄을 두 갈래가 — 이제 된다. 같은 갈래 둘은 여전히 안 된다.
    _add(path, [(3, MNA, "가투자", 11)])
    with pytest.raises(sqlite3.IntegrityError):
        _add(path, [(4, MNA, "가투자", 11)])

    # 내리기 — 겹친 줄이 있으면 거절하고 아무것도 안 바꾼다
    down = _alembic(path, "downgrade", "0113_send_via_agent")
    assert down.returncode != 0
    assert "투자사 줄 번호 11" in down.stdout + down.stderr
    assert "uq_sourcing_contacts_vc_contact_bucket" in _indexes(path)

    con = sqlite3.connect(path)
    con.execute("UPDATE sourcing_contacts SET vc_contact_id = NULL WHERE id = 3")
    con.commit()
    con.close()
    down = _alembic(path, "downgrade", "0113_send_via_agent")
    assert down.returncode == 0, down.stdout + down.stderr
    idx = _indexes(path)
    assert "uq_sourcing_contacts_vc_contact" in idx
    assert "uq_sourcing_contacts_vc_contact_bucket" not in idx


# ── ⑧ 대표 · 개인참여 (사용자 결정 2026-10-22) ──────────────────────────────

def test_ceo_and_solo_branches_check_uncheck_and_show_chips(board, db, users):
    contact = vc(db, users)
    # 운영처럼 이미 있는 갈래(한 칸 공백 글자) — 새 줄은 그 글자로 선다(탭이 둘로 안 갈린다)
    sc(db, name="가대표", phone="010-1111-0001", bucket="딜 소싱 참여 투자사 대표")

    r = pick(board, contact.id, "대표, 개인참여")
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["sourcing_branches"] == "대표, 개인참여"
    assert [(c["text"], c["on"]) for c in got["sourcing_chips"]] == [
        ("대표·제안 대기중", True), ("개인참여·제안 대기중", True)]
    assert got["sourcing_join"] == ""                 # 시리즈 A 칸은 그대로
    [ceo] = [x for x in rows_of(db, "딜 소싱 참여 투자사 대표") if x.vc_contact_id]
    [solo] = rows_of(db, SOLO)
    assert (ceo.participation, solo.participation) == (PENDING, PENDING)
    assert ceo.vc_contact_id == solo.vc_contact_id == contact.id
    assert rows_of(db, CEO) == []                     # 두 칸 공백 글자로 새 탭을 안 세웠다

    # 개인참여를 풀면 그 줄은 `불참여` — 지우지 않는다
    r = pick(board, contact.id, "대표")
    assert r.status_code == 200, r.text
    assert r.json()["sourcing_filter"] == "대표|개인참여 불참여"
    [solo] = rows_of(db, SOLO)
    assert (solo.vc_contact_id, solo.participation) == (contact.id, "불참여")

    # 엑셀 글자도 같은 짧은 이름
    from app.services import sourcing_sync
    assert sourcing_sync.parse_branches("대표, 개인참여 불참여") == {
        sourcing_sync.CEO_BUCKET: True, sourcing_sync.SOLO_BUCKET: False}


def test_the_table_and_edit_panel_offer_all_five_branches(board, db, users):
    vc(db, users)
    body = board.get(f"/contacts?sheet={SHEET}").text
    assert 'data-choices="Series A,M&amp;A,구주·후속,대표,개인참여"' in body
    boxes = [v for v in ("Series A", "M&amp;A", "구주·후속", "대표", "개인참여")
             if f'class="f-branch" value="{v}"' in body]
    assert len(boxes) == 5


def test_shared_fields_follow_into_ceo_and_solo_rows(board, db, users):
    from app.models import SourcingContact, VcContact

    contact = vc(db, users)
    pick(board, contact.id, "Series A, 개인참여")
    board.patch(f"/api/contacts/{contact.id}", json={"firm": "새이름벤처스"})
    [sa], [solo] = rows_of(db, SYNC), rows_of(db, SOLO)
    assert (sa.firm, solo.firm) == ("새이름벤처스", "새이름벤처스")
    # 개인참여 줄에서 고치면 투자사 줄과 시리즈 A 줄도
    r = board.patch(f"/api/sourcing/{solo.id}", json={"email": "solo@example.com"})
    assert r.status_code == 200, r.text
    db.expire_all()
    assert db.get(VcContact, contact.id).email == "solo@example.com"
    assert db.get(SourcingContact, sa.id).email == "solo@example.com"


def test_auto_link_pass_runs_on_ceo_and_solo_tabs(board, db, users):
    exact = vc(db, users, name="가투자", phone="010-1111-0001")
    a = sc(db, name="가투자", phone="01011110001", bucket=CEO, participation="참여")
    b = sc(db, name="나투자", phone="010-1111-0002", bucket=SOLO, participation="참여")
    for tab, row, want in ((CEO, a, exact.id), (SOLO, b, None)):
        page = board.get("/sourcing", params={"tab": tab}).text
        assert 'data-testid="sourcing-vc-auto-link"' in page, tab
        assert 'data-testid="sourcing-add-unsynced"' not in page, tab
        r = board.post("/sourcing/vc-auto-link", data={"tab": tab}, follow_redirects=False)
        assert r.status_code == 303
        db.expire_all()
        assert rows_of(db, tab)[0].vc_contact_id == want, tab
        assert rows_of(db, tab)[0].participation == "참여"       # 참여 칸은 그대로
    assert board.get("/api/contacts").json()["contacts"][0]["sourcing_branches"] == "대표"
    # 잇지 않는 갈래에서 누르면 막고 잇는 갈래 이름을 다 말한다
    r = board.post("/sourcing/vc-auto-link", data={"tab": OTHER}, follow_redirects=False)
    from urllib.parse import unquote
    assert "Series A · M&A · 구주·후속 · 대표 · 개인참여" in unquote(r.headers["location"])


def test_ceo_and_solo_sending_rules_stay_as_they_were(board, db, users):
    """잇기만 늘었다 — 방은 시리즈 A 처럼 투자사 방을 이어받지 않고,
    Series A 딜 공유는 시리즈 A 에만, 호칭 · 개수는 갈래 이름 그대로."""
    from app.services import series_a_share, sourcing_link, sourcing_msg

    contact = vc(db, users)
    pick(board, contact.id, "대표, 개인참여")
    ceo, solo = rows_of(db, CEO)[0], rows_of(db, SOLO)[0]
    assert sourcing_link.linked_rooms(db, [ceo, solo]) == {}
    assert not series_a_share.is_share_bucket(CEO) and not series_a_share.is_share_bucket(SOLO)
    assert (sourcing_msg.honorific(CEO), sourcing_msg.deal_count(CEO)) == ("대표", 5)
    assert (sourcing_msg.honorific(SOLO), sourcing_msg.deal_count(SOLO)) == ("심사역", 2)
