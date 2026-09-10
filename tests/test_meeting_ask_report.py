"""미팅 요청 표시가 **읽을 수 있게** 보이는가 — `/ir` 표와 업무 보고.

#162 가 「자료를 보내 놓고 미팅 요청을 안 한 건」을 세기 시작했다. 세는 것은
맞았는데 **보이는 자리 둘이 못 쓰게 돼 있었다.**

## 1) `/ir` 「전달한 자료」 표가 깨졌다

헤드리스 크롬(CDP)으로 잰 값이다. 판이 절반 폭(544px)이라 표가 544px 였다.

    칸      준 폭    속이 요구한 폭(`scrollWidth`)    결과
    상태    104px    147px                            `미팅 요청 보냄 · 202…` 에서 끊김
    단추    120px    207px                            `삭제` 가 통째로 안 보임

`table-layout: fixed` 라 칸은 안 늘어나고, `td.row-acts` 의
`white-space: nowrap` 이 접히지도 않게 막아서 **말없이 잘렸다.** 단추가
눌러야 하는 것인데 화면에 없었다.

## 2) 업무 보고에는 **숫자만** 있었다

`미팅 요청 안 보냄 N명` 은 있는데 **누구인지**는 없어서, 업무 보고를 보다가
딜 진행 관리로 건너가 다시 찾아야 했다.

## 여기서 못박는 것

- 「전달한 자료」 표가 **칸을 넘치지 않는다**(좁은 화면 포함)
- 세 단추가 다 **눌린다**
- 업무 보고에서 **누가** 안 보냈는지 알 수 있다
- 업무 보고 숫자와 `/ir` 숫자가 **같다**
- 화면과 **엑셀이 같다**
- 달이 걸치는 건을 어떻게 보이기로 했는가

날짜는 전부 인자로 넣거나 오늘에서 재서 만든다 — 특정 날에만 깨지지 않게.
"""
from __future__ import annotations

import re
from datetime import date, timedelta
from pathlib import Path

import pytest

from .conftest import DEMO_PASSWORD

CSS = Path("app/static/css/app.css")
IR_HTML = Path("app/templates/ir.html")
REPORT_HTML = Path("app/templates/report.html")

#: 헤드리스 크롬에서 잰 값 — 이 칸이 실제로 요구하는 폭(px).
#: `/ir` 에서 `미팅 요청 보냄 · 2026-09-10` 이 147px, 단추 셋이 207px 였다.
#: 업무 보고에는 한 갈래가 더 있어서(`미팅 요청 안 보냄 · 2026-09-17까지`)
#: 170px 이 필요하다 — `/ir` 은 그 글을 안 쓴다.
NEEDS = {"상태": 147, "단추": 207, "보고 상태": 170}


@pytest.fixture()
def logged(client, users):
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    return client


def _contact(db, users, name, *, firm="가나벤처스", user_key="u1"):
    from app.models import SheetOwner, VcContact

    if db.query(SheetOwner).filter_by(label="내 명단").first() is None:
        db.add(SheetOwner(label="내 명단", user_id=users["u1"].id))
    row = VcContact(user_id=users[user_key].id, name=name, title="심사역",
                    firm=firm, source_sheet="내 명단", channel_kakao=1,
                    connect_stage="connected", kakao_room_name=f"{name} 심사역님")
    db.add(row)
    db.commit()
    return row


def _delivered(db, users, contact, company, *, when, user_key="u1",
               requested=None):
    from app.models import IrRequest

    row = IrRequest(user_id=users[user_key].id, contact_id=contact.id,
                    company_name=company,
                    requested_at=(requested or when).isoformat(),
                    status="delivered", delivered_at=when.isoformat())
    db.add(row)
    db.commit()
    return row


def _meeting_send(db, users, contact, *, when):
    """앱이 실제로 내보낸 미팅 요청 한 건(stage 3 · sent)."""
    from app.models import SendItem, SendJob

    job = SendJob(user_id=users["u1"].id, kind="deal_intro", status="done")
    db.add(job)
    db.flush()
    db.add(SendItem(job_id=job.id, contact_id=contact.id, stage=3,
                    room_name=contact.kakao_room_name, message="미팅 가능하실지요",
                    status="sent", sent_at=f"{when.isoformat()}T11:00:00+09:00"))
    db.commit()


def _css() -> str:
    """주석을 지운 CSS — 규칙 위의 설명이 선택자로 세어지지 않게."""
    return re.sub(r"/\*.*?\*/", "", CSS.read_text(encoding="utf-8"), flags=re.S)


def _delivered_table_head() -> list:
    """「전달한 자료」 표의 `<th>` 들을 순서대로."""
    html = IR_HTML.read_text(encoding="utf-8")
    start = html.index('id="delivered-table"')
    head = html[start:html.index("</thead>", start)]
    # `<thead>` 가 같이 잡히지 않게 — `<th` 다음은 공백이나 `>` 다.
    return re.findall(r"<th(\s[^>]*)?>", head)


def _width(attrs: str) -> int:
    hit = re.search(r"width:\s*(\d+)px", attrs)
    return int(hit.group(1)) if hit else 0


# --- 1) 「전달한 자료」 표가 안 넘친다 ---------------------------------------

def test_상태_칸이_미팅_요청_한_줄을_담을_만큼_넓다():
    """★ `미팅 요청 보냄 · 2026-09-10` 이 147px 인데 104px 를 받고 있었다.

    **글자를 줄여 때우면 안 된다** — `미팅 요청 보냄` 을 `보냄` 으로 줄이면
    무엇을 보냈다는 말인지 알 수 없다. 자리를 만드는 쪽으로 푼다.
    """
    widths = [_width(a) for a in _delivered_table_head()]
    assert widths[3] >= NEEDS["상태"], (
        f"상태 칸이 {widths[3]}px — {NEEDS['상태']}px 이 필요하다(잰 값)")


def test_단추_칸이_단추_셋을_담을_만큼_넓다():
    """★ `미팅 잡기`·`보냈음으로 표시`·`삭제` 가 207px 인데 120px 를 받고 있었다.

    모자라면 `삭제` 가 통째로 화면 밖으로 나가 **누를 수가 없었다.**
    """
    widths = [_width(a) for a in _delivered_table_head()]
    assert widths[4] >= NEEDS["단추"], (
        f"단추 칸이 {widths[4]}px — {NEEDS['단추']}px 이 필요하다(잰 값)")


def test_담당자와_기업_칸에도_남는_폭이_있다():
    """고정폭만 늘리면 이번에는 이름이 짜부라진다.

    표에 `min-width` 를 줘서, 폭이 모자라면 칸을 눌러 부수는 대신 `.table-wrap`
    안에서 가로로 굴린다. 고정폭 셋 + 담당자·기업 120px 씩은 있어야 한다.
    """
    widths = [_width(a) for a in _delivered_table_head()]
    fixed = sum(widths)                       # 0 인 칸(담당자·기업)은 안 더해진다
    hit = re.search(r"\.table-wrap\s+#delivered-table\s*\{[^}]*min-width:\s*(\d+)px",
                    _css())
    assert hit, ".table-wrap 안의 #delivered-table 에 min-width 가 없다"
    assert int(hit.group(1)) >= fixed + 240, (
        f"min-width {hit.group(1)}px — 고정폭 {fixed}px 에 담당자·기업 몫이 없다")


def test_판이_한_줄을_통째로_쓴다():
    """★ 원인은 칸이 아니라 **판이 좁은 것**이었다.

    절반 폭(544px)에서는 어떤 너비를 줘도 다섯 칸이 안 들어간다 — 상태·단추를
    필요한 만큼 주면 담당자·기업에 44px 씩밖에 안 남는다.
    """
    html = IR_HTML.read_text(encoding="utf-8")
    start = html.index('id="delivered-table"')
    panel = html.rindex('<section class="panel', 0, start)
    assert 'class="panel wide"' in html[panel:panel + 60], \
        "「전달한 자료」 판이 한 줄을 다 쓰지 않는다"
    assert re.search(r"\.dash-grid\s+\.panel\.wide\s*\{[^}]*grid-column:\s*1\s*/\s*-1",
                     _css()), ".panel.wide 가 한 줄을 다 쓰는 규칙이 없다"


def test_단추_줄은_모자라면_접힌다():
    """`white-space: nowrap` 은 모자란 폭을 **말없이 삼킨다**.

    칸을 넓혀 놓았지만, 글자 크기나 문구가 조금만 달라져도 다시 소리 없이
    잘리는 자리다. 잘려 안 보이는 단추보다 두 줄이 낫다 — 다만 단추 하나가
    가운데서 끊기지는 않아야 한다.
    """
    css = _css()
    assert re.search(r"#delivered-table\s+td\.row-acts\s*\{[^}]*white-space:\s*normal",
                     css), "단추 칸이 여전히 nowrap 이라 넘치면 잘린다"
    assert re.search(r"#delivered-table\s+td\.row-acts\s+\.linkbtn\s*\{[^}]*"
                     r"white-space:\s*nowrap", css), \
        "단추 낱개가 가운데서 끊긴다"


def test_상태_배지는_여전히_한_덩어리다():
    """넓히면서 **원래 이유를 없애면 안 된다** — 배지가 중간에 끊기던 문제.

    `ir.html:473` 에 왜 칸 너비를 줬는지 적혀 있다. 그 이유는 그대로 둔다.
    """
    assert re.search(r"\.status-badge\s*\{[^}]*white-space:\s*nowrap", _css())
    widths = [_width(a) for a in _delivered_table_head()]
    assert widths[3], "상태 칸의 너비를 아예 없애면 배지가 제일 좁게 눌린다"


# --- 2) 세 단추가 다 눌린다 ---------------------------------------------------

def test_세_단추가_한_줄에_다_있다(db, users, logged):
    contact = _contact(db, users, "가담당")
    row = _delivered(db, users, contact, "샘플애그",
                     when=date.today() - timedelta(days=30))

    body = logged.get("/ir").text
    assert "js-book-meeting" in body                       # 미팅 잡기
    assert f"/ir/contacts/{contact.id}/meeting-asked" in body   # 보냈음으로 표시
    assert f"/ir/requests/{row.id}/delete" in body          # 삭제


def test_보냈음으로_표시가_먹는다(db, users, logged):
    contact = _contact(db, users, "가담당")
    _delivered(db, users, contact, "샘플애그",
               when=date.today() - timedelta(days=30))

    resp = logged.post(f"/ir/contacts/{contact.id}/meeting-asked",
                       follow_redirects=False)
    assert resp.status_code in (302, 303), resp.status_code

    body = logged.get("/ir").text
    assert "7일 지남" not in body
    assert "미팅 요청 보냄 · " in body


def test_삭제가_먹는다(db, users, logged):
    """★ 이 단추가 화면 밖으로 나가 있어서 **누를 수가 없었다.**"""
    contact = _contact(db, users, "가담당")
    # `샘플애그` 는 요청 등록 폼의 예시 글자로도 화면에 있다 — 그걸로 세면
    # 지워도 안 지워진 것처럼 보인다.
    row = _delivered(db, users, contact, "샘플로지",
                     when=date.today() - timedelta(days=30))

    assert "샘플로지" in logged.get("/ir").text
    resp = logged.post(f"/ir/requests/{row.id}/delete", follow_redirects=False)
    assert resp.status_code in (302, 303), resp.status_code
    assert "샘플로지" not in logged.get("/ir").text


# --- 3) 업무 보고에서 **누구인지** 보인다 -------------------------------------

def _old_month() -> date:
    """**확실히 지난 달**의 어느 날.

    `오늘이 든 달의 1일` 을 기준으로 잡으면 매달 1~7일에만 깨지는 검사가 된다
    — 그 날짜에 자료를 전달했다고 해도 아직 `7일` 이 안 지났기 때문이다.
    여기서 돌려주는 날은 오늘이 며칠이든 재촉할 때가 이미 지나 있다.
    """
    return date.today().replace(day=1) - timedelta(days=45)


def _report(logged, when):
    return logged.get(f"/report?month={when.year}-{when.month:02d}").text


def _bucket_rows(html: str, label: str) -> list:
    """`이 달의 반응` 의 한 갈래에 그려진 `<tr>` 들.

    갈래 머리(`bucket-head`)를 짚어 찾는다 — 위 안내문에도 같은 이름이
    나오므로 글자만 찾으면 엉뚱한 자리를 연다.
    """
    hit = re.search(r'bucket-head"[^>]*>\s*' + re.escape(label), html)
    assert hit, f"갈래를 못 찾았다: {label}"
    chunk = html[hit.end():]
    end = chunk.find("bucket-head")
    chunk = chunk[:end] if end > 0 else chunk
    if "<tbody>" not in chunk:
        return []
    body = chunk[chunk.index("<tbody>"):chunk.index("</tbody>")]
    return re.findall(r"<tr[^>]*>.*?</tr>", body, flags=re.S)


def test_업무_보고에_누가_안_보냈는지_적힌다(db, users, logged):
    """★ 숫자만으로는 대상을 알 수 없다 — 이름 옆에 적혀야 한다."""
    when = _old_month()
    late = _contact(db, users, "가담당", firm="가나벤처스")
    _delivered(db, users, late, "샘플애그", when=when)

    html = _report(logged, when)
    rows = _bucket_rows(html, "IR 요청 투자사")
    assert len(rows) == 1, rows
    assert "가담당" in rows[0] and "가나벤처스" in rows[0]
    assert "미팅 요청 안 보냄" in rows[0], "누구 줄인지 알 수 없다"
    assert "overdue-row" in rows[0], "지난 줄이 눈에 안 걸린다"


def test_보낸_담당자에게는_보냄과_날짜가_적힌다(db, users, logged):
    when = _old_month()
    contact = _contact(db, users, "나담당")
    _delivered(db, users, contact, "샘플메디", when=when)
    _meeting_send(db, users, contact, when=when + timedelta(days=1))

    rows = _bucket_rows(_report(logged, when), "IR 요청 투자사")
    asked_at = (when + timedelta(days=1)).isoformat()
    assert f"미팅 요청 보냄 · {asked_at}" in rows[0]
    assert "overdue-row" not in rows[0]


def test_아직_안_보낸_요청_줄에는_안_붙는다(db, users, logged):
    """자료도 아직 안 줬는데 "미팅 요청 안 보냄" 은 앞뒤가 없는 말이다."""
    from app.models import IrRequest

    when = _old_month()
    contact = _contact(db, users, "다담당")
    db.add(IrRequest(user_id=users["u1"].id, contact_id=contact.id,
                     company_name="샘플페이", requested_at=when.isoformat(),
                     status="open"))
    db.commit()

    rows = _bucket_rows(_report(logged, when), "IR 요청 투자사")
    assert "미팅 요청" not in rows[0], rows[0]


def test_담당자_하나가_여러_건_받으면_줄마다_같은_표시가_붙는다(db, users, logged):
    """나가는 카톡은 한 통이라 **숫자는 1**이다. 줄은 셋이라 표시도 셋이다.

    `/ir` 의 `전달한 자료` 표와 같은 방식이라 두 화면이 같아 보인다.
    """
    when = _old_month()
    contact = _contact(db, users, "가담당")
    for name in ("샘플애그", "샘플메디", "샘플페이"):
        _delivered(db, users, contact, name, when=when)

    html = _report(logged, when)
    rows = _bucket_rows(html, "IR 요청 투자사")
    assert len(rows) == 3
    assert all("미팅 요청 안 보냄" in r for r in rows)
    assert _n(html, r"미팅 요청 안 보냄</span> ?<b[^>]*>(\d+)명") == 1, \
        "담당자 수(1명)가 아니라 줄 수로 세고 있다"


# --- 4) 두 화면의 숫자가 같다 -------------------------------------------------

def _n(html: str, pattern: str) -> int:
    hit = re.search(pattern, re.sub(r"\s+", " ", html))
    assert hit, f"숫자를 못 찾았다: {pattern}"
    return int(hit.group(1))


def test_업무_보고_숫자와_ir_숫자가_같다(db, users, logged):
    """★ 두 화면이 각자 세면 반드시 갈린다 — 이 저장소가 반복해 겪은 사고다.

    담당자를 셋 만들어 셋 다 다른 상태로 둔다(지남 · 아직 · 보냄).

    **자료를 전부 한 달 안에 둔다.** `/ir` 은 남아 있는 요청 전부를 세고
    업무 보고는 **그 달치**를 세므로, 달이 섞이면 두 수가 달라도 맞다.
    여기서 보려는 것은 그 차이가 아니라 **판정이 한 곳인가**다 — 같은 자료를
    주면 두 화면이 같은 수를 말해야 한다.
    """
    when = _old_month()
    late = _contact(db, users, "가담당")
    _delivered(db, users, late, "샘플애그", when=when)
    _delivered(db, users, late, "샘플메디", when=when)   # 같은 사람, 두 줄

    # 요청은 그 달에 받았고 **자료는 오늘 보냈다** — 아직 재촉할 때가 아니다.
    # 달을 가르는 것은 `requested_at`, 재촉할 때를 재는 것은 `delivered_at`.
    soon = _contact(db, users, "나담당")
    _delivered(db, users, soon, "샘플페이", when=date.today(), requested=when)

    done = _contact(db, users, "다담당")
    _delivered(db, users, done, "샘플로지", when=when)
    _meeting_send(db, users, done, when=when + timedelta(days=1))

    ir = logged.get("/ir").text
    report = _report(logged, when)

    assert _n(ir, r"미팅 요청 안 보냄 (\d+)명") == \
        _n(report, r"미팅 요청 안 보냄</span> ?<b[^>]*>(\d+)명")
    assert _n(ir, r"일 지남 (\d+)명\s*\)") == \
        _n(report, r"일 지남 (\d+)명\s*\)")


def test_보고_줄에_뜨는_사람과_센_사람이_같다(db, users, logged):
    """숫자와 목록이 갈리면 어느 쪽을 믿을지 알 수 없다.

    `안 보냄` 표시가 붙은 **담당자**(줄이 아니다)의 수가 그 숫자여야 한다.
    """
    when = _old_month()
    for name in ("가담당", "나담당"):
        c = _contact(db, users, name)
        _delivered(db, users, c, "샘플애그", when=when)
        _delivered(db, users, c, "샘플메디", when=when)   # 사람마다 두 줄
    done = _contact(db, users, "다담당")
    _delivered(db, users, done, "샘플페이", when=when)
    _meeting_send(db, users, done, when=when)

    html = _report(logged, when)
    rows = _bucket_rows(html, "IR 요청 투자사")
    who = {n for n in ("가담당", "나담당", "다담당")
           for r in rows if n in r and "미팅 요청 안 보냄" in r}
    assert len(who) == _n(html, r"미팅 요청 안 보냄</span> ?<b[^>]*>(\d+)명")
    assert len(rows) == 5, "줄 수와 사람 수가 같으면 이 검사가 아무것도 안 지킨다"


# --- 5) 화면과 엑셀이 같다 ----------------------------------------------------

def _reactions_sheet(logged, when):
    import io

    import openpyxl

    resp = logged.get(f"/api/export/report.xlsx?month={when.year}-{when.month:02d}")
    assert resp.status_code == 200, resp.status_code
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    name = [s for s in wb.sheetnames if "반응" in s][0]
    return wb[name]


def _cells(ws) -> list:
    return [str(c.value) for row in ws.iter_rows() for c in row if c.value]


def test_엑셀에도_같은_말이_실린다(db, users, logged):
    """★ 화면에만 있으면 파일로 받아 보는 사람은 그 건을 영영 못 본다."""
    when = _old_month()
    late = _contact(db, users, "가담당")
    _delivered(db, users, late, "샘플애그", when=when)
    done = _contact(db, users, "나담당")
    _delivered(db, users, done, "샘플메디", when=when)
    _meeting_send(db, users, done, when=when + timedelta(days=1))

    cells = _cells(_reactions_sheet(logged, when))
    joined = "\n".join(cells)
    asked_at = (when + timedelta(days=1)).isoformat()
    assert "미팅 요청 안 보냄" in joined
    assert f"미팅 요청 보냄 · {asked_at}" in joined
    # 상태와 미팅 요청이 **한 칸**에 있다(칸을 새로 만들지 않았다).
    assert any("전달함" in c and "미팅 요청" in c for c in cells)


def test_화면_글자와_엑셀_글자가_한_자도_안_다르다(db, users, logged):
    """글자를 두 곳에서 지으면 갈린다 — 짓는 곳은 `report.meeting_ask_note` 하나다."""
    from app.services import report as report_svc

    when = _old_month()
    late = _contact(db, users, "가담당")
    _delivered(db, users, late, "샘플애그", when=when)
    # 세 갈래가 **한 달 안에** 다 서게 한다 — 요청은 그 달, 전달은 오늘.
    soon = _contact(db, users, "나담당")
    _delivered(db, users, soon, "샘플메디", when=date.today(), requested=when)
    done = _contact(db, users, "다담당")
    _delivered(db, users, done, "샘플페이", when=when)
    _meeting_send(db, users, done, when=when + timedelta(days=1))

    data = report_svc.monthly(db, when.year, when.month, users["u1"])
    said = {r["ask"] for b in data["buckets"] if b["key"] == "ir"
            for r in b["rows"] if r["ask"]}
    assert len(said) == 3, said

    html = _report(logged, when)
    joined = "\n".join(_cells(_reactions_sheet(logged, when)))
    for text in said:
        assert text in html, f"화면에 없다: {text}"
        assert text in joined, f"엑셀에 없다: {text}"


# --- 6) 달이 걸치는 건 -------------------------------------------------------

def test_지난달_자료에_이번_달_미팅_요청이면_지난달_보고에_보냄으로_뜬다(db, users, logged):
    """★ 어떻게 정했는가.

    8월에 받아 전달한 요청인데 미팅 요청은 9월에 나갔다면, **8월 보고에
    `미팅 요청 보냄 · 9월 날짜` 로 뜬다.** 달을 자르지 않는다.

      · 이 줄은 기록이 아니라 **아직 남은 일**이다. 이미 보낸 건이 지난달
        보고에 영영 "안 보냄" 으로 남아 있으면 거짓 경보고, 거짓으로 뜨는
        숫자는 곧 아무도 안 본다.
      · 숫자(`ir_meeting_ask_missing`)가 이미 그렇게 센다. 줄만 달을 잘라
        보이면 **같은 화면 안에서 숫자와 목록이 갈린다.**
    """
    last = (date.today().replace(day=1) - timedelta(days=1)).replace(day=1)
    contact = _contact(db, users, "가담당")
    _delivered(db, users, contact, "샘플애그", when=last)
    asked = date.today().replace(day=1)          # 이번 달에 보냈다
    _meeting_send(db, users, contact, when=asked)

    html = _report(logged, last)
    rows = _bucket_rows(html, "IR 요청 투자사")
    assert f"미팅 요청 보냄 · {asked.isoformat()}" in rows[0]
    assert _n(html, r"미팅 요청 안 보냄</span> ?<b[^>]*>(\d+)명") == 0
    assert "overdue-row" not in rows[0]


def test_달을_자르지_않는다는_것이_숫자와_줄에_똑같이_적용된다(db, users, logged):
    """줄만 달을 자르면 숫자와 목록이 갈린다. 둘이 같은 판정을 읽는지 본다."""
    last = (date.today().replace(day=1) - timedelta(days=1)).replace(day=1)
    still = _contact(db, users, "가담당")
    _delivered(db, users, still, "샘플애그", when=last)
    late = _contact(db, users, "나담당")
    _delivered(db, users, late, "샘플메디", when=last)
    _meeting_send(db, users, late, when=date.today().replace(day=1))

    html = _report(logged, last)
    rows = _bucket_rows(html, "IR 요청 투자사")
    marked = [r for r in rows if "미팅 요청 안 보냄" in r]
    assert len(marked) == _n(html, r"미팅 요청 안 보냄</span> ?<b[^>]*>(\d+)명") == 1
    assert "가담당" in marked[0]


def test_그_달에_받은_수는_안_변한다(db, users, logged):
    """변하는 것은 "지금 남은 일" 쪽뿐이다 — 그 달에 무엇을 받았는지는 기록이다."""
    last = (date.today().replace(day=1) - timedelta(days=1)).replace(day=1)
    contact = _contact(db, users, "가담당")
    _delivered(db, users, contact, "샘플애그", when=last)
    _meeting_send(db, users, contact, when=date.today().replace(day=1))

    html = re.sub(r"\s+", " ", _report(logged, last))
    assert re.search(r"요청받음</span> ?<b[^>]*> ?1 ?</b>", html), html[:200]


# --- 안내문이 그리로 보낸다 ---------------------------------------------------

def test_보고_안내문이_어디를_보라고_말한다(db, users, logged):
    """숫자만 던져 놓으면 사람은 그 다음에 무엇을 할지 모른다."""
    when = _old_month()
    contact = _contact(db, users, "가담당")
    _delivered(db, users, contact, "샘플애그", when=when)

    html = _report(logged, when)
    assert 'href="#reactions"' in html
    assert 'id="reactions"' in html


def test_보고_표의_상태_칸도_넉넉하다():
    """`/ir` 에서 끊겼던 것과 **같은 글**이 여기에도 실린다."""
    html = REPORT_HTML.read_text(encoding="utf-8")
    start = html.index("이 달의 반응")
    head = html[start:html.index("</thead>", start)]
    widths = [_width(a) for a in re.findall(r"<th(\s[^>]*)?>", head)]
    status = [w for w in widths if w]
    assert max(status) >= NEEDS["보고 상태"], (
        f"보고 표의 상태 칸이 좁다({status}) — 날짜가 `2026-09-17까…` 로 끊긴다")
