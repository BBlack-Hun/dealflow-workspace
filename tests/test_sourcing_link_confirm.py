"""딜 소싱 ↔ 투자사 관리 현황 잇기 — **이름+번호일 때만 저절로, 아니면 사람이 확정**(#28).

사용자가 정했다(2026-10-16).

  ① 저절로 잇는 것은 번호가 투자사 쪽 딱 한 줄에 맞고 그 줄의 이름도 같을 때뿐이다
     (`sourcing_sync.auto_link`). 이름은 `_key`(띄어쓰기 빼고 소문자) — 투자사 쪽
     O 가 같은 분을 찾는 잣대(`find_row`)와 같다. 이미 이어 둔 줄은 풀지 않는다.
  ② 이름만 · 번호만 같거나 번호가 여러 줄에 맞으면 잇지 않고 **후보**로 띄운다
     (`link_candidates`) — 이미 다른 소싱 줄과 이어진 투자사 줄 · 감춘 줄 ·
     투자사로 세지 않는 명단의 줄은 후보가 아니다.
  ③ [이 분으로 확정](`POST /api/sourcing/{id}/link`) — 서버가 후보를 다시 세어 그
     안에 있을 때만 잇는다. 시리즈 A 갈래 · 안 이어진 줄만. 수정 로그에 남는다.
  ④ 투자사 관리 현황에 `참여 요청일` — 이어진 줄의 `참여 요청일` 을 날짜만(시트가
     붙여 온 `00:00:00` 같은 시각은 뗀다), 읽기만. `Series A 딜 소싱` 바로 뒤.
     엑셀에도 두 칸이 같이 나간다.
  ⑤ [연결 풀기](`POST /api/sourcing/{id}/unlink`) — 관리자 · 이어진 투자사 줄의 담당만
     (그 줄의 주인 번호로 — 누구나 고치는 `담당자` 칸으로는 판정하지 않는다). 두 줄
     다 지우지 않고 잇기만 푼다. 수정 로그에 남는다(사용자 결정 2026-10-17).
  ⑥ 시트를 다시 가져와도(`scripts/import_sourcing.py`) **사람이 확정한 잇기**는
     이름 · 번호가 둘 다 맞지 않아도 그대로 되살린다 — 다시 후보로 뜨지 않는다.
     0104 전부터 이어져 있던 줄도 확정으로 친다. 감춰졌거나 투자사로 세지 않는
     명단으로 간 투자사 줄에는 되살리지 않는다.

이름 · 회사 · 번호는 전부 가상값이다(공개 저장소).
"""
from __future__ import annotations

import json
import re
import shutil
import subprocess
from pathlib import Path

import pytest

from tests.test_sourcing_vc_sync import (  # noqa: F401 - 픽스처를 함께 쓴다
    MNA, SHEET, SYNC, admin, board, column, login, mark, sc, sync_rows, vc)

OTHER_SHEET = "윤서아 투자사"


def picks(db, row):
    from app.services import sourcing_sync

    db.expire_all()
    return sourcing_sync.link_candidates(db, [row]).get(row.id)


def confirm(client, row_id, contact_id):
    return client.post(f"/api/sourcing/{row_id}/link", json={"contact_id": contact_id})


def other_list(db, users):
    from app.models import SheetOwner

    db.add(SheetOwner(label=OTHER_SHEET, user_id=users["u2"].id, layout="investor"))
    db.commit()


# ── ① 저절로 잇는 것은 이름 + 번호 ──────────────────────────────────────────

def test_name_and_phone_together_link_on_their_own(board, db, users):
    """번호 표기(하이픈) · 이름 띄어쓰기가 달라도 같은 분이다(`_key` · 숫자만)."""
    contact = vc(db, users, name="홍길동", phone="010-5555-6666")
    r = board.post(f"/api/sourcing?bucket={SYNC}",
                   json={"name": "홍 길동", "phone": "01055556666"})
    assert r.status_code == 200, r.text
    assert SHEET in r.json()["linked_msg"]
    [row] = sync_rows(db)
    assert row.vc_contact_id == contact.id
    assert column(board, contact.id) == "O"


def test_a_phone_alone_no_longer_links(board, db, users):
    """번호는 맞는데 이름이 다르다 — 잇지 않고 후보(번호만 같음)로 띄운다."""
    contact = vc(db, users, name="김서연", phone="010-5555-6666")
    r = board.post(f"/api/sourcing?bucket={SYNC}",
                   json={"name": "홍길동", "phone": "010-5555-6666"})
    assert r.status_code == 200, r.text
    assert "linked_msg" not in r.json()
    [row] = sync_rows(db)
    assert row.vc_contact_id is None
    assert column(board, contact.id) == ""
    [cand] = picks(db, row)
    assert (cand["id"], cand["name_match"], cand["phone_match"]) == (contact.id, False, True)
    assert cand["label"] == "김서연 심사역 (가나다벤처스)"
    assert cand["sheet"] == SHEET


def test_a_name_alone_is_only_a_candidate(board, db, users):
    """이름만 같다(번호가 다르거나 없다) — 예전처럼 잇지 않고, 이제는 후보로 뜬다."""
    contact = vc(db, users, name="홍길동", phone="010-1234-5678")
    other = sc(db, phone="010-9999-0000")
    no_phone = sc(db, phone=None, position=2)
    for row in (other, no_phone):
        [cand] = picks(db, row)
        assert (cand["id"], cand["name_match"], cand["phone_match"]) == (contact.id, True, False)
    assert all(r.vc_contact_id is None for r in sync_rows(db))


def test_a_phone_on_several_investor_rows_is_left_to_a_person(board, db, users):
    """같은 분이 두 팀원 명단에 한 줄씩 — 이름 · 번호가 둘 다 맞아도 어느 줄인지
    코드가 고르지 않는다. 둘 다 후보로 뜨고, 사람이 고른 줄과 잇는다."""
    other_list(db, users)
    mine = vc(db, users)
    theirs = vc(db, users, who="u2", sheet=OTHER_SHEET)
    board.post(f"/api/sourcing?bucket={SYNC}",
               json={"name": "홍길동", "phone": "010-1234-5678"})
    [row] = sync_rows(db)
    assert row.vc_contact_id is None
    got = picks(db, row)
    assert {c["id"] for c in got} == {mine.id, theirs.id}
    assert all(c["name_match"] and c["phone_match"] for c in got)
    assert {c["sheet"] for c in got} == {SHEET, OTHER_SHEET}

    r = confirm(board, row.id, theirs.id)
    assert r.status_code == 200, r.text
    assert r.json()["vc_link"]["sheet"] == OTHER_SHEET
    [row] = sync_rows(db)
    assert row.vc_contact_id == theirs.id


def test_candidates_are_ordered_both_then_phone_then_name(board, db, users):
    both = vc(db, users, name="홍길동", phone="010-1234-5678")
    phone_only = vc(db, users, name="김서연", phone="010-1234-5678")
    name_only = vc(db, users, name="홍길동", phone="010-2222-3333", firm="마바벤처스")
    row = sc(db)
    assert [c["id"] for c in picks(db, row)] == [both.id, phone_only.id, name_only.id]


def test_rows_already_taken_or_not_investors_are_not_candidates(board, db, users):
    """이미 다른 소싱 줄과 이어진 투자사 줄은 후보가 아니다 — 한 분께 딜 소싱 제안이
    두 번 나간다. 감춘 줄 · 투자사로 세지 않는 명단의 줄도 아니다(저절로 잇기와
    같은 모집단)."""
    from app.models import SheetOwner

    db.add(SheetOwner(label="스타트업 명단", user_id=users["u1"].id,
                      layout="startup", is_hidden=1))
    db.commit()
    taken = vc(db, users, phone="010-1111-0001")
    sc(db, name="홍길동", phone="010-1111-0001", vc_contact_id=taken.id)
    vc(db, users, phone="010-1111-0002", is_hidden=1)
    vc(db, users, phone="010-1111-0003", sheet="스타트업 명단")
    free = vc(db, users, phone="010-1111-0004")

    row = sc(db, phone="010-9999-0000", position=2)
    assert [c["id"] for c in picks(db, row)] == [free.id]


def test_existing_links_are_not_undone(board, db, users):
    """번호로만 이어 둔 옛 줄(이주 0093 · #28 전)은 그대로다 — 이름이 달라도."""
    contact = vc(db, users, name="홍길순")
    row = sc(db, vc_contact_id=contact.id)
    board.patch(f"/api/sourcing/{row.id}", json={"memo": "통화함"})
    [row] = sync_rows(db)
    assert row.vc_contact_id == contact.id
    assert column(board, contact.id) == "O"
    assert picks(db, row) is None                 # 이어진 줄에는 후보 자리가 없다
    page = board.get(f"/sourcing?tab={SYNC}").text
    assert f"투자사 관리 현황: {SHEET} 명단" in page
    assert 'data-testid="sourcing-vc-picks"' not in page


def test_fixing_the_name_links_it_and_answers_with_the_badge(board, db, users):
    """이름도 잣대라 이름을 고쳐도 다시 본다. 응답이 연결 표시 · 채운 칸을 싣는다."""
    contact = vc(db, users, name="홍길동", phone="010-7777-8888", email="hong@example.com")
    row = sc(db, name="홍길둥", phone="010-7777-8888", email=None)
    assert picks(db, row)[0]["phone_match"] is True

    r = board.patch(f"/api/sourcing/{row.id}", json={"name": "홍길동"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body.get("linked_msg")
    assert body["vc_link"]["sheet"] == SHEET
    assert body["fields"]["email"] == "hong@example.com"
    assert column(board, contact.id) == "O"


def test_a_change_that_does_not_link_answers_with_new_candidates(board, db, users):
    contact = vc(db, users, name="김서연", phone="010-7777-8888")
    row = sc(db, name="홍길동", phone="010-0000-0000")
    assert picks(db, row) == []

    r = board.patch(f"/api/sourcing/{row.id}", json={"phone": "010-7777-8888"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert "linked_msg" not in body and "vc_link" not in body
    assert [c["id"] for c in body["vc_picks"]] == [contact.id]

    # 이름 · 번호가 그대로인 고침(메모)은 후보를 다시 세지 않는다
    r = board.patch(f"/api/sourcing/{row.id}", json={"memo": "통화함"})
    assert "vc_picks" not in r.json()


def test_other_buckets_never_get_candidates(board, db, users):
    vc(db, users)
    row = sc(db, bucket=MNA)
    assert picks(db, row) is None
    page = board.get(f"/sourcing?tab={MNA}").text
    assert 'data-testid="sourcing-vc-picks"' not in page


# ── ③ [이 분으로 확정] ──────────────────────────────────────────────────────

def test_confirm_links_fills_blank_cells_and_shows_the_badge(board, db, users):
    contact = vc(db, users, name="홍길동", phone="010-1234-5678",
                 email="hong@example.com", firm="가나다벤처스")
    row = sc(db, phone="010-9999-0000", email=None, firm="가나다벤처스(주)")

    page = board.get(f"/sourcing?tab={SYNC}").text
    box = re.search(r"data-vc-picks='([^']*)'", page)
    assert box, "후보 칸이 안 섰다"
    assert [c["id"] for c in json.loads(box.group(1).replace("&#34;", '"'))] == [contact.id]

    r = confirm(board, row.id, contact.id)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["vc_link"]["sheet"] == SHEET
    assert "?sheet=" in body["vc_link"]["href"]
    assert set(body["vc_link"]["diff"]) == {"휴대폰", "회사"}   # 다르게 적힌 값은 그대로
    assert "두 화면이 같이 바뀌는 칸" in body["vc_link"]["title"]
    assert body["fields"]["email"] == "hong@example.com"       # 빈 칸만 채운다
    assert body["fields"]["firm"] == "가나다벤처스(주)"

    [row] = sync_rows(db)
    assert row.vc_contact_id == contact.id
    assert column(board, contact.id) == "O"
    page = board.get(f"/sourcing?tab={SYNC}").text
    assert f"투자사 관리 현황: {SHEET} 명단" in page
    assert 'data-testid="sourcing-vc-picks"' not in page


def test_confirm_is_in_the_edit_log(board, db, users):
    from app.models import EditLog

    contact = vc(db, users, name="김서연")
    row = sc(db)                                   # 번호만 같다 — 후보
    rid, cid = row.id, contact.id
    assert confirm(board, rid, cid).status_code == 200
    logs = db.query(EditLog).filter_by(table_name="sourcing_contacts", row_id=rid,
                                       action="update").all()
    changes = [c for log in logs for c in json.loads(log.changes_json)]
    assert {"field": "vc_contact_id", "before": None, "after": cid} in changes


def test_confirm_refuses_what_is_not_a_candidate(board, db, users):
    contact = vc(db, users, name="홍길동", phone="010-1234-5678")
    stranger = vc(db, users, name="박지훈", phone="010-4444-5555")
    row = sc(db, phone="010-9999-0000")

    for cid in (stranger.id, None, 999999):
        r = confirm(board, row.id, cid)
        assert r.status_code == 400, (cid, r.text)
        assert "후보가 아닙니다" in r.json()["detail"]
    assert sync_rows(db)[0].vc_contact_id is None

    # 그 사이 다른 소싱 줄이 그 투자사 줄과 이어졌다 — 화면의 옛 후보는 막힌다
    sc(db, name="홍길동", phone="010-1234-5678", position=2, vc_contact_id=contact.id)
    r = confirm(board, row.id, contact.id)
    assert r.status_code == 400
    assert sync_rows(db)[0].vc_contact_id is None


def test_two_people_confirming_the_same_investor_row_at_once(board, db, users, monkeypatch):
    """두 사람이 같은 투자사 줄을 거의 동시에 [이 분으로 확정] — 둘 다 후보를 세는
    순간에는 그 줄이 비어 있었다. 늦게 커밋한 쪽은 유일 색인에 걸리는데, 500 이 아니라
    다른 막힘과 같은 400 '새로고침해 보세요' 로 답하고 아무것도 잇지 않는다."""
    from app.services import sourcing_sync

    contact = vc(db, users, name="김서연")          # 번호만 같다 — 두 줄 다의 후보
    first = sc(db)
    second = sc(db, position=2)
    stale = {second.id: picks(db, second)}         # 늦은 쪽 화면이 본 후보
    assert confirm(board, first.id, contact.id).status_code == 200

    real = sourcing_sync.link_candidates
    monkeypatch.setattr(sourcing_sync, "link_candidates",
                        lambda db_, rows: {r.id: stale.get(r.id) or [] for r in rows}
                        if any(r.id in stale for r in rows) else real(db_, rows))
    r = confirm(board, second.id, contact.id)
    assert r.status_code == 400, r.text
    assert "새로고침" in r.json()["detail"]
    rows = sync_rows(db)
    assert [x.vc_contact_id for x in rows] == [contact.id, None]


def test_confirm_refuses_linked_rows_and_other_buckets(board, db, users):
    first = vc(db, users, name="홍길동", phone="010-1234-5678")
    second = vc(db, users, name="홍길동", phone="010-2222-3333", firm="마바벤처스")
    linked = sc(db, vc_contact_id=first.id)
    r = confirm(board, linked.id, second.id)
    assert r.status_code == 400
    assert "이미" in r.json()["detail"]
    assert sync_rows(db)[0].vc_contact_id == first.id

    other = sc(db, bucket=MNA, phone="010-2222-3333")
    r = confirm(board, other.id, second.id)
    assert r.status_code == 400
    assert SYNC in r.json()["detail"]

    assert confirm(board, 999999, second.id).status_code == 404


def test_confirm_is_open_to_anyone_logged_in_like_editing_the_table(board, db, users):
    """권한은 딜 소싱 표를 고치는 것과 같다 — 로그인한 누구나(남의 투자사 줄이어도).
    로그인하지 않았으면 못 한다."""
    contact = vc(db, users, name="김서연")          # u1 의 투자사 줄
    row = sc(db)                                   # 번호만 같다 — 후보

    board.post("/logout")
    assert confirm(board, row.id, contact.id).status_code != 200
    assert sync_rows(db)[0].vc_contact_id is None

    login(board, "01000000002")                    # u2
    assert confirm(board, row.id, contact.id).status_code == 200
    assert sync_rows(db)[0].vc_contact_id == contact.id


# ── ④ 투자사 관리 현황의 `참여 요청일` ──────────────────────────────────────

def _requested(client, contact_id):
    rows = client.get("/api/contacts").json()["contacts"]
    return next(r for r in rows if r["id"] == contact_id)["sourcing_requested_at"]


def test_the_requested_date_shows_only_for_linked_rows_as_written(board, db, users):
    linked = vc(db, users, name="홍길동")
    alone = vc(db, users, name="김서연", phone="010-2222-3333")
    other_bucket = vc(db, users, name="박지훈", phone="010-3333-4444")
    sc(db, vc_contact_id=linked.id, requested_at="2026/05/21")
    # 다른 갈래의 줄은 이어진 줄이 아니다(`Series A 딜 소싱` 칸과 같은 판정)
    sc(db, name="박지훈", phone="010-3333-4444", bucket=MNA,
       vc_contact_id=other_bucket.id, requested_at="2026-01-01")

    assert _requested(board, linked.id) == "2026/05/21"      # 적힌 표기 그대로
    assert _requested(board, alone.id) == ""
    assert _requested(board, other_bucket.id) == ""

    body = board.get(f"/contacts?sheet={SHEET}").text
    assert 'data-filters="sourcing_date:참여 요청일"' in body
    assert 'data-f-sourcing_date="2026/05/21"' in body
    assert '<td class="sourcing-requested" title="2026/05/21">2026/05/21</td>' in body


def test_the_requested_date_drops_the_time_the_sheet_carried(board, db, users):
    """시트 가져오기가 날짜 칸을 `2026-08-11 00:00:00` 처럼 시각까지 적어 두었다 —
    두 화면 다 날짜만 보인다(저장된 글자는 그대로). 날짜가 아닌 글은 그대로."""
    from app.services import sourcing_sync

    assert sourcing_sync.requested_label("2026-08-11 00:00:00") == "2026-08-11"
    assert sourcing_sync.requested_label(" 2026/5/2 14:30 ") == "2026/5/2"
    assert sourcing_sync.requested_label("2026/05/21") == "2026/05/21"
    assert sourcing_sync.requested_label("5월 초") == "5월 초"
    assert sourcing_sync.requested_label(None) == ""

    contact = vc(db, users)
    row = sc(db, vc_contact_id=contact.id, requested_at="2026-08-11 00:00:00")
    assert _requested(board, contact.id) == "2026-08-11"
    page = board.get(f"/sourcing?tab={SYNC}").text
    assert '<td class="cell" data-field="requested_at">2026-08-11</td>' in page
    assert 'data-f-requested_at="2026-08-11"' in page
    assert "00:00:00" not in page
    assert sync_rows(db)[0].requested_at == "2026-08-11 00:00:00"   # 고쳐 쓰지 않는다
    assert row.id


def test_the_requested_date_sits_right_after_series_a_and_is_read_only(board, db, users):
    """머리글 · 행 칸 차례가 같아야 한다 — `Series A 딜 소싱` 바로 뒤. 칸에는
    `.cell` · `data-field` 가 없다(눌러 고치는 `inline_edit.js` 가 안 잡는다)."""
    contact = vc(db, users)
    sc(db, vc_contact_id=contact.id, requested_at="2026-07-01 00:00:00")
    body = board.get(f"/contacts?sheet={SHEET}").text

    head = body[body.index("<thead>"):body.index("</thead>")]
    labels = [re.sub(r"<[^>]+>", "", x).strip()
              for x in re.findall(r"<th[\s>].*?</th>", head, re.S)]
    at = labels.index("딜소싱 참여 투자사")
    assert labels[at + 1:at + 3] == ["Series A 딜 소싱", "참여 요청일"], labels

    rows = body[body.index("<tbody>"):body.index("</tbody>")]
    first_row = rows[:rows.index("</tr>")]
    tds = re.findall(r"<td(?:\s[^>]*)?>", first_row)
    i = next(n for n, td in enumerate(tds) if 'data-field="sourcing_join"' in td)
    assert tds[i + 1].startswith('<td class="sourcing-requested"'), tds[i + 1]
    assert 'class="cell' not in tds[i + 1] and "data-field" not in tds[i + 1]
    assert len(re.findall(r"<th[\s>]", head)) == len(re.findall(r"<td[\s>]", first_row))

    # 수정창에는 없다 — 고칠 칸이 아니다(`딜소싱 참여 투자사` 다음은 여전히 O/X)
    panel = re.findall(r'id="f-(\w+)"', body)
    assert "sourcing_requested_at" not in panel
    assert panel[panel.index("sourcing_note") + 1] == "sourcing_join"


def test_patch_ignores_the_requested_date(board, db, users):
    contact = vc(db, users)
    sc(db, vc_contact_id=contact.id, requested_at="2026/05/21")
    r = board.patch(f"/api/contacts/{contact.id}",
                    json={"sourcing_requested_at": "1999-01-01", "memo": "통화함"})
    assert r.status_code == 200, r.text
    assert r.json()["sourcing_requested_at"] == "2026/05/21"
    [row] = sync_rows(db)
    assert row.requested_at == "2026/05/21"


def test_O_on_the_investor_side_answers_with_the_new_date(board, db, users):
    from app import clock

    contact = vc(db, users)
    r = mark(board, contact.id, "O")
    assert r.status_code == 200, r.text
    assert r.json()["sourcing_requested_at"] == clock.today().isoformat()
    assert _requested(board, contact.id) == clock.today().isoformat()


def _xlsx_rows(content):
    import io

    import openpyxl

    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    head = next(i for i, r in enumerate(rows) if "이름" in r)
    return rows[head], rows[head + 1:]


def test_the_excel_carries_series_a_and_the_requested_date(board, db, users):
    """엑셀(`전체` · 명단별)에도 `Series A 딜 소싱` · `참여 요청일` 두 칸이 화면과
    같은 값으로 나란히 나간다."""
    linked = vc(db, users, name="홍길동")
    vc(db, users, name="김서연", phone="010-2222-3333")
    sc(db, vc_contact_id=linked.id, requested_at="2026-07-01 00:00:00")

    for url in ("/api/export/contacts.xlsx", f"/api/export/contacts.xlsx?sheet={SHEET}"):
        r = board.get(url)
        assert r.status_code == 200, (url, r.text[:200])
        head, rows = _xlsx_rows(r.content)
        at = head.index("Series A 딜 소싱")
        assert head[at + 1] == "참여 요청일", (url, head)
        got = {row[head.index("이름")]: (row[at] or "", row[at + 1] or "") for row in rows}
        assert got == {"홍길동": ("O", "2026-07-01"), "김서연": ("", "")}, (url, got)


# ── ⑤ [연결 풀기] ───────────────────────────────────────────────────────────

def unlink(client, row_id):
    return client.post(f"/api/sourcing/{row_id}/unlink")


def test_unlink_frees_the_link_and_keeps_both_rows(board, db, users, admin):
    """관리자가 푼다 — 두 줄 다 남고, 투자사 쪽 칸이 비고, 그 줄은 다시 후보로 뜬다.
    수정 로그에 `vc_contact_id: 번호 → 빈칸` 으로 남는다."""
    from app.models import EditLog, SourcingContact, VcContact

    contact = vc(db, users, name="김서연")          # 번호만 같다 — 후보
    row = sc(db)
    rid, cid = row.id, contact.id
    assert confirm(board, rid, cid).status_code == 200
    assert column(board, cid) == "O"

    login(board, admin.phone)
    r = unlink(board, rid)
    assert r.status_code == 200, r.text
    body = r.json()
    assert "vc_link" not in body
    assert [c["id"] for c in body["vc_picks"]] == [cid]       # 다시 후보로

    db.expire_all()
    assert db.get(SourcingContact, rid).vc_contact_id is None
    assert db.get(VcContact, cid) is not None                  # 지우지 않는다
    login(board, "01000000001")                                 # 그 투자사 줄의 담당
    assert column(board, cid) == ""
    logs = db.query(EditLog).filter_by(table_name="sourcing_contacts", row_id=rid,
                                       action="update").all()
    changes = [c for log in logs for c in json.loads(log.changes_json)]
    assert {"field": "vc_contact_id", "before": cid, "after": None} in changes


def test_unlink_is_for_the_admin_and_the_owner_only(board, db, users):
    """누가 푸나 — 관리자 · 이어진 투자사 줄의 담당 팀원(그 줄의 `user_id`). 딜 소싱
    `담당자` 칸에 이름이 적혀 있는 것만으로는 못 푼다 — 그 칸은 누구나 고친다.
    그 밖의 팀원에게는 단추가 안 서고, 보내도 막힌다(403)."""
    contact = vc(db, users, name="김서연")          # u1(강민준) 의 투자사 줄
    row = sc(db, vc_contact_id=contact.id, assignee_name="윤서아")
    rid = row.id

    login(board, "01000000002")                    # u2 — 딜 소싱 담당 칸에만 있다
    page = board.get(f"/sourcing?tab={SYNC}").text
    assert "js-vc-unlink" not in page
    r = unlink(board, rid)
    assert r.status_code == 403
    assert "관리자" in r.json()["detail"]
    assert sync_rows(db)[0].vc_contact_id == contact.id

    login(board, "01000000001")                    # u1 — 이어진 투자사 줄의 담당
    page = board.get(f"/sourcing?tab={SYNC}").text
    assert "js-vc-unlink" in page
    assert unlink(board, rid).status_code == 200
    assert sync_rows(db)[0].vc_contact_id is None

    other = vc(db, users, name="박지훈", phone="010-4444-5555", who="u2")
    sync_rows(db)[0].vc_contact_id = other.id     # u2 의 투자사 줄과 이어졌다
    db.commit()
    page = board.get(f"/sourcing?tab={SYNC}").text   # u1 — 그 투자사 줄의 주인이 아니다
    assert "js-vc-unlink" not in page
    r = unlink(board, rid)
    assert r.status_code == 403
    assert sync_rows(db)[0].vc_contact_id == other.id


def test_writing_my_name_in_the_assignee_cell_does_not_open_unlink(board, db, users):
    """`담당자` 칸은 로그인한 누구나 PATCH 로 고친다 — 거기 제 이름을 적고 [연결 풀기]를
    누르는 길이 막혀야 한다(남이 확정해 둔 잇기를 거두는 일이다)."""
    contact = vc(db, users, name="김서연")          # u1 의 투자사 줄
    row = sc(db, vc_contact_id=contact.id, assignee_name="강민준")
    rid = row.id

    login(board, "01000000002")                    # u2 — 담당도 주인도 아니다
    r = board.patch(f"/api/sourcing/{rid}", json={"assignee_name": users["u2"].name})
    assert r.status_code == 200, r.text
    assert sync_rows(db)[0].assignee_name == users["u2"].name
    assert "js-vc-unlink" not in board.get(f"/sourcing?tab={SYNC}").text
    r = unlink(board, rid)
    assert r.status_code == 403
    assert sync_rows(db)[0].vc_contact_id == contact.id


def test_unlink_refuses_rows_that_are_not_linked(board, db, users, admin):
    login(board, admin.phone)
    row = sc(db)
    r = unlink(board, row.id)
    assert r.status_code == 400
    assert "이어진" in r.json()["detail"]
    assert unlink(board, 999999).status_code == 404


# ── ⑥ 시트를 다시 가져와도 사람이 확정한 잇기는 남는다 ──────────────────────

def _reimport(tmp_path, monkeypatch, rows):
    """`scripts/import_sourcing.py --apply` 를 시리즈 A 갈래 한 장짜리 파일로 돌린다."""
    import sys

    import openpyxl

    from scripts import import_sourcing

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SYNC
    ws.append(["참여 요청일", "이름", "휴대폰", "회사"])
    for row in rows:
        ws.append(list(row))
    path = tmp_path / "sourcing.xlsx"
    wb.save(path)
    monkeypatch.setattr(sys, "argv", ["import_sourcing.py", str(path), "--apply"])
    import_sourcing.main()


def test_reimport_keeps_links_a_person_confirmed(board, db, users, tmp_path, monkeypatch):
    """확정한 줄은 이름 · 번호가 둘 다 맞지 않아도 다시 이어진다 — 후보로 다시 뜨지
    않는다. 저절로 이어지는 줄(이름 + 번호)은 지금 규칙으로 다시 잇는다. 번호만으로
    이어져 있던 옛 줄(이주 0093 · #28 전)도 남는다 — 0104 가 그 줄을 확정(1)으로
    적었다(`test_links_made_before_0104_count_as_confirmed`)."""
    confirmed = vc(db, users, name="김서연", phone="010-1234-5678")   # 번호만 같다
    auto = vc(db, users, name="박지훈", phone="010-3333-4444")
    old = vc(db, users, name="최옛날", phone="010-5555-6666")
    row = sc(db)                                                        # 홍길동
    assert confirm(board, row.id, confirmed.id).status_code == 200
    assert board.post(f"/api/sourcing?bucket={SYNC}",                  # 저절로 이어진다
                      json={"name": "박지훈", "phone": "010-3333-4444"}).status_code == 200
    # 0104 를 지난 옛 잇기 — 이름이 달라도 확정으로 적혀 있다
    sc(db, name="최옛줄", phone="010-5555-6666", position=3, vc_contact_id=old.id,
       vc_link_confirmed=1)
    linked = {r.name: r.vc_contact_id for r in sync_rows(db)}
    assert linked == {"홍길동": confirmed.id, "박지훈": auto.id, "최옛줄": old.id}

    _reimport(tmp_path, monkeypatch, [
        ("2026/05/21", "홍길동", "010-1234-5678", "가나다벤처스"),
        ("2026/05/22", "박지훈", "010-3333-4444", "가나다벤처스"),
        ("2026/05/23", "최옛줄", "010-5555-6666", "가나다벤처스")])

    rows = {r.name: r for r in sync_rows(db)}
    assert rows["홍길동"].vc_contact_id == confirmed.id
    assert rows["박지훈"].vc_contact_id == auto.id
    assert rows["최옛줄"].vc_contact_id == old.id
    assert picks(db, rows["홍길동"]) is None                  # 후보 자리 없음 — 이어졌다
    assert picks(db, rows["최옛줄"]) is None
    assert column(board, confirmed.id) == "O"
    assert column(board, old.id) == "O"

    # 되살린 잇기도 사람이 확정한 것이다 — 한 번 더 가져와도 남는다.
    _reimport(tmp_path, monkeypatch, [("2026/05/21", "홍길동", "010-1234-5678", "")])
    [again] = sync_rows(db)
    assert again.vc_contact_id == confirmed.id


def test_reimport_restores_only_to_the_same_person(board, db, users, tmp_path, monkeypatch):
    """확정해 둔 분이 시트에서 빠졌으면 잇지 않는다. 같은 이름이 둘이면 번호로 고른다."""
    confirmed = vc(db, users, name="김서연", phone="010-1234-5678")
    row = sc(db)                                                        # 홍길동 010-1234-5678
    assert confirm(board, row.id, confirmed.id).status_code == 200

    _reimport(tmp_path, monkeypatch, [("", "홍길동", "010-9999-0000", ""),
                                      ("", "홍길동", "010-1234-5678", "")])
    got = {r.phone: r.vc_contact_id for r in sync_rows(db)}
    assert got == {"010-9999-0000": None, "010-1234-5678": confirmed.id}

    _reimport(tmp_path, monkeypatch, [("", "박지훈", "010-7777-8888", "")])
    [left] = sync_rows(db)
    assert left.vc_contact_id is None


@pytest.mark.parametrize("gone", ["hidden", "not_an_investor_list"])
def test_reimport_does_not_restore_to_a_row_that_can_no_longer_be_linked(
        board, db, users, tmp_path, monkeypatch, gone):
    """확정해 둔 투자사 줄이 그사이 감춰졌거나 투자사로 세지 않는 명단으로 갔으면
    되살리지 않는다 — 저절로 잇기 · 후보와 같은 모집단(`_investors`)만 잇는다."""
    from app.models import SheetOwner, VcContact

    confirmed = vc(db, users, name="김서연", phone="010-1234-5678")   # 번호만 같다
    row = sc(db)
    assert confirm(board, row.id, confirmed.id).status_code == 200
    target = db.get(VcContact, confirmed.id)
    if gone == "hidden":
        target.is_hidden = 1
    else:
        db.add(SheetOwner(label="스타트업 명단", user_id=users["u1"].id,
                          layout="startup", is_hidden=1))
        target.source_sheet = "스타트업 명단"
    db.commit()

    _reimport(tmp_path, monkeypatch, [("", "홍길동", "010-1234-5678", "")])
    [again] = sync_rows(db)
    assert again.vc_contact_id is None


def test_links_made_before_0104_count_as_confirmed(tmp_path):
    """0104 가 칸을 더할 때 **이미 이어져 있던 줄은 확정(1)** 으로 적는다 — 이주(0093)와
    #28 전의 저절로 잇기(번호만)가 이은 줄이고, 지금까지 사람이 그대로 받아들여 쓴
    잇기다. 0 으로 두면 시트를 다시 가져올 때 조용히 풀렸다(`confirmed_links` 는 1 만
    기억한다). 안 이어진 줄은 0 이다."""
    import sqlite3

    from tests.test_migrations import _alembic

    db_path = tmp_path / "m.db"
    done = _alembic(db_path, "upgrade", "0103_template_apply_backups")
    assert done.returncode == 0, done.stdout + done.stderr
    # 운영 DB 는 이 칸이 없는 채로 0104 를 맞는다(0001 의 create_all 은 이미 칸이 있다)
    con = sqlite3.connect(db_path)
    try:
        con.execute("ALTER TABLE sourcing_contacts DROP COLUMN vc_link_confirmed")
        con.executemany(
            "INSERT INTO sourcing_contacts (id, bucket, position, name, room_verified, "
            "participation, vc_contact_id, created_at, updated_at) "
            "VALUES (?, ?, ?, ?, 'unverified', '참여', ?, '2026-10-01', '2026-10-01')",
            [(1, SYNC, 1, "옛 잇기", 7), (2, SYNC, 2, "안 이음", None),
             (3, MNA, 3, "다른 갈래", None)])
        con.commit()
    finally:
        con.close()

    done = _alembic(db_path, "upgrade", "0104_sourcing_vc_link_confirmed")
    assert done.returncode == 0, done.stdout + done.stderr
    con = sqlite3.connect(db_path)
    try:
        got = dict(con.execute("SELECT id, vc_link_confirmed FROM sourcing_contacts"))
    finally:
        con.close()
    assert got == {1: 1, 2: 0, 3: 0}


# ── 화면 — 후보 · 확정 · 읽기 칸 ────────────────────────────────────────────

JS_DIR = Path(__file__).resolve().parent / "js"


@pytest.mark.skipif(shutil.which("node") is None, reason="node 미설치 — 브라우저 로직 테스트 생략")
def test_the_candidate_box_and_confirm_button():
    """로컬에서는 `node tests/js/sourcing_vc_link_test.js` 로도 돈다."""
    result = subprocess.run([shutil.which("node"), str(JS_DIR / "sourcing_vc_link_test.js")],
                            capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stdout + result.stderr
