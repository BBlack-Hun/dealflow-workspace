"""`Series A 이상 딜 소싱` 탭의 [딜소개 문구 올리기] — 관리자가 파일로 `딜소개 문구` 를 채운다.

사용자 결정(2026-10-16):

  ① 관리자만 — 단추도 길도. 팀원 · 투자컨설턴트는 단추가 없고 403.
  ② 줄은 시트 가져오기와 **같은 함수**(`series_a_import.match`)로 맞춘다 — 열쇠 →
     열쇠 없는 줄의 지금 값 열쇠 → `(기업명, 수신 일자)`(양쪽 다 하나일 때만).
  ③ 바꾸는 것은 `딜소개 문구` 하나 — 줄을 만들지 않고, `import_key` 도 안 적는다.
  ④ 기본은 빈 칸만 채움 · [덮어쓰기] 를 켜야 다른 글을 덮음 · 파일의 빈 칸은
     아무것도 안 지움 · 같은 글은 그대로.
  ⑤ 미리보기는 아무것도 안 바꾸고, 반영은 파일에서 다시 센 것을 돌려준다.
  ⑥ 바꾼 줄은 수정 로그에 누가 바꿨는지와 함께 남는다.
  ⑦ 브라우저(node).

값은 전부 지어낸 것이다 — 저장소가 공개다.
"""
from __future__ import annotations

import csv
import io
import json
import shutil
import subprocess
from pathlib import Path

import pytest

from .conftest import DEMO_PASSWORD

ROOT = Path(__file__).resolve().parent.parent
URL = "/api/import/series-a-pitch"

HEAD = ["NO", "수신 일자", "Type", "담당팀원", "공유 심사역", "회사", "기업명",
        "사업분야 대분류", "소분류", "기업구분", "상세 내용", "딜소개 문구"]


def _line(day, sharer, firm, name, pitch, detail="", sector=""):
    return ["", day, "딜 소싱 네트워크", "강민준", sharer, firm, name,
            sector, "", "Series A", detail, pitch]


def _sheet(head=HEAD):
    return [
        head,
        # ① 열쇠로 맞는 줄 · 앱 문구가 빔 → 채움. 다른 칸(상세 · 대분류)은 파일에
        #    값이 있어도 안 들어가야 한다.
        _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", "애그 새 문구",
              detail="파일의 상세", sector="애그테크"),
        # ② 열쇠로 맞는 줄 · 앱에 다른 문구 → 기본은 둠, 덮어쓰기면 덮음.
        _line("2026-07-02", "김철수 수석", "다라인베스트", "샘플메디", "메디 새 문구"),
        # ③ 열쇠로 맞는 줄 · 같은 문구 → 그대로.
        _line("2026-07-03", "이영수 상무", "가나벤처스", "샘플바이오", "바이오 문구"),
        # ④ 열쇠 없는 줄 — 지금 값으로 셈한 열쇠가 같다(이름은 `(주)` 차이) → 채움.
        _line("2026-07-04", "박영희 이사", "마바벤처스", "샘플페이", "페이 새 문구"),
        # ⑤ 열쇠 없는 줄 — `(기업명, 수신 일자)` 만 같고 양쪽 다 하나 → 채움.
        _line("2026-07-05", "홍길동 이사", "가나벤처스", "샘플로봇", "로봇 새 문구"),
        # ⑥ 열쇠 없는 줄이 둘 — 어느 것인지 모른다 → 애매해서 건너뜀.
        _line("2026-07-06", "김서연 이사", "가나벤처스", "샘플푸드", "푸드 새 문구"),
        # ⑦ 파일의 문구가 빔 → 아무것도 안 지운다(덮어쓰기를 켜도).
        _line("2026-07-07", "박영희 이사", "가나벤처스", "샘플에너지", ""),
        # ⑧ 앱에 없는 딜 → 못 맞춘 줄. 줄을 만들지 않는다.
        _line("2026-07-08", "박영희 이사", "가나벤처스", "샘플없음", "없는 기업 문구"),
        # ⑨ ① 과 같은 딜이 또 — 앞 줄만 쓴다(시트 안 중복).
        _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", "겹친 줄 문구"),
        # 기업명이 빈 줄(메모) — 건너뛴다.
        ["", "", "", "", "", "", "", "", "", "", "메모 줄", ""],
    ]


def _key(day, sharer, firm, name):
    from app.services import series_a_import as svc

    return svc.import_key({"name": name, "received_at": day,
                           "shared_by": sharer, "source_firm": firm})


@pytest.fixture()
def table(db):
    """앱의 표 — 위 파일 줄들과 맞을 줄들. `{표시: 번호}` 를 돌려준다."""
    from app.models import SeriesACompany

    rows = {
        "agri": SeriesACompany(
            name="샘플애그", received_at="2026-07-01", shared_by="박영희 이사",
            source_firm="가나벤처스", detail="앱의 상세",
            import_key=_key("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그")),
        "medi": SeriesACompany(
            name="샘플메디", received_at="2026-07-02", shared_by="김철수 수석",
            source_firm="다라인베스트", deal_pitch="메디 옛 문구",
            import_key=_key("2026-07-02", "김철수 수석", "다라인베스트", "샘플메디")),
        "bio": SeriesACompany(
            name="샘플바이오", received_at="2026-07-03", shared_by="이영수 상무",
            source_firm="가나벤처스", deal_pitch="바이오 문구",
            import_key=_key("2026-07-03", "이영수 상무", "가나벤처스", "샘플바이오")),
        # 열쇠 없는 줄들 — 앱에서 더한 줄 · 열쇠 칸이 생기기 전에 들어온 줄.
        "pay": SeriesACompany(name="(주)샘플페이", received_at="2026-07-04",
                              shared_by="박영희 이사", source_firm="마바벤처스"),
        "robot": SeriesACompany(name="샘플로봇", received_at="2026-07-05"),
        "food1": SeriesACompany(name="샘플푸드", received_at="2026-07-06",
                                shared_by="한 사람"),
        "food2": SeriesACompany(name="샘플푸드", received_at="2026-07-06",
                                shared_by="다른 사람"),
        "energy": SeriesACompany(
            name="샘플에너지", received_at="2026-07-07", shared_by="박영희 이사",
            source_firm="가나벤처스", deal_pitch="에너지 문구",
            import_key=_key("2026-07-07", "박영희 이사", "가나벤처스", "샘플에너지")),
    }
    db.add_all(rows.values())
    db.commit()
    return {k: r.id for k, r in rows.items()}


@pytest.fixture()
def portal(db, users):
    """관리자 · 팀원 · 투자컨설턴트를 **따로** 로그인한 클라이언트."""
    from fastapi.testclient import TestClient

    from app.main import create_app
    from app.models import User
    from app.services import auth as auth_svc

    pw = auth_svc.hash_password(DEMO_PASSWORD)
    db.add_all([
        User(id=91, name="관리자시험", phone="01000000091", role="admin", password_hash=pw),
        User(id=92, name="컨설턴트시험", phone="01000000092", role="consultant",
             password_hash=pw),
    ])
    db.commit()
    app = create_app()

    def sign_in(phone):
        c = TestClient(app)
        c.post("/login", data={"phone": phone, "password": DEMO_PASSWORD})
        return c

    return {"admin": sign_in("01000000091"), "member": sign_in("01000000001"),
            "consultant": sign_in("01000000092")}


def _csv(rows) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _xlsx(tabs) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in tabs:
        ws = book.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _post(client, data: bytes, name="sheet.csv", **form):
    fields = {"dry_run": "true", "overwrite": "false", "sheet": ""}
    fields.update({k: str(v).lower() if isinstance(v, bool) else v for k, v in form.items()})
    return client.post(URL, files={"file": (name, data)}, data=fields)


def _snapshot(db):
    """표 전체를 칸까지 — 무엇이 바뀌었는지 견준다."""
    from app.models import SeriesACompany

    db.expire_all()
    cols = [c.key for c in SeriesACompany.__table__.columns]
    return {r.id: {c: getattr(r, c) for c in cols}
            for r in db.query(SeriesACompany).all()}


def _pitches(db):
    from app.models import SeriesACompany

    db.expire_all()
    return {r.id: r.deal_pitch for r in db.query(SeriesACompany).all()}


# ── ① 관리자만 ──────────────────────────────────────────────────────────────

def test_관리자에게만_단추와_창이_보인다(portal, table):
    admin = portal["admin"].get("/companies?tab=series-a").text
    for mark in ('id="sa-pitch-btn"', 'id="sa-pitch-panel"', "js/series_a_pitch.js",
                 'id="sa-pitch-overwrite"', "이미 적힌 문구도 덮어쓰기"):
        assert mark in admin, mark
    member = portal["member"].get("/companies?tab=series-a").text
    for mark in ("sa-pitch-btn", "sa-pitch-panel", "series_a_pitch.js"):
        assert mark not in member, f"팀원 화면에 {mark} 가 있습니다"


def test_관리자가_아니면_403_이고_아무것도_안_바뀐다(portal, table, db):
    before = _snapshot(db)
    data = _csv(_sheet())
    res = _post(portal["member"], data, dry_run=False)
    assert res.status_code == 403, res.text
    assert _post(portal["consultant"], data, dry_run=False).status_code == 403
    assert _snapshot(db) == before
    # 판정은 한 함수 — 화면과 라우터가 같이 지난다.
    src = (ROOT / "app" / "routers" / "series_a.py").read_text("utf-8")
    assert src.count("can_upload_pitch(user)") >= 2, "화면과 라우터가 다른 판정을 읽습니다"


# ── ② · ⑤ 미리보기 ──────────────────────────────────────────────────────────

def test_미리보기는_수를_세고_아무것도_안_바꾼다(portal, table, db):
    from app.models import EditLog

    before = _snapshot(db)
    logs = db.query(EditLog).count()
    res = _post(portal["admin"], _csv(_sheet()))
    assert res.status_code == 200, res.text
    d = res.json()
    assert d["dry_run"] is True
    assert (d["file_rows"], d["no_name"]) == (9, 1), d
    assert d["matched"] == 6, d                 # 애그 · 메디 · 바이오 · 페이 · 로봇 · 에너지
    assert (d["unmatched"], d["unmatched_with_pitch"]) == (1, 1), d
    assert d["ambiguous"] == 1 and d["duplicates"] == 1, d
    assert (d["changed"], d["filled"], d["overwritten"]) == (3, 3, 0), d
    assert (d["kept"], d["same"], d["blank"]) == (1, 1, 1), d

    changes = {c["name"]: c for c in d["changes"]}
    assert set(changes) == {"샘플애그", "(주)샘플페이", "샘플로봇"}, changes
    agri = changes["샘플애그"]
    assert (agri["received_at"], agri["shared_by"], agri["source_firm"]) == \
        ("2026-07-01", "박영희 이사", "가나벤처스")
    assert (agri["before"], agri["after"], agri["how"]) == ("", "애그 새 문구", "fill")
    assert agri["id"] == table["agri"]
    assert d["unmatched_rows"] == [{"name": "샘플없음", "received_at": "2026-07-08",
                                    "shared_by": "박영희 이사", "source_firm": "가나벤처스",
                                    "has_pitch": True}]
    assert [r["name"] for r in d["ambiguous_rows"]] == ["샘플푸드"]

    assert _snapshot(db) == before, "미리보기가 표를 바꿨습니다"
    db.expire_all()
    assert db.query(EditLog).count() == logs, "미리보기가 수정 로그를 남겼습니다"


def test_긴_목록은_잘라_싣고_수는_다_준다(portal, table, monkeypatch):
    monkeypatch.setattr("app.routers.data_io.DIFF_ROWS", 1)
    d = _post(portal["admin"], _csv(_sheet())).json()
    assert d["changed"] == 3 and len(d["changes"]) == 1, d


def test_문구는_줄여_싣는다(portal, table, db):
    from app.services.import_diff import VALUE_WIDTH

    long = "아주 긴 딜소개 문구 " * 30
    sheet = [HEAD, _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", long)]
    d = _post(portal["admin"], _csv(sheet)).json()
    after = d["changes"][0]["after"]
    assert after.endswith("…") and len(after) <= VALUE_WIDTH + 1, after


# ── ③ · ④ 반영 ──────────────────────────────────────────────────────────────

def test_반영은_빈_칸만_채우고_딜소개_문구_말고는_안_바꾼다(portal, table, db):
    before = _snapshot(db)
    res = _post(portal["admin"], _csv(_sheet()), dry_run=False)
    assert res.status_code == 200, res.text
    d = res.json()
    assert d["dry_run"] is False and d["changed"] == 3, d
    assert "딜소개 문구 3줄을 바꿨습니다" in d["summary"], d["summary"]

    after = _snapshot(db)
    assert set(after) == set(before), "줄이 새로 생기거나 사라졌습니다"
    pitch = {rid: after[rid]["deal_pitch"] for rid in after}
    assert pitch[table["agri"]] == "애그 새 문구"           # 앞 줄(겹친 뒤 줄은 버림)
    assert pitch[table["pay"]] == "페이 새 문구"            # 열쇠 없는 줄 — 지금 값 열쇠
    assert pitch[table["robot"]] == "로봇 새 문구"          # 열쇠 없는 줄 — (기업명, 수신 일자)
    assert pitch[table["medi"]] == "메디 옛 문구", "덮어쓰기를 안 켰는데 덮었습니다"
    assert pitch[table["bio"]] == "바이오 문구"
    assert pitch[table["energy"]] == "에너지 문구", "파일의 빈 칸이 문구를 지웠습니다"
    assert pitch[table["food1"]] is None and pitch[table["food2"]] is None, "애매한 줄을 골랐습니다"

    # 딜소개 문구 말고는 한 칸도 안 바뀐다 — 파일의 상세 · 대분류도, 열쇠(import_key)도.
    for rid in before:
        old = {k: v for k, v in before[rid].items() if k not in ("deal_pitch", "updated_at")}
        new = {k: v for k, v in after[rid].items() if k not in ("deal_pitch", "updated_at")}
        assert old == new, (rid, old, new)
    assert after[table["pay"]]["import_key"] is None
    assert after[table["agri"]]["detail"] == "앱의 상세"
    assert after[table["agri"]]["sector_major"] is None


def test_덮어쓰기를_켜면_다른_문구를_덮고_빈_칸은_여전히_안_지운다(portal, table, db):
    d = _post(portal["admin"], _csv(_sheet()), dry_run=False, overwrite=True).json()
    assert (d["changed"], d["filled"], d["overwritten"], d["kept"]) == (4, 3, 1, 0), d
    pitch = _pitches(db)
    assert pitch[table["medi"]] == "메디 새 문구"
    assert pitch[table["energy"]] == "에너지 문구", "덮어쓰기에서 파일의 빈 칸이 문구를 지웠습니다"
    assert pitch[table["bio"]] == "바이오 문구"


def test_반영은_파일에서_다시_세어_실제로_바꾼_것을_돌려준다(portal, table, db):
    """미리보기와 반영 사이에 누가 칸을 채웠으면 그 줄은 안 덮는다 — 반영이 다시 센다."""
    from app.models import SeriesACompany

    data = _csv(_sheet())
    assert _post(portal["admin"], data).json()["changed"] == 3
    row = db.get(SeriesACompany, table["robot"])
    row.deal_pitch = "그 사이 적은 문구"
    db.commit()
    d = _post(portal["admin"], data, dry_run=False).json()
    assert (d["changed"], d["kept"]) == (2, 2), d
    assert _pitches(db)[table["robot"]] == "그 사이 적은 문구"


def test_다시_올리면_바뀔_것이_없다(portal, table, db):
    data = _csv(_sheet())
    _post(portal["admin"], data, dry_run=False)
    d = _post(portal["admin"], data, dry_run=False).json()
    assert d["changed"] == 0 and d["same"] == 4, d


# ── ⑥ 수정 로그 ─────────────────────────────────────────────────────────────

def test_바꾼_줄은_누가_바꿨는지_수정_로그에_남는다(portal, table, db):
    from app.models import EditLog

    _post(portal["admin"], _csv(_sheet()), dry_run=False)
    db.expire_all()
    logs = [r for r in db.query(EditLog).all() if r.table_name == "series_a_companies"]
    assert sorted(r.row_id for r in logs) == sorted(
        [table["agri"], table["pay"], table["robot"]]), logs
    for log in logs:
        assert log.actor_user_id == 91 and log.scope == "shared" and log.action == "update"
        assert json.loads(log.changes_json) == [{"field": "deal_pitch", "changed": True}]
        assert log.path == URL


# ── ② 맞추기는 가져오기와 같은 함수 ─────────────────────────────────────────

def test_열쇠_없는_줄은_가져오기와_같은_줄에_맞는다(table, db):
    """이 길과 가져오기(`upsert`)가 **같은 `match`** 를 지난다 — 같은 파일은 두 길에서
    같은 줄에 닿는다. 가져오기는 그 줄들에 열쇠를 적고(linked), 이 길은 안 적는다."""
    from app.services import series_a_import as svc
    from app.services import series_a_pitch_upload as pitch

    records = svc.parse(_sheet()).records
    p = pitch.plan(db, records)
    filled = {c.row.id for c in p.changes}

    # 열쇠 없던 줄 중 가져오기가 열쇠를 적은 줄(새로 만든 `샘플없음` 은 빼고).
    loose = {table[k] for k in ("pay", "robot", "food1", "food2")}
    result = svc.upsert(db, records)
    linked = {r.id for r in svc.table(db) if r.import_key and r.id in loose}
    db.rollback()
    assert result.linked == 2 and result.ambiguous == 1 and result.duplicates == 1, result
    assert linked == {table["pay"], table["robot"]}
    assert linked <= filled
    assert len(p.ambiguous) == result.ambiguous

    src = (ROOT / "app" / "services" / "series_a_pitch_upload.py").read_text("utf-8")
    assert "sai.match(" in src and "sai.cell_plan(" in src
    src = (ROOT / "app" / "services" / "series_a_import.py").read_text("utf-8")
    assert "found = match(table(db), records)" in src


def test_맞추기는_줄을_건드리지_않는다(table, db):
    from app.services import series_a_import as svc

    rows = svc.table(db)
    found = svc.match(rows, svc.parse(_sheet()).records)
    assert sum(1 for *_x, linked in found.pairs if linked) == 2
    assert not db.dirty, "match 가 줄을 고쳤습니다"
    assert all(r.import_key is None for r in rows
               if r.id in (table["pay"], table["robot"]))


# ── 파일 ────────────────────────────────────────────────────────────────────

def test_xlsx_는_이_시트로_보이는_탭을_고른다(portal, table, db):
    other = [["NO", "기업명", "사업분야 대분류"], ["1", "샘플다른", "바이오"]]
    data = _xlsx([("IR 기업 현황", other), ("Series A 이상 기업", _sheet())])
    d = _post(portal["admin"], data, name="book.xlsx", dry_run=False).json()
    assert d["sheet_used"] == "Series A 이상 기업", d
    assert d["sheets"] == ["IR 기업 현황", "Series A 이상 기업"]
    assert d["changed"] == 3
    assert _pitches(db)[table["agri"]] == "애그 새 문구"


def test_이_시트로_보이는_탭이_둘이면_고르라고_한다(portal, table):
    data = _xlsx([("사본", _sheet()), ("Series A 이상 기업", _sheet())])
    res = _post(portal["admin"], data, name="book.xlsx")
    assert res.status_code == 400
    body = res.json()
    assert "탭" in body["detail"] and "골라" in body["detail"], body
    assert body["sheets"] == ["사본", "Series A 이상 기업"]

    res = _post(portal["admin"], data, name="book.xlsx", sheet="사본")
    assert res.status_code == 200 and res.json()["sheet_used"] == "사본"


def test_딜소개_문구_칸이_없으면_400(portal, table, db):
    before = _snapshot(db)
    head = [h for h in HEAD if h != "딜소개 문구"]
    sheet = [head, _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", "")[:-1]]
    res = _post(portal["admin"], _csv(sheet), dry_run=False)
    assert res.status_code == 400
    assert "딜소개 문구" in res.json()["detail"], res.json()
    assert _snapshot(db) == before


@pytest.mark.parametrize("firm_head", [None, "투자사"])
def test_회사_칸이_없으면_400_줄을_맞추는_칸이다(portal, table, db, firm_head):
    """줄은 가져오기 열쇠(기업명 · 수신 일자 · 공유 심사역 · **회사**)로 맞춘다. `회사` 칸
    없이 올리면 열쇠가 하나도 안 맞아 문구가 통째로 안 들어간다 — 조용히 '못 맞춘 줄'
    로 끝내지 않고 칸이 없다고 막는다. `투자사` 라는 옛 머리글은 같은 칸이다."""
    before = _snapshot(db)
    at = HEAD.index("회사")
    if firm_head is None:
        head = HEAD[:at] + HEAD[at + 1:]
        line = _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", "애그 새 문구")
        sheet = [head, line[:at] + line[at + 1:]]
        res = _post(portal["admin"], _csv(sheet), dry_run=False)
        assert res.status_code == 400
        assert "`회사`" in res.json()["detail"], res.json()
        assert _snapshot(db) == before
        return
    head = HEAD[:at] + [firm_head] + HEAD[at + 1:]
    sheet = [head, _line("2026-07-01", "박영희 이사", "가나벤처스", "샘플애그", "애그 새 문구")]
    res = _post(portal["admin"], _csv(sheet), dry_run=False)
    assert res.status_code == 200, res.text
    assert _pitches(db)[table["agri"]] == "애그 새 문구"


def test_큰_파일은_열어_보기_전에_막고_워크북은_한_번만_연다(portal, table, monkeypatch):
    """10MB 상한은 파일을 **열기 전에** 본다 — 탭 이름을 읽으려고 먼저 열면 상한이
    지키려던 메모리를 이미 쓴다. 상한 안의 파일도 워크북은 한 번만 연다."""
    from app.services import spreadsheet as sp

    opened = []
    real = sp._load_workbook
    monkeypatch.setattr(sp, "_load_workbook", lambda data: opened.append(1) or real(data))

    res = _post(portal["admin"], b"x" * (sp.MAX_UPLOAD_BYTES + 1), name="big.xlsx")
    assert res.status_code == 400
    assert "MB" in res.json()["detail"]
    assert opened == []

    data = _xlsx([("IR 기업 현황", [["NO", "기업명"]]), ("Series A 이상 기업", _sheet())])
    res = _post(portal["admin"], data, name="book.xlsx")
    assert res.status_code == 200, res.text
    assert res.json()["sheets"] == ["IR 기업 현황", "Series A 이상 기업"]
    assert opened == [1]


def test_머리글_줄이_없으면_400(portal, table):
    res = _post(portal["admin"], _csv([["기업명", "딜소개 문구"], ["샘플애그", "문구"]]))
    assert res.status_code == 400
    assert "공유 심사역" in res.json()["detail"], res.json()


@pytest.mark.parametrize("name, content", [
    ("broken.xlsx", b"this is not a zip"),
    ("old.xls", b"\xd0\xcf\x11\xe0 old excel"),
    ("empty.csv", b""),
])
def test_못_읽는_파일은_400(portal, table, name, content):
    res = _post(portal["admin"], content, name=name)
    assert res.status_code == 400, res.text
    assert res.json()["detail"]


def test_바이트로_읽어도_탭_고르기는_가져오기와_같다():
    """`read_upload`(화면) · `read_rows`(스크립트)가 같은 `pick_tab` 을 지난다."""
    from app.services import series_a_import as svc

    data = _xlsx([("메모", [["아무거나"]]), ("Series A 이상 기업", _sheet())])
    used, rows = svc.read_upload("book.xlsx", data)
    assert used == "Series A 이상 기업"
    assert svc.parse(rows).records[0]["name"] == "샘플애그"
    used, rows = svc.read_upload("sheet.csv", _csv(_sheet()))
    assert used == "" and len(svc.parse(rows).records) == 9
    with pytest.raises(ValueError):
        svc.read_upload("book.xlsx", data, sheet="없는 탭")


# ── ⑦ 브라우저 ──────────────────────────────────────────────────────────────

@pytest.mark.skipif(shutil.which("node") is None,
                    reason="node 미설치 — 브라우저 로직 테스트 생략")
def test_브라우저_쪽은_파일을_그대로_돌려_본다():
    script = ROOT / "tests" / "js" / "series_a_pitch_test.js"
    out = subprocess.run([shutil.which("node"), str(script)],
                         capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stdout + out.stderr
