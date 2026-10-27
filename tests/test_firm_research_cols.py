"""투자사 회사 단위 칸 — `챗GPT 조사 소유펀드`(옛 `소유 펀드`) 이름 · `펀드 종류` · `소분류`(0124).

  1. 이름 바꾸기 — 화면은 새 이름, 엑셀 올리기는 옛 머리글도 받는다
  2. 새 칸 둘 — 저장 · 같은 회사 줄 전파 · 표 · 수정창 · 엑셀의 자리(`선호 투자분야` 바로 왼쪽)
  3. `소분류` 검사 — 분야 기준 목록의 소분류만 · 5개까지 · 빈 값은 지움
  4. 업로드 — 빈 칸은 그대로 · 어긋난 줄은 안 넣고 까닭을 보인다
  5. 0124 오르내리기

회사 · 펀드 · 사람 이름은 전부 지어낸 것이다 — 저장소가 공개다.
"""
from __future__ import annotations

import io
import re
import sqlite3

import pytest

from tests.conftest import DEMO_PASSWORD

from .test_migrations import _alembic

TIP = "2025년 이후 기사 기준 인터넷 조사"


def _contact(db, user_id, name, firm):
    from app.models import VcContact

    row = VcContact(user_id=user_id, name=name, firm=firm)
    db.add(row)
    db.commit()
    return row


@pytest.fixture()
def member(client, db, users):
    from app.services import taxonomy

    taxonomy.ensure_seed(db)
    client.post("/login", data={"phone": "01000000001", "password": DEMO_PASSWORD})
    return client


@pytest.fixture()
def admin(client, db, users):
    from app.services import taxonomy

    taxonomy.ensure_seed(db)
    users["u2"].role = "admin"
    db.commit()
    client.post("/login", data={"phone": "01000000002", "password": DEMO_PASSWORD})
    return client


def _xlsx(rows):
    import openpyxl

    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _heads(html):
    head = html[html.index('id="contacts-table"'):html.index("</thead>")]
    cells = re.findall(r"<th\b([^>]*)>(.*?)</th>", head, re.S)
    return [(attrs, re.sub(r"<[^>]+>", "", cell).strip()) for attrs, cell in cells]


# ── 1. 이름 바꾸기 ─────────────────────────────────────────────────────────

def test_화면은_새_이름이다(member, db, users):
    _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    html = member.get("/contacts?sheet=all").text
    names = [n for _a, n in _heads(html)]
    assert "챗GPT 조사 소유펀드" in names and "소유 펀드" not in names
    # 수정창 이름도 표 머리글과 같다.
    assert "<span>챗GPT 조사 소유펀드</span>" in html
    assert "<span>소유 펀드</span>" not in html


@pytest.mark.parametrize("head", ["소유 펀드", "챗GPT 조사 소유펀드"])
def test_업로드는_옛_머리글도_새_머리글도_받는다(head):
    from app.services import investor_firms as f

    parsed = f.parse_upload([["투자사명", head], ["테스트벤처스", "테스트 1호"]])
    assert [(p.name, p.values) for p in parsed] == [("테스트벤처스", {"owned_funds": "테스트 1호"})]


@pytest.mark.parametrize("head", ["소유 펀드", "챗GPT 조사 소유펀드"])
def test_현황_업로드는_회사_단위_머리글이_있어도_다른_칸을_안_건드린다(head):
    """현황 업로드(`parse_sheet_a`)는 회사 단위 칸을 읽지 않는다 — 그 값은 관리자 업로드로만.
    옛 · 새 머리글과 새 칸 둘이 있어도 이웃 칸(라운드 · 단계 · 분야)이 제자리다."""
    from app.services.sheet_import import parse_sheet_a

    out = parse_sheet_a([
        ["이름", "회사", "라운드 사이즈(투자운영금액)", head, "선호 투자단계",
         "펀드 종류", "소분류", "선호 투자분야"],
        ["가나다", "테스트벤처스", "10~30억", "테스트 1호", "Seed", "청년창업", "로봇", "AI SaaS"],
    ], 2026)
    c = out.contacts[0]
    assert (c.round_size, c.stages, c.sectors) == ("10~30억", "Seed", ["AI SaaS"])
    assert out.activity_columns == []


# ── 2. 새 칸 둘 ────────────────────────────────────────────────────────────

def test_한_줄에서_적으면_같은_회사의_모든_줄에_선다(member, db, users):
    from app.routers.contacts import contact_rows

    a = _contact(db, users["u1"].id, "가나다", "(주)테스트벤처스")
    b = _contact(db, users["u1"].id, "라마바", "테스트 벤처스")
    other = _contact(db, users["u1"].id, "사아자", "테스트인베스트")

    r = member.patch(f"/api/contacts/{a.id}", json={"fund_types": "모태 문화계정,  청년창업"})
    assert r.status_code == 200, r.text
    assert r.json()["fund_type_list"] == ["모태 문화계정", "청년창업"]
    r = member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": "로봇, 의료기기"})
    assert r.status_code == 200, r.text
    assert r.json()["sub_sectors"] == "로봇, 의료기기"
    assert r.json()["sub_sector_list"] == ["로봇", "의료기기"] and r.json()["firm_key"]

    rows = {row["name"]: row for row in contact_rows(db, users["u1"])}
    assert rows["라마바"]["fund_types"] == "모태 문화계정,  청년창업"
    assert rows["라마바"]["sub_sectors"] == "로봇, 의료기기"
    assert rows["사아자"]["fund_types"] == "" and rows["사아자"]["sub_sectors"] == ""
    got = member.get(f"/api/contacts/{b.id}").json()["contact"]
    assert (got["fund_types"], got["sub_sectors"]) == ("모태 문화계정,  청년창업", "로봇, 의료기기")
    assert member.get(f"/api/contacts/{other.id}").json()["contact"]["sub_sectors"] == ""


def test_남의_줄은_못_고친다(member, db, users):
    theirs = _contact(db, users["u2"].id, "가나다", "테스트벤처스")
    assert member.patch(f"/api/contacts/{theirs.id}",
                        json={"fund_types": "청년창업"}).status_code == 404


def test_회사명이_없는_줄에는_적을_수_없다(member, db, users):
    a = _contact(db, users["u1"].id, "가나다", None)
    r = member.patch(f"/api/contacts/{a.id}", json={"fund_types": "청년창업"})
    assert r.status_code == 400 and "펀드 종류" in r.json()["detail"]


def test_표에서_선호_투자분야_바로_왼쪽이고_풍선말이_있다(member, db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    f.set_funds(db, "테스트벤처스", "청년창업, 소부장", users["u1"].id, "fund_types")
    f.set_funds(db, "테스트벤처스", "로봇", users["u1"].id, "sub_sectors")
    db.commit()
    html = member.get("/contacts?sheet=all").text
    heads = _heads(html)
    names = [n for _a, n in heads]
    at = names.index("선호 투자분야")
    assert names[at - 2:at] == ["펀드 종류", "소분류"]
    for attrs, name in heads[at - 2:at]:
        assert f'title="{TIP}"' in attrs, name
    assert 'data-filters="fund_type:펀드 종류"' in html
    assert 'data-filters="sub_sector:소분류"' in html
    row = re.search(r'id="contacts-table".*?<tbody>(.*?)</tr>', html, re.S).group(1)
    fields = re.findall(r'data-field="([a-z_]+)"', row)
    at = fields.index("sectors")
    assert fields[at - 2:at] == ["fund_types", "sub_sectors"]
    assert 'data-f-fund_type="청년창업|소부장"' in row
    assert 'data-f-sub_sector="로봇"' in row
    # 수정창도 같은 차례 — `선호 투자단계` · `펀드 종류` · `소분류` · `선호 투자분야`.
    panel = re.findall(r'id="f-(stages|fund_types|sub_sectors|sectors)"', html)
    assert panel == ["stages", "fund_types", "sub_sectors", "sectors"]


def test_달마다_늘어나는_명단의_수정창에도_같은_자리다():
    from app.services import contact_columns as cc

    cols = cc.INVESTOR_MONTHLY_LAYOUT.extra
    keys = [c.key for c in cols]
    at = keys.index("sectors")
    assert keys[at - 2:at] == ["fund_types", "sub_sectors"]
    assert {c.key: c.label for c in cols}["owned_funds"] == "챗GPT 조사 소유펀드"


def test_엑셀_내려받기의_차례와_값(member, db, users):
    import openpyxl

    a = _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    member.patch(f"/api/contacts/{a.id}", json={
        "owned_funds": "테스트 1호", "fund_types": "청년창업", "sub_sectors": "로봇"})
    r = member.get("/api/export/contacts.xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    head = [c.value for c in ws[1]]
    at = head.index("챗GPT 조사 소유펀드")
    assert head[at - 1:at + 5] == ["라운드 사이즈(투자운영금액)", "챗GPT 조사 소유펀드",
                                   "선호 투자단계", "펀드 종류", "소분류", "선호 투자분야"]
    line = next(row for row in ws.iter_rows(min_row=2, values_only=True) if "가나다" in row)
    assert (line[at], line[at + 2], line[at + 3]) == ("테스트 1호", "청년창업", "로봇")


# ── 3. 소분류 검사 ─────────────────────────────────────────────────────────

def test_소분류는_목록의_글자로_맞춰_적는다(member, db, users):
    from app.services import investor_firms as f

    a = _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    r = member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": "ai saas, AI SaaS ,로봇"})
    assert r.status_code == 200, r.text
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == "AI SaaS, 로봇"


@pytest.mark.parametrize("value, words", [
    ("로봇, 없는분야", ["없는분야"]),
    ("AI·SaaS·데이터", ["AI·SaaS·데이터"]),            # 대분류는 소분류가 아니다
    ("로봇, 센서, 광학, 반도체, 배터리, 의료기기", ["5개"]),
])
def test_목록_밖이거나_5개를_넘으면_400이고_안_바뀐다(member, db, users, value, words):
    from app.services import investor_firms as f

    a = _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    assert member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": "로봇"}).status_code == 200
    r = member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": value, "memo": "같이 보낸 메모"})
    assert r.status_code == 400
    assert "소분류" in r.json()["detail"]
    for word in words:
        assert word in r.json()["detail"]
    db.expire_all()
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == "로봇"
    assert a.memo is None, "막힌 요청의 다른 칸이 저장됐다"


def test_소분류_5개는_된다_빈_값은_지운다(member, db, users):
    from app.services import investor_firms as f

    a = _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    five = "로봇, 센서, 광학, 반도체, 배터리"
    assert member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": five}).status_code == 200
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == five
    r = member.patch(f"/api/contacts/{a.id}", json={"sub_sectors": ""})
    assert r.status_code == 200 and r.json()["sub_sector_list"] == []
    db.expire_all()
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == ""


# ── 4. 업로드 ──────────────────────────────────────────────────────────────

def _upload(admin, rows, dry_run):
    files = {"file": ("firms.xlsx", _xlsx(rows), "application/octet-stream")}
    return admin.post("/api/import/owned-funds", files=files,
                      data={"dry_run": "true" if dry_run else "false"})


def test_업로드의_빈_칸은_그대로_두고_어긋난_줄은_안_넣는다(admin, db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    _contact(db, users["u1"].id, "라마바", "테스트인베스트")
    _contact(db, users["u1"].id, "사아자", "테스트캐피탈")
    for firm in ("테스트벤처스", "테스트인베스트", "테스트캐피탈"):
        f.set_funds(db, firm, "옛 펀드", users["u1"].id)
        f.set_funds(db, firm, "옛 종류", users["u1"].id, "fund_types")
        f.set_funds(db, firm, "센서", users["u1"].id, "sub_sectors")
    db.commit()
    rows = [["투자사명", "챗GPT 조사 소유펀드", "펀드 종류", "소분류"],
            ["테스트벤처스", "", "청년창업", ""],              # 펀드 종류만 바뀐다
            ["테스트인베스트", "새 펀드", "", "없는분야"],      # 소분류가 목록 밖 — 통째로 안 넣는다
            ["테스트캐피탈", "", "", ""]]                        # 다 비었다 — 건너뜀

    preview = _upload(admin, rows, True)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["counts"]["changes"] == 1 and body["counts"]["errors"] == 1
    assert body["counts"]["blank"] == 1
    assert body["matched"][0]["after"] == {"fund_types": "청년창업"}
    assert body["matched"][0]["before"] == {"fund_types": "옛 종류"}
    assert body["errors"][0]["name"] == "테스트인베스트" and "없는분야" in body["errors"][0]["reason"]
    assert body["labels"]["sub_sectors"] == "소분류"

    done = _upload(admin, rows, False)
    assert done.status_code == 200, done.text
    assert done.json()["changed"] == 1
    db.expire_all()
    assert f.funds_of(db, "테스트벤처스", "fund_types") == "청년창업"
    assert f.funds_of(db, "테스트벤처스") == "옛 펀드", "빈 칸이 값을 지웠다"
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == "센서", "빈 칸이 값을 지웠다"
    assert f.funds_of(db, "테스트인베스트") == "옛 펀드", "어긋난 줄의 다른 칸을 넣었다"
    assert f.funds_of(db, "테스트캐피탈", "fund_types") == "옛 종류"


def test_업로드의_소분류도_5개까지고_목록_글자로_맞춘다(admin, db, users):
    from app.services import investor_firms as f

    _contact(db, users["u1"].id, "가나다", "테스트벤처스")
    _contact(db, users["u1"].id, "라마바", "테스트인베스트")
    rows = [["투자사명", "소분류"],
            ["테스트벤처스", "ai saas, 로봇"],
            ["테스트인베스트", "로봇, 센서, 광학, 반도체, 배터리, 의료기기"]]
    done = _upload(admin, rows, False)
    assert done.status_code == 200, done.text
    assert done.json()["errors"] == 1
    db.expire_all()
    assert f.funds_of(db, "테스트벤처스", "sub_sectors") == "AI SaaS, 로봇"
    assert f.funds_of(db, "테스트인베스트", "sub_sectors") == ""


def test_업로드_머리글이_없으면_새_이름으로_안내한다():
    from app.services import investor_firms as f

    with pytest.raises(f.UploadError, match="챗GPT 조사 소유펀드"):
        f.parse_upload([["회사", "펀드"]])


# ── 5. 0124 오르내리기 ─────────────────────────────────────────────────────

BEFORE = "0123_session_token_hash"


def _cols(db):
    con = sqlite3.connect(db)
    try:
        return {r[1] for r in con.execute('PRAGMA table_info("investor_firms")')}
    finally:
        con.close()


def test_0124_는_칸_둘을_붙이고_되돌리면_뗀다(tmp_path):
    db = tmp_path / "old.db"
    assert _alembic(db, "upgrade", BEFORE).returncode == 0
    # 운영 DB 의 모양 — 0001 의 `create_all()` 은 지금 모델이라 칸이 이미 있다.
    con = sqlite3.connect(db)
    try:
        con.execute("ALTER TABLE investor_firms DROP COLUMN fund_types")
        con.execute("ALTER TABLE investor_firms DROP COLUMN sub_sectors")
        con.execute("INSERT INTO investor_firms (firm_key, display_name, owned_funds, "
                    "created_at, updated_at) VALUES ('테스트벤처스', '테스트벤처스', "
                    "'테스트 1호', '2026-10-27T00:00:00+09:00', '2026-10-27T00:00:00+09:00')")
        con.commit()
    finally:
        con.close()
    assert {"fund_types", "sub_sectors"} & _cols(db) == set()

    up = _alembic(db, "upgrade", "head")
    assert up.returncode == 0, up.stdout + up.stderr
    assert {"fund_types", "sub_sectors"} <= _cols(db)

    down = _alembic(db, "downgrade", BEFORE)
    assert down.returncode == 0, down.stdout + down.stderr
    assert {"fund_types", "sub_sectors"} & _cols(db) == set()
    con = sqlite3.connect(db)
    try:
        assert con.execute("SELECT owned_funds FROM investor_firms").fetchall() == [("테스트 1호",)]
    finally:
        con.close()
